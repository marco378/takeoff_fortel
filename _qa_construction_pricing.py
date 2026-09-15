"""Portal QA for per-construction pricing, driven through the real assessor flow.

Uploads a real Radlett Surface Finishes Plan, includes every offered construction exactly as
an assessor would, then checks what the ASSESSOR sees -- regions, flags, the costing card --
and what the CLIENT gets: the draft quotation's priced rows and its combined area line.

Both JOBS_FILE and DRAWINGS_DIR point at scratch. DRAWINGS_DIR defaults to the repo's own
drawings/, so a QA upload without it lands in the corpus and the next robustness run reports
extra files, which reads exactly like a code change moved sheets.
"""
import json, os, shutil, subprocess, sys, tempfile, time, urllib.request, urllib.error

PORT = int(os.environ.get("QA_PORT", "8791"))
BASE = f"http://127.0.0.1:{PORT}"
SHEET = "drawings/inderjit_13sep/radlett_1_Surface_Finishes_Plan.pdf"
SCRATCH = tempfile.mkdtemp(prefix="qa_construction_")
SHOTS = os.path.join(SCRATCH, "shots")
os.makedirs(SHOTS, exist_ok=True)
FAILS = []


def ck(name, cond, got=""):
    print(f"  [{'PASS' if cond else 'FAIL'}] {name} {got}")
    if not cond:
        FAILS.append(name)


def get(path, raw=False):
    with urllib.request.urlopen(BASE + path, timeout=120) as r:
        body = r.read()
    return body if raw else json.loads(body)


def post(path, payload):
    req = urllib.request.Request(
        BASE + path, data=json.dumps(payload).encode(),
        headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            return r.status, json.loads(r.read())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read() or b"{}")


env = dict(os.environ,
           APPROVAL_PORT=str(PORT),
           JOBS_FILE=os.path.join(SCRATCH, "jobs.json"),
           DRAWINGS_DIR=os.path.join(SCRATCH, "drawings"),
           LEARNING_ENVIRONMENT="qa")
os.makedirs(env["DRAWINGS_DIR"], exist_ok=True)
server = subprocess.Popen([".venv/bin/python", "approval_server.py"], env=env,
                          stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
try:
    for _ in range(120):
        try:
            urllib.request.urlopen(BASE + "/status", timeout=2).read()
            break
        except Exception:
            time.sleep(0.5)
    else:
        raise SystemExit("server never came up")

    # Upload through the real endpoint: field name is `pdf`, and the response is 202 async.
    boundary = "----qaboundary"
    with open(SHEET, "rb") as fh:
        pdf = fh.read()
    parts = []
    for key, value in (("project_name", "QA Construction Pricing"),
                       ("project_ref", "QA-CP-001")):
        parts.append(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{key}\"\r\n\r\n"
                     f"{value}\r\n".encode())
    parts.append(f"--{boundary}\r\nContent-Disposition: form-data; name=\"pdf\"; "
                 f"filename=\"radlett_sfp.pdf\"\r\n"
                 f"Content-Type: application/pdf\r\n\r\n".encode() + pdf + b"\r\n")
    parts.append(f"--{boundary}--\r\n".encode())
    req = urllib.request.Request(
        BASE + "/upload", data=b"".join(parts),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"}, method="POST")
    with urllib.request.urlopen(req, timeout=300) as r:
        accepted = json.loads(r.read())
    job_id = accepted.get("job_id") or accepted.get("id")
    ck("upload accepted", bool(job_id), f"-> {job_id}")

    for _ in range(600):
        job = get(f"/job/{job_id}")
        if job.get("status") != "processing":
            break
        time.sleep(1)
    ck("job finished processing", job.get("status") != "processing", f"-> {job.get('status')}")

    regions = job.get("yard_regions") or (job.get("result") or {}).get("yard_regions") or []
    ck("the sheet offers constructions to the assessor", len(regions) > 1, f"-> {len(regions)}")
    print("     offered:", [(r.get("region_id"), r.get("surface_name"),
                             r.get("depth_mm"), r.get("area_m2")) for r in regions])

    status, body = post(f"/yard-regions/{job_id}", {"decisions": [
        {"region_id": str(r["region_id"]), "action": "keep"} for r in regions]})
    ck("including every region is accepted", status == 200, f"-> {status} {str(body)[:200]}")

    job = get(f"/job/{job_id}")
    flags = list(job.get("flags") or [])
    breakdown = [f for f in flags if "INCLUDED BY CONSTRUCTION" in str(f)]
    ck("the assessor is shown the per-construction breakdown", bool(breakdown))
    for f in breakdown:
        print("     flag:", f)

    costing = job.get("costing") or {}
    ck("the costing card is priced per construction",
       bool(costing.get("constructions")), f"-> {len(costing.get('constructions') or [])}")
    print("     card:", costing.get("area_m2"), costing.get("rate"),
          costing.get("total_gbp"), costing.get("constructions_note"))
    depths = {c.get("depth_mm") for c in (costing.get("constructions") or [])}
    ck("more than one thickness is priced on this sheet", len(depths) > 1, f"-> {sorted(d for d in depths if d)}")
    card_sum = sum(float(c.get("total_gbp") or 0) for c in (costing.get("constructions") or []))
    ck("the card total is the sum of its constructions",
       abs(card_sum - float(costing.get("total_gbp") or 0)) < 1.0,
       f"-> {card_sum:,.2f} vs {costing.get('total_gbp')}")

    q = get(f"/quotation/{job_id}.json")
    rows = [(li["description"], li["qty"], li.get("rate")) for li in q["line_items"]
            if li.get("line_role") == "concrete_slab"]
    print("     quote rows:")
    for r in rows:
        print("      ", r)
    ck("the quotation prints one slab row per construction", len(rows) == len(depths),
       f"-> {len(rows)} rows for {len(depths)} thicknesses")
    ck("the quotation's rows carry different rates",
       len({r for _d, _q, r in rows}) == len(rows), f"-> {[r for _d, _q, r in rows]}")
    ck("the quotation declares the combined total is a sum",
       any("separate constructions" in d for d in q["declarations"]))

    xlsx = get(f"/quotation/{job_id}.xlsx", raw=True)
    xpath = os.path.join(SCRATCH, "quote.xlsx")
    open(xpath, "wb").write(xlsx)
    from openpyxl import load_workbook
    ws = load_workbook(xpath)[load_workbook(xpath).sheetnames[0]]
    labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)
              if isinstance(ws.cell(r, 1).value, str) and ws.cell(r, 1).value.startswith("Total")]
    ck("the workbook ends the section with the combined area line",
       any("Yard Slab Area" in str(l) for l in labels), f"-> {[l for l in labels if 'Area' in str(l)]}")

    marked = get(f"/marked-pdf/{job_id}.pdf", raw=True)
    ck("the marked PDF still exports", marked[:4] == b"%PDF", f"-> {len(marked)} bytes")
    open(os.path.join(SHOTS, "marked.pdf"), "wb").write(marked)

    # What the assessor and the client actually look at, as pictures.
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1440, "height": 1100})
        page.goto(f"{BASE}/portal", wait_until="networkidle")
        # The portal is one page with a job list; the assessor's detail view is whatever
        # selectJob renders, so drive that rather than guessing at a per-job URL.
        page.wait_for_timeout(1500)
        page.evaluate(f"selectJob({job_id!r})")
        page.wait_for_timeout(3000)
        page.screenshot(path=os.path.join(SHOTS, "01_assessor_after_include.png"),
                        full_page=True)
        # The flag list is what the assessor actually reads about this decision, and it sits
        # below the fold. A screenshot that stops at the measurement is evidence of the number,
        # not of the breakdown, so scroll to the flags and capture them too.
        page.evaluate("document.getElementById('flagList')"
                      ".scrollIntoView({block:'center'})")
        page.wait_for_timeout(800)
        page.screenshot(path=os.path.join(SHOTS, "03_assessor_flags.png"))
        flag_text = page.evaluate("document.getElementById('flagList').innerText")
        ck("the breakdown flag is on the assessor's screen, not just in the API",
           "INCLUDED BY CONSTRUCTION" in flag_text)
        ck("the stale 'none counted' offer text is gone from the screen",
           "NONE COUNTED" not in flag_text.upper()
           and "no area emitted" not in flag_text)
        page.goto(f"{BASE}/quotation/{job_id}.html", wait_until="networkidle")
        page.screenshot(path=os.path.join(SHOTS, "02_quotation_per_construction.png"),
                        full_page=True)
        browser.close()
    for name in ("01_assessor_after_include.png", "02_quotation_per_construction.png",
                 "03_assessor_flags.png"):
        path = os.path.join(SHOTS, name)
        size = os.path.getsize(path) if os.path.exists(path) else 0
        ck(f"screenshot captured: {name}", size > 20000, f"-> {size} bytes")
    print("     screenshots:", SHOTS)
finally:
    server.terminate()
    try:
        server.wait(timeout=10)
    except Exception:
        server.kill()

print(f"\n{'ALL PASS' if not FAILS else 'FAILURES: ' + ', '.join(FAILS)}")
print("scratch:", SCRATCH)
sys.exit(1 if FAILS else 0)
