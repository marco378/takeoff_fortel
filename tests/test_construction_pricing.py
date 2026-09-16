#!/usr/bin/env python3
"""Per-construction pricing: one quote line per construction, summed at the end.

Sections in this module (printed in this order):
  - per-construction pricing

Imported by ci_tests.py, which owns the running total. Do not reorder.

The landmine these checks exist for: a Surface Finishes Plan names several constructions at
several thicknesses, and until 15 Sep 2026 the quotation priced their combined area at ONE
depth. The two things that must stay true are opposites, so both are checked end to end --
different thicknesses must produce different rates on separate rows, and constructions that
all share the zone's own thickness must produce EXACTLY what a job with no breakdown produces.
"""
import copy as _copy, io as _io
from openpyxl import load_workbook as _load_workbook

from tests import ck

print("per-construction pricing")
from quotation import (generate_quotation, quotation_xlsx, quotation_text,
                       PROVISIONAL_LABEL)
import construction_pricing as _cp

_SPEC = {"depth_mm": 200, "conc_rate": 120.0, "conc_wastage": 0.05, "mesh": "A393",
         "layers": 2, "steel_rate_t": 900.0, "steel_wastage": 0.05, "lap_acc": 0.15,
         "dpm": 1.2, "curing": 0.8, "labour": 12.0, "trim": 1.5, "margin": 0.1}
_ZONE_RATE, _ = _cp.rate_buildup(**{k: _SPEC[k] for k in _cp.RATE_FIELDS})


def _job(constructions, area_m2=23234.0):
    zone = {"category": "external_yard", "area_m2": area_m2}
    if constructions is not None:
        zone["constructions"] = _copy.deepcopy(constructions)
    return {
        "file": "WP5-SFP.pdf", "pdf_path": "drawings/WP5-SFP.pdf", "area_m2": area_m2,
        "zones": [zone],
        "costing": {"area_m2": area_m2, "rate": _ZONE_RATE,
                    "total_gbp": round(area_m2 * _ZONE_RATE, 2), "spec": dict(_SPEC),
                    "assumed": True, "breakdown": {}},
    }


_THREE = [
    {"name": "HGV Slab Construction", "depth_mm": 200, "detail_ref": "DD-C-1010",
     "area_m2": 4000.0, "region_ids": ["surface-finish-1"]},
    {"name": "Container Slab Construction", "depth_mm": 375, "detail_ref": "DD-C-1020",
     "area_m2": 17689.3, "region_ids": ["surface-finish-2", "surface-finish-2-part-2"]},
    {"name": "Rail Crossing", "depth_mm": 225, "detail_ref": "DD-C-1030",
     "area_m2": 1544.7, "region_ids": ["surface-finish-3"]},
]


def _concrete_rows(q):
    return [(li["description"], li["qty"], li["rate"]) for li in q["line_items"]
            if li.get("line_role") == "concrete_slab"]


_q3 = generate_quotation(_job(_THREE), project="Radlett WP5", client="Fortel", ref="T-CP")
_rows3 = _concrete_rows(_q3)

ck("three constructions produce three priced slab rows", len(_rows3) == 3,
   f"-> {len(_rows3)}")
ck("each priced row carries its own rate", len({r for _d, _q, r in _rows3}) == 3,
   f"-> {sorted(r for _d, _q, r in _rows3)}")
ck("a thicker construction prices higher per m2",
   dict((d.split(' — ')[0], r) for d, _q, r in _rows3).get("Container Slab Construction", 0)
   > dict((d.split(' — ')[0], r) for d, _q, r in _rows3).get("HGV Slab Construction", 1e9))
ck("each row is named by its construction",
   all(any(c["name"] in d for c in _THREE) for d, _q, _r in _rows3),
   f"-> {[d for d, _q, _r in _rows3]}")
ck("each row states its own thickness",
   all(f"{int(c['depth_mm'])}mm" in " ".join(d for d, _q, _r in _rows3) for c in _THREE))
ck("quantities are the constructions' own areas",
   sorted(q for _d, q, _r in _rows3) == sorted(c["area_m2"] for c in _THREE))
ck("no construction area is lost or double counted",
   round(sum(q for _d, q, _r in _rows3), 1) == round(sum(c["area_m2"] for c in _THREE), 1))
ck("the combined value is the sum of the separately priced rows",
   round(sum(q * r for _d, q, r in _rows3), 2)
   == round(sum(_cp.price_constructions(_THREE, _SPEC)[i]["total_gbp"] for i in range(3)), 2))
ck("the quotation declares that the total is a sum of constructions",
   any("separate constructions" in d for d in _q3["declarations"]))

# The gold-path guard. Constructions that all share the zone's own thickness describe the
# SAME slab the old single-line quote described, so the document must be identical -- not
# merely close. If this fails, the change has moved a price on a job it had no business
# touching, which is exactly the failure this whole module is watching for.
_same = [{"name": "A", "depth_mm": 200, "area_m2": 13234.0, "region_ids": ["a"]},
         {"name": "B", "depth_mm": 200, "area_m2": 10000.0, "region_ids": ["b"]}]
_q_same = generate_quotation(_job(_same), project="P", client="C", ref="R")
_q_plain = generate_quotation(_job(None), project="P", client="C", ref="R")
ck("one thickness throughout prices exactly as an unfanned job",
   _q_same["subtotal_gbp"] == _q_plain["subtotal_gbp"],
   f"-> {_q_same['subtotal_gbp']} vs {_q_plain['subtotal_gbp']}")
ck("one thickness throughout keeps one priced slab row",
   _concrete_rows(_q_same) == _concrete_rows(_q_plain),
   f"-> {_concrete_rows(_q_same)}")
ck("a job with no construction breakdown is untouched",
   len(_concrete_rows(_q_plain)) == 1 and _concrete_rows(_q_plain)[0][1] == 23234.0)

# A thickness the legend did not state is never borrowed from a neighbouring construction.
_no_depth = [{"name": "HGV Slab Construction", "depth_mm": 375, "area_m2": 4000.0,
              "region_ids": ["a"]},
             {"name": "Unlabelled Apron", "depth_mm": None, "area_m2": 900.0,
              "region_ids": ["b"]}]
_q_nd = generate_quotation(_job(_no_depth, area_m2=4900.0),
                           project="P", client="C", ref="R")
_nd_priced = _cp.price_constructions(_no_depth, _SPEC)
# Until 16 Sep this fell back to the zone thickness and priced on it. That is what put
# £13,182.98 on Radlett's client document at a thickness the drawing never gave. A construction
# with no stated thickness now borrows nothing — not the zone's, and certainly not a
# neighbouring construction's.
ck("a construction with no stated thickness borrows no thickness at all",
   _nd_priced[1]["depth_mm"] is None and _nd_priced[1]["depth_assumed"],
   f"-> {_nd_priced[1]['depth_mm']}")
ck("...so it is not priced", _nd_priced[1]["rate"] is None
   and _nd_priced[1]["total_gbp"] is None)
ck("...while the construction that DOES state one is priced normally",
   isinstance(_nd_priced[0]["rate"], (int, float)) and _nd_priced[0]["depth_mm"] == 375)
ck("the thickness gap is declared on the quotation",
   any("THICKNESS NOT STATED" in d for d in _q_nd["declarations"]),
   f"-> {[d for d in _q_nd['declarations'] if 'THICKNESS' in d]}")

# The legend thickness is a drawing fact, so it must not be recorded as an assessor entry.
_specs3 = {s["area_m2"]: s for s in _q3["specifications"]}
_container_spec = _specs3.get(17689.3, {})
_depth_field = (_container_spec.get("fields") or {}).get("depth_mm") or {}
ck("a legend thickness is recorded with engineer-drawing provenance",
   _depth_field.get("source") == "engineer_drawing", f"-> {_depth_field.get('source')}")
ck("the legend thickness is not marked provisional",
   _depth_field.get("provisional") is False)
ck("the construction's detail reference is cited with it",
   "DD-C-1020" in str((_depth_field.get("evidence") or {}).get("text") or ""),
   f"-> {(_depth_field.get('evidence') or {}).get('text')}")

# The take-off sheet: each construction totals separately, then once combined.
_wb = _load_workbook(_io.BytesIO(quotation_xlsx(_q3)), data_only=False)
_ws = _wb[_wb.sheetnames[0]]
_area_totals = [(_ws.cell(r, 1).value, _ws.cell(r, 2).value)
                for r in range(1, _ws.max_row + 1)
                if isinstance(_ws.cell(r, 1).value, str)
                and _ws.cell(r, 1).value.startswith("Total")
                and "Area" in _ws.cell(r, 1).value]
ck("each construction gets its own Total Area Take Off",
   sum(1 for label, _f in _area_totals if label == "Total Area Take Off:") == 3,
   f"-> {[l for l, _f in _area_totals]}")
ck("the sheet ends the section with Aryan's combined area line",
   any(label == "Total External/Service Yard Slab Area:" for label, _f in _area_totals))
_combined = next((f for label, f in _area_totals
                  if label == "Total External/Service Yard Slab Area:"), "")
ck("the combined line sums the constructions' own totals, not a restated zone area",
   str(_combined).startswith("=B") and str(_combined).count("+") == 2,
   f"-> {_combined}")

_wb_plain = _load_workbook(_io.BytesIO(quotation_xlsx(_q_plain)), data_only=False)
_ws_plain = _wb_plain[_wb_plain.sheetnames[0]]
ck("a job with no constructions gets no combined area line",
   not any(isinstance(_ws_plain.cell(r, 1).value, str)
           and "Yard Slab Area" in _ws_plain.cell(r, 1).value
           for r in range(1, _ws_plain.max_row + 1)))

# The portal card and the client document must not disagree about the same job.
from approval_server import _roll_up_constructions as _roll
_rolled = _roll(_job(_THREE)["costing"], _job(_THREE))
ck("the portal costing totals the same money as the quotation's slab rows",
   abs(_rolled["total_gbp"] - round(sum(q * r for _d, q, r in _rows3), 2)) < 0.05,
   f"-> {_rolled['total_gbp']}")
ck("the portal costing keeps the constructions' own areas",
   _rolled["constructions_area_m2"] == round(sum(c["area_m2"] for c in _THREE), 1))
ck("the portal rate is labelled as an implied average, not a priced rate",
   "not a rate anything was priced at" in _rolled.get("constructions_note", ""))
# Within half a penny per m2 -- the rate is rounded to 2dp for display, so this is the
# tightest the identity can be stated, and it is stated against the area rather than a
# flat currency figure so the check does not loosen as the job gets bigger.
ck("the implied average multiplies back out to the stated total",
   abs(_rolled["rate"] * _rolled["area_m2"] - _rolled["total_gbp"])
   <= 0.005 * _rolled["area_m2"],
   f"-> {_rolled['rate']} x {_rolled['area_m2']} vs {_rolled['total_gbp']}")
ck("a costing with no constructions is returned untouched",
   _roll(_job(None)["costing"], _job(None)) == _job(None)["costing"])
ck("a costing with no spec is returned untouched rather than guessed",
   _roll({"area_m2": 10.0, "rate": 5.0, "total_gbp": 50.0},
         _job(_THREE)) == {"area_m2": 10.0, "rate": 5.0, "total_gbp": 50.0})

# A mixed drawing: only the Yard fans; the Dock alongside it keeps its own single line.
_mixed = _job(_THREE)
_mixed["zones"].append({"category": "dock", "area_m2": 930.0})
_mixed["area_m2"] = 24164.0
_mixed["costing"]["area_m2"] = 24164.0
_q_mixed = generate_quotation(_mixed, project="P", client="C", ref="R")
_sections = [li["section"] for li in _q_mixed["line_items"]
             if li.get("line_role") == "concrete_slab"]
ck("a Dock zone beside a fanned Yard stays one row in its own section",
   _sections.count("Dock slabs") == 1 and _sections.count("External yard slabs") == 3,
   f"-> {_sections}")

_text = quotation_text(_q3)
ck("the plain-text quotation names every construction",
   all(c["name"] in _text for c in _THREE))
ck("the plain-text quotation does not repeat the sum declaration",
   _text.count("separate constructions") == 1, f"-> {_text.count('separate constructions')}")

# The shape production actually produces, and the one every fixture above got wrong.
#
# A Surface Finishes Plan arrives UNMEASURED with NO zones -- the assessor's Yard-region
# review is what measures it -- so its breakdown sits on the RESULT, not on a zone. Portal QA
# on the real Radlett sheet found four constructions (180/200/225/375 mm) being priced as one
# 25,730 m2 slab at the 190 mm default, with the per-construction flag printed above it. Every
# check above passed while that was true, because every fixture above put the constructions on
# a zone the live sheet does not have.
def _zoneless_job(constructions, area_m2=25730.7):
    job = {"file": "radlett_sfp.pdf", "pdf_path": "drawings/radlett_sfp.pdf",
           "area_m2": area_m2, "zones": [],
           "costing": {"area_m2": area_m2, "rate": _ZONE_RATE,
                       "total_gbp": round(area_m2 * _ZONE_RATE, 2), "spec": dict(_SPEC),
                       "assumed": True, "breakdown": {}},
           "perimeter_lm": 640.0}
    if constructions is not None:
        job["constructions"] = _copy.deepcopy(constructions)
    return job


_RADLETT = [
    {"name": "Intermodal Terminal HGV Slab Construction 200mm thick", "depth_mm": 200,
     "area_m2": 2134.9, "region_ids": ["surface-finish-3"]},
    {"name": "Intermodal Terminal HGV Slab Construction 225mm thick", "depth_mm": 225,
     "area_m2": 1564.6, "region_ids": ["surface-finish-4"]},
    {"name": "Intermodal Terminal HGV Slab Construction 180mm thick", "depth_mm": 180,
     "area_m2": 2908.6, "region_ids": ["surface-finish-5"]},
    {"name": "Intermodal Terminal Container Slab Construction 375mm thick", "depth_mm": 375,
     "area_m2": 19122.6, "region_ids": ["surface-finish-6"]},
]

_qz = generate_quotation(_zoneless_job(_RADLETT), project="Radlett", client="Fortel", ref="Z")
_rowsz = _concrete_rows(_qz)
ck("a zone-less reviewed sheet still prices one row per construction", len(_rowsz) == 4,
   f"-> {len(_rowsz)}")
ck("a zone-less sheet's rows carry four different rates",
   len({r for _d, _q, r in _rowsz}) == 4, f"-> {sorted(r for _d, _q, r in _rowsz)}")
ck("no zone-less row is priced at the default thickness the sheet never states",
   not any("190mm" in d for d, _q, _r in _rowsz), f"-> {[d for d, _q, _r in _rowsz]}")
ck("the zone-less quantities are the constructions' own",
   sorted(q for _d, q, _r in _rowsz) == sorted(c["area_m2"] for c in _RADLETT))
ck("a zone-less sheet gets the combined-area declaration",
   any("separate constructions" in d for d in _qz["declarations"]))

_wbz = _load_workbook(_io.BytesIO(quotation_xlsx(_qz)), data_only=False)
_wsz = _wbz[_wbz.sheetnames[0]]
ck("a zone-less sheet gets the combined area line in the workbook",
   any(isinstance(_wsz.cell(r, 1).value, str) and "Yard Slab Area" in _wsz.cell(r, 1).value
       for r in range(1, _wsz.max_row + 1)))

# One slab has one perimeter. Fanning must not multiply the perimeter row by the number of
# constructions, nor drop it: it is the same measured outline either way.
def _perimeter_rows(q):
    # Perimeter is a MEASUREMENT row, not a priced line item -- looking for it in
    # line_items finds nothing on every job and passes vacuously.
    return [m for m in (q.get("measurements") or [])
            if m.get("description") == "Slab perimeter"]


_qz_plain = generate_quotation(_zoneless_job(None), project="Radlett", client="Fortel", ref="Z")
ck("the fixture actually has a perimeter row to protect",
   len(_perimeter_rows(_qz_plain)) == 1, f"-> {len(_perimeter_rows(_qz_plain))}")
ck("fanning does not multiply the slab perimeter row",
   len(_perimeter_rows(_qz)) == len(_perimeter_rows(_qz_plain)),
   f"-> {len(_perimeter_rows(_qz))} vs {len(_perimeter_rows(_qz_plain))}")
ck("...and the perimeter quantity is unchanged by fanning",
   [m["qty"] for m in _perimeter_rows(_qz)] == [m["qty"] for m in _perimeter_rows(_qz_plain)]
   == [640.0],
   f"-> {[m['qty'] for m in _perimeter_rows(_qz)]} vs "
   f"{[m['qty'] for m in _perimeter_rows(_qz_plain)]}")

ck("a zone-less sheet with no constructions is untouched",
   len(_concrete_rows(_qz_plain)) == 1
   and _concrete_rows(_qz_plain)[0][1] == 25730.7)

_rolled_z = _roll(_zoneless_job(_RADLETT)["costing"], _zoneless_job(_RADLETT))
ck("the portal card rolls up a zone-less breakdown too",
   len(_rolled_z.get("constructions") or []) == 4,
   f"-> {len(_rolled_z.get('constructions') or [])}")
ck("the zone-less card total matches the quotation's slab rows",
   abs(_rolled_z["total_gbp"] - round(sum(q * r for _d, q, r in _rowsz), 2)) < 0.05,
   f"-> {_rolled_z['total_gbp']}")
ck("the zone-less card is not still the single blended default",
   _rolled_z["total_gbp"] != _zoneless_job(_RADLETT)["costing"]["total_gbp"])

# 16 Sep: Aryan tested production and found two faults in one row. Both are checked here on
# the shape the real Radlett sheet produces -- a legend that states a thickness for four
# constructions and NONE for the fifth (Rail Crossing).
#
# Fault 1, the one he named: a row whose thickness WAS read off the legend, cited to the sheet
# and marked confirmed still printed "PROVISIONAL — NO DETAILS PROVIDED" beside its price.
# Fault 2, the one he did not name and which matters more: the construction with no stated
# thickness was priced at the 190 mm default -- £13,182.98 on the client document at a
# thickness the drawing never gave -- directly beneath its own specification block reading
# "nothing assumed, nothing priced".
_MIXED = [
    {"name": "Intermodal Terminal HGV Slab Construction 200mm thick", "depth_mm": 200,
     "area_m2": 2993.7, "region_ids": ["surface-finish-3"]},
    {"name": "Intermodal Terminal Rail Crossing Slab Construction", "depth_mm": None,
     "area_m2": 292.5, "region_ids": ["surface-finish-7"]},
]
_q_mixed_depth = generate_quotation(_zoneless_job(_MIXED, area_m2=3286.2),
                                    project="Radlett", client="Fortel", ref="M")
_rows_mixed = _concrete_rows(_q_mixed_depth)
_rail = next((r for r in _rows_mixed if "Rail Crossing" in r[0]), None)
_hgv = next((r for r in _rows_mixed if "HGV" in r[0]), None)

ck("a construction with no stated thickness is NOT priced", _rail is not None and _rail[2] is None,
   f"-> {_rail}")
ck("...and its description does not invent a thickness",
   _rail is not None and "mm. th" not in _rail[0], f"-> {_rail[0] if _rail else None}")
ck("...and it says why the rate is blank",
   _rail is not None and "thickness not stated" in _rail[0].lower())
ck("...and its quantity is still carried, because it was measured",
   _rail is not None and _rail[1] == 292.5)
ck("...and the row asks the assessor for a rate", any(
    li.get("assessor_rate_required") for li in _q_mixed_depth["line_items"]
    if li.get("line_role") == "concrete_slab" and "Rail Crossing" in li["description"]))
ck("a construction WITH a stated thickness still prices", _hgv is not None and _hgv[2] == 59.41,
   f"-> {_hgv}")
ck("the unpriced construction contributes no money to the quotation",
   all(li.get("value") in (None, 0) for li in _q_mixed_depth["line_items"]
       if "Rail Crossing" in li.get("description", "") and li.get("line_role") == "concrete_slab"))
ck("the sheet declares the thickness gap as measured-but-not-priced",
   any("MEASURED BUT NOT PRICED" in d for d in _q_mixed_depth["declarations"]),
   f"-> {[d for d in _q_mixed_depth['declarations'] if 'THICKNESS' in d]}")

_rolled_mixed = _roll(_zoneless_job(_MIXED, area_m2=3286.2)["costing"],
                      _zoneless_job(_MIXED, area_m2=3286.2))
ck("the portal card prices only what can be priced",
   _rolled_mixed["total_gbp"] == round(2993.7 * 59.41, 2),
   f"-> {_rolled_mixed['total_gbp']}")
ck("...and names on the card what is NOT in that total",
   "NOT IN THIS TOTAL" in _rolled_mixed["constructions_note"]
   and "Rail Crossing" in _rolled_mixed["constructions_note"])
ck("...and carries the unpriced quantity",
   _rolled_mixed.get("constructions_unpriced_m2") == 292.5)

# Fault 1: the reason must name what came off the drawing.
_reason = next((li["provisional_reason"] for li in _q_mixed_depth["line_items"]
                if li.get("line_role") == "concrete_slab" and "HGV" in li["description"]), "")
ck("a row whose thickness came from the drawing does not claim NO DETAILS PROVIDED",
   _reason != PROVISIONAL_LABEL, f"-> {_reason!r}")
ck("...it names the thickness as from the drawing", "thickness from the drawing" in _reason)
ck("...and still names mesh and mix as assumed",
   "mesh" in _reason and "mix" in _reason and "assumed" in _reason)
ck("...and the row remains provisional, because the price rests on assumed mesh and mix",
   any(li["provisional"] for li in _q_mixed_depth["line_items"]
       if li.get("line_role") == "concrete_slab" and "HGV" in li["description"]))

# The gold guard for the label: nothing confirmed anywhere -> the original string, exactly.
_q_plain_reason = next(
    (li["provisional_reason"] for li in _qz_plain["line_items"]
     if li.get("line_role") == "concrete_slab"), None)
ck("a row with no extracted detail keeps the original wording byte-for-byte",
   _q_plain_reason == PROVISIONAL_LABEL, f"-> {_q_plain_reason!r}")

# And the gold guard for money: every construction stating its own thickness is untouched by
# any of this — same rows, same rates, same subtotal as before the 16 Sep change.
_q_all_stated = generate_quotation(_zoneless_job(_RADLETT), project="R", client="F", ref="A")
ck("a sheet where every construction states a thickness still prices all of them",
   all(r[2] is not None for r in _concrete_rows(_q_all_stated)),
   f"-> {[r[2] for r in _concrete_rows(_q_all_stated)]}")
ck("...and none of its rows carries the thickness-not-stated wording",
   not any("thickness not stated" in r[0].lower() for r in _concrete_rows(_q_all_stated)))


# ── The workbook's own arithmetic ────────────────────────────────────────────────────────
# Aryan, 16 Sep 2026, after reading the exported sheet: "#DIV/0! errors in the 180mm section",
# "formulas that appear to be using the DPM gauge where the steel rate should be used", "I
# don't think the final quotation total should be treated as validated yet". He was right on
# all three. The rate build-up block was a transcription of the reference workbook's LAYOUT at
# fixed row numbers, written before one measured area could fan into several constructions, so
# every cross-reference into the take-off pointed at whatever had since moved into that row.
#
# Nothing above this line could see it: openpyxl stores a formula as text, so a test that
# reads cells passes straight through a division by an empty cell. These checks EVALUATE the
# workbook (tests/_xlsx_eval.py).
from tests._xlsx_eval import Sheet as _Sheet

_wb_mixed = _load_workbook(_io.BytesIO(quotation_xlsx(_q_mixed_depth)))
_ws_mixed = _wb_mixed[_wb_mixed.sheetnames[0]]
_sheet_mixed = _Sheet(_ws_mixed)
_errors_mixed = _sheet_mixed.error_cells()
ck("every formula in the exported workbook evaluates — no #DIV/0! anywhere",
   not _errors_mixed, f"-> {_errors_mixed}")

_wb_all = _load_workbook(_io.BytesIO(quotation_xlsx(_q_all_stated)))
_ws_all = _wb_all[_wb_all.sheetnames[0]]
_sheet_all = _Sheet(_ws_all)
_errors_all = _sheet_all.error_cells()
ck("...and on the four-construction sheet too", not _errors_all, f"-> {_errors_all}")


def _buildup_blocks(ws):
    """(header row, TOTAL RATE/M2 row) for each rate build-up written into columns G-J."""
    blocks, header = [], None
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 7).value
        if isinstance(label, str) and label.startswith("Rate build-up — "):
            header = row
        elif label == "TOTAL RATE/M2" and header:
            blocks.append((header, row))
            header = None
    return blocks


_blocks_all = _buildup_blocks(_ws_all)
ck("one rate build-up is written per priced construction, not per BOQ section",
   len(_blocks_all) == len(_concrete_rows(_q_all_stated)),
   f"-> {len(_blocks_all)} blocks for {len(_concrete_rows(_q_all_stated))} priced rows")

# THE invariant. The build-up is Fortel's working for the rate in column D; if the two
# disagree the sheet contradicts itself, which is exactly why Aryan would not sign off the
# total. Checked by evaluating the formulas, not by reading the numbers back.
_rates_all = sorted(r for _d, _q, r in _concrete_rows(_q_all_stated))
_totals_all = sorted(_sheet_all.value(f"H{total_row}") for _h, total_row in _blocks_all)
ck("each build-up's TOTAL RATE/M2 equals the rate priced beside it",
   _totals_all == _rates_all, f"-> {_totals_all} vs {_rates_all}")

# The workbook says so itself, in a cell the assessor can see.
_checks_all = [_sheet_all.value(f"H{header + 16}") for header, _t in _blocks_all]
ck("...and the workbook's own agreement cell reads OK for every block",
   set(_checks_all) == {"OK"}, f"-> {_checks_all}")

# No component may be a negative cost: Total Trimming used to subtract a cell that had become
# the Joints rate, pricing trimming at -£3.26/m2.
_components_all = [_sheet_all.value(f"H{row}")
                   for header, total in _blocks_all for row in range(header + 1, total)]
ck("no component of any build-up is negative",
   all(not isinstance(v, (int, float)) or v >= 0 for v in _components_all),
   f"-> {[v for v in _components_all if isinstance(v, (int, float)) and v < 0]}")

# Steel was zero on every sheet ever exported: the "Steel Rate/T" label was written into
# column I with column J left empty, so the mesh line multiplied by a blank cell.
_steel_totals = [_sheet_all.value(f"H{header + 7}") for header, _t in _blocks_all]
ck("the steel line is priced, not multiplied by an empty input cell",
   all(isinstance(v, (int, float)) and v > 0 for v in _steel_totals), f"-> {_steel_totals}")

# The construction that could not be priced must not acquire a build-up here by the back door.
_blocks_mixed = _buildup_blocks(_ws_mixed)
ck("a construction with no stated thickness gets no rate build-up",
   len(_blocks_mixed) == 1, f"-> {len(_blocks_mixed)} blocks")
ck("...and no build-up block names it",
   not any("Rail Crossing" in str(_ws_mixed.cell(header, 7).value)
           for header, _t in _blocks_mixed))

# A job with no constructions at all is the shape every other client sheet takes.
_wb_plain = _load_workbook(_io.BytesIO(quotation_xlsx(_qz_plain)))
_ws_plain = _wb_plain[_wb_plain.sheetnames[0]]
_sheet_plain = _Sheet(_ws_plain)
ck("a sheet with no construction breakdown still exports a workbook that evaluates",
   not _sheet_plain.error_cells(), f"-> {_sheet_plain.error_cells()}")
_blocks_plain = _buildup_blocks(_ws_plain)
ck("...with exactly one build-up, for its one priced slab", len(_blocks_plain) == 1,
   f"-> {len(_blocks_plain)}")
ck("...whose TOTAL RATE/M2 is that slab's rate",
   _sheet_plain.value(f"H{_blocks_plain[0][1]}")
   == next(r for _d, _q, r in _concrete_rows(_qz_plain)),
   f"-> {_sheet_plain.value(f'H{_blocks_plain[0][1]}')}")


# ── What the sheet SAYS about what it extracted ──────────────────────────────────────────
# Aryan, 16 Sep: "'PROVISIONAL — NO DETAILS PROVIDED' makes it sound like nothing was
# extracted" and "the notes say 'no engineer construction-detail drawing supplied' even though
# the thicknesses are explicitly being taken from the Surface Finishes drawings". Both were
# fixed for the portal and the HTML on 16 Sep and BOTH SURVIVED IN THE XLSX, which is the
# output he reads. The fix is only real if it is checked in the workbook.
def _column_f(ws):
    return [ws.cell(row, 6).value for row in range(1, ws.max_row + 1)
            if isinstance(ws.cell(row, 6).value, str) and ws.cell(row, 6).value]


_f_all = _column_f(_ws_all)
ck("the workbook's provisional column names what came off the drawing",
   any("from the drawing" in text for text in _f_all), f"-> {sorted(set(_f_all))[:2]}")
ck("...and no row on a sheet with extracted thicknesses claims NO DETAILS PROVIDED",
   not any(text == PROVISIONAL_LABEL for text in _f_all),
   f"-> {[t for t in _f_all if t == PROVISIONAL_LABEL]}")
ck("a sheet with nothing extracted still prints the original wording in the workbook",
   PROVISIONAL_LABEL in _column_f(_ws_plain))

_titles_all = [_ws_all.cell(row, 1).value for row in range(1, _ws_all.max_row + 1)
               if isinstance(_ws_all.cell(row, 1).value, str)
               and "Provisional Cost" in _ws_all.cell(row, 1).value]
ck("the section heading does not say (No Details) over rows that have them",
   _titles_all and all("No Details" not in title for title in _titles_all),
   f"-> {_titles_all}")

_notes_all = " ".join(_q_all_stated["declarations"])
ck("the notes do not claim no construction detail was supplied for a stated thickness",
   "no engineer construction-detail drawing supplied)" not in _notes_all)
ck("...they say which part of the build-up is assumed instead",
   "PART-ASSUMED" in _notes_all and "ASSUMED" in _notes_all)
ck("a job with nothing extracted keeps Fortel's original assumption wording",
   "no engineer construction-detail drawing supplied" in
   " ".join(_qz_plain["declarations"]))

# One count printed two ways is what produced "4" in one place and "5" in another.
_counts = [d for d in _q_mixed_depth["declarations"] if "separate constructions" in d]
ck("the constructions note states measured AND priced, never one of them alone",
   _counts and "2 separate constructions measured, of which 1 is priced" in _counts[0],
   f"-> {_counts}")
ck("...and a sheet where all of them price says so plainly",
   any("4 separate constructions," in d for d in _q_all_stated["declarations"]),
   f"-> {[d for d in _q_all_stated['declarations'] if 'separate constructions' in d]}")

# The same label reaches the client through four renderers. On 16 Sep it was fixed in the JSON
# only, and the XLSX, the HTML and the plain text each went on printing the old constant.
from quotation import quotation_html as _quotation_html

_html_all = _quotation_html(_q_all_stated)
ck("the HTML quotation names what came off the drawing, like the workbook",
   "thickness from the drawing" in _html_all)
ck("...and no HTML row claims NO DETAILS PROVIDED on a sheet that has them",
   PROVISIONAL_LABEL not in _html_all)
_text_all = quotation_text(_q_all_stated)
ck("the plain-text quotation says the same thing",
   "thickness from the drawing" in _text_all and PROVISIONAL_LABEL not in _text_all)
ck("a sheet with nothing extracted still reads NO DETAILS PROVIDED in HTML",
   PROVISIONAL_LABEL in _quotation_html(_qz_plain))


# ── One named area, drawn in two places ──────────────────────────────────────────────────
# Inderjit, 16 Sep handover call: he measures one footpath with + Area, then wants the SECOND
# footpath inside that same footpath element rather than in the yard. + Region always went to
# the main slab, so every footpath became its own element and got its own quotation line --
# "all the footpaths needs to be together". Aryan approved a destination selector beside
# + Region; it works by giving the second piece the SAME element id, and the grouping below is
# what turns that into one priced line. If this stops summing, the selector is a lie.
def _element_job(elements, area_m2=5000.0):
    return {
        "file": "warwick_sfp.pdf", "pdf_path": "drawings/warwick_sfp.pdf", "area_m2": area_m2,
        "zones": [{"category": "external_yard", "area_m2": area_m2}],
        "area_elements": elements,
        "costing": {"area_m2": area_m2, "rate": _ZONE_RATE,
                    "total_gbp": round(area_m2 * _ZONE_RATE, 2), "spec": dict(_SPEC),
                    "assumed": True, "breakdown": {}},
    }


_two_pieces = [
    {"element_id": "area-footpath", "name": "Footpath", "category": "external_yard",
     "area_m2": 204.39, "boq_scope": "main"},
    {"element_id": "area-footpath", "name": "Footpath", "category": "external_yard",
     "area_m2": 198.0, "boq_scope": "main"},
]
_q_two = generate_quotation(_element_job(_two_pieces), project="W", client="F", ref="W1")
_fp_rows = [li for li in _q_two["line_items"]
            if li.get("line_role") == "concrete_slab" and "Footpath" in li["description"]]
ck("two pieces of one named area make ONE priced line, not two",
   len(_fp_rows) == 1, f"-> {len(_fp_rows)} rows")
ck("...and that line carries the sum of both pieces",
   bool(_fp_rows) and abs(float(_fp_rows[0]["qty"]) - 402.39) < 0.01,
   f"-> {_fp_rows[0]['qty'] if _fp_rows else None}")

# And two genuinely different areas must still stay apart — the selector must not merge things
# the assessor kept separate.
_two_areas = [
    {"element_id": "area-fp-100", "name": "Footpath Unit 100", "category": "external_yard",
     "area_m2": 204.39, "boq_scope": "main"},
    {"element_id": "area-fp-200", "name": "Footpath Unit 200", "category": "external_yard",
     "area_m2": 198.0, "boq_scope": "main"},
]
_q_sep = generate_quotation(_element_job(_two_areas), project="W", client="F", ref="W2")
_sep_rows = [li for li in _q_sep["line_items"]
             if li.get("line_role") == "concrete_slab" and "Footpath Unit" in li["description"]]
ck("two DIFFERENT named areas still price as two separate lines",
   len(_sep_rows) == 2, f"-> {[r['description'][:26] for r in _sep_rows]}")

# The control itself.
_portal_src = open("assessor_portal.html", encoding="utf-8").read()
ck("the portal offers a destination for the next + Region outline",
   'id="regionTarget"' in _portal_src)
ck("...which defaults to the main slab",
   '<option value="">＋ Region → Main slab</option>' in _portal_src)
ck("...and a region sent to a named area reuses that area's element id",
   "elementId: target.elementId" in _portal_src)


# ── A sheet that states several thicknesses prices none of them ──────────────────────────
# PLP Warwick Site A, on the 16 Sep handover call: the quotation priced 2,222.5 m² at 190 mm
# for £100,188.08 while the job's own flag read "SPEC CONFLICT — slab thickness: the drawing
# states 180 mm… and 200 mm… and 150 mm… nothing is assumed for pricing". The extractor was
# telling the truth and the price was contradicting it: 190 came from DEFAULT_SPEC. The sheet
# does contain 190, under "Ground Floor Slab Construction (By Others)" — not Fortel's scope —
# so it was doubly not ours to charge for. Same rule as the Rail Crossing: measured, not priced.
from slab_spec import build_brief_spec as _build_brief_spec
from quotation import depth_unresolved_on_drawing as _depth_unresolved

_CONFLICT_NOTES = {"depth_mm": {
    "source": "drawing_states_more_than_one",
    "note": "STATED ON THE DRAWING as 150, 180, 200 mm for different surfaces"}}


def _sheet_job(field_notes, area_m2=2222.5):
    brief = _build_brief_spec("external_yard", field_notes=field_notes or {})
    return {
        "file": "plp_site_a.pdf", "pdf_path": "drawings/plp_site_a.pdf", "area_m2": area_m2,
        "zones": [{"category": "external_yard", "area_m2": area_m2}],
        "brief_spec": brief, "brief_specs": {"external_yard": brief},
        "spec_field_notes": field_notes or {},
        "costing": {"area_m2": area_m2, "rate": _ZONE_RATE,
                    "total_gbp": round(area_m2 * _ZONE_RATE, 2), "spec": dict(_SPEC),
                    "assumed": True, "breakdown": {}},
    }


ck("a thickness the drawing gives several answers for counts as unresolved",
   _depth_unresolved(_build_brief_spec("external_yard", field_notes=_CONFLICT_NOTES)))
ck("...while a sheet that simply never mentioned it does NOT change behaviour",
   not _depth_unresolved(_build_brief_spec("external_yard")))

_q_conflict = generate_quotation(_sheet_job(_CONFLICT_NOTES), project="PLP", client="F", ref="P1")
_conflict_rows = _concrete_rows(_q_conflict)
ck("a sheet stating several thicknesses is measured but NOT priced",
   len(_conflict_rows) == 1 and _conflict_rows[0][2] is None,
   f"-> {[r[2] for r in _conflict_rows]}")
ck("...its quantity is still carried, because it WAS measured",
   bool(_conflict_rows) and abs(float(_conflict_rows[0][1]) - 2222.5) < 0.01,
   f"-> {_conflict_rows[0][1] if _conflict_rows else None}")
ck("...its description does not quote the 190mm default beside the blank rate",
   bool(_conflict_rows) and "190" not in _conflict_rows[0][0],
   f"-> {_conflict_rows[0][0][:70] if _conflict_rows else None}")
ck("...and it says why the rate is blank",
   bool(_conflict_rows) and "more than one thickness" in _conflict_rows[0][0])
ck("...so no money rests on a thickness nobody chose",
   not any(isinstance(li.get("value"), (int, float)) and li["value"] > 0
           for li in _q_conflict["line_items"] if li.get("line_role") == "concrete_slab"))

# THE guard against over-reach: an ordinary sheet with no conflict must price exactly as it
# did before. This rule may only ever remove a number the drawing never justified.
_q_ordinary = generate_quotation(_sheet_job(None), project="PLP", client="F", ref="P2")
_ordinary_rows = _concrete_rows(_q_ordinary)
ck("a sheet with no thickness conflict still prices exactly as before",
   len(_ordinary_rows) == 1 and _ordinary_rows[0][2] == _ZONE_RATE,
   f"-> {[r[2] for r in _ordinary_rows]} vs {_ZONE_RATE}")
ck("...and still states its thickness in the description",
   bool(_ordinary_rows) and "200mm" in _ordinary_rows[0][0],
   f"-> {_ordinary_rows[0][0][:60] if _ordinary_rows else None}")
