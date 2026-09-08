#!/usr/bin/env python3
"""approval_server upload/approve gate, queue, /snapshot, watchdog, jobs file.

Sections in this module (printed in this order):
  - approval_server: upload format handling + approve hard-block
  - approval_server: bounded 26-PDF queue excludes wait time from watchdog budget
  - approval_server: /snapshot status codes for all four measurement states (Aryan field report — 'session which renders screenshots is not working properly')
  - approval_server: watchdog-vs-completion race (Aryan field report — 'server is unstable')
  - approval_server: approval_jobs.json concurrent read/write does not raise (Aryan field report — 'the server is unstable')

Imported by ci_tests.py, which owns the running total. Do not reorder.
"""
import os, shutil
from pathlib import Path
from tests import ck
import copy as _copy, fitz as _fitz_fast_refusal
from io import BytesIO as _BytesIO
from openpyxl import load_workbook as _load_workbook
from quotation import FORTEL_CHANNEL_ROW, FORTEL_TRANSITION_ROW, quotation_html, quotation_json, quotation_text, quotation_xlsx
from slab_spec import empty_brief_spec as _empty_brief_spec
from tests.test_marked_export import _MARKED_MANIFEST_NAME, _fitz_marked, _json_marked, _marked_job
from tests.test_quotation import _demo_result

print("approval_server: upload format handling + approve hard-block")
try:
    import approval_server as _AS
    import fitz as _fitz3, zipfile as _zipfile, tempfile as _tempfile, io as _io3

    _tmpdir = Path(_tempfile.mkdtemp(prefix="ci_upload_"))

    # .zip with two PDFs -> both extracted, ranked by drawing_priority, no zip-slip
    _pdf_a = _tmpdir / "Proposed Site Plan.pdf"
    _pdf_b = _tmpdir / "External Construction Thickness Layout.pdf"
    for _pp in (_pdf_a, _pdf_b):
        _dd = _fitz3.open(); _dd.new_page(); _dd.save(str(_pp))
    _zip_path = _tmpdir / "pack.zip"
    with _zipfile.ZipFile(_zip_path, "w") as _zf:
        _zf.write(_pdf_a, _pdf_a.name)
        _zf.write(_pdf_b, _pdf_b.name)
    _extracted, _zflags = _AS._safe_extract_zip(_zip_path, _tmpdir)
    ck("zip extraction pulls both PDFs", len(_extracted) == 2)
    _ranked_zip = _AS._rank_pdfs_by_priority(_extracted)
    ck("zip PDFs ranked — construction-thickness beats site plan",
       "Construction_Thickness" in _ranked_zip[0].name or "Thickness" in _ranked_zip[0].name)

    # zip-slip guard: a malicious entry name must never escape dest_dir
    _evil_zip = _tmpdir / "evil.zip"
    with _zipfile.ZipFile(_evil_zip, "w") as _zf:
        _zf.writestr("../../etc/evil.pdf", b"%PDF-1.4 fake")
    _esc_before = set(_tmpdir.parent.glob("evil.pdf"))
    _extracted_evil, _eflags = _AS._safe_extract_zip(_evil_zip, _tmpdir)
    ck("zip-slip entry sanitised to a safe basename (stays inside dest_dir)",
       all(str(p).startswith(str(_tmpdir.resolve())) for p in _extracted_evil))

    # encrypted / zero-byte PDF -> rejected reason, not a crash
    _zero = _tmpdir / "zero.pdf"; _zero.write_bytes(b"")
    _doc, _reason = _AS._open_pdf_safely(_zero)
    ck("zero-byte PDF -> rejected with reason (not a crash)", _doc is None and "zero-byte" in _reason)

    _enc = _tmpdir / "enc.pdf"
    _ed = _fitz3.open(); _ed.new_page()
    _ed.save(str(_enc), encryption=_fitz3.PDF_ENCRYPT_AES_256, owner_pw="x", user_pw="y")
    _doc2, _reason2 = _AS._open_pdf_safely(_enc)
    ck("encrypted PDF -> rejected with reason (not a crash)",
       _doc2 is None and "encrypted" in _reason2.lower())

    _orig_jobs_file_up = _AS.JOBS_FILE
    _orig_jobs_archive_file_up = _AS.JOBS_ARCHIVE_FILE
    _orig_backup_dir_up = _AS.BACKUP_DIR
    _orig_drawings_dir_up = _AS.DRAWINGS_DIR
    _orig_quotations_dir_up = _AS.QUOTATIONS_DIR
    _orig_server_file_up = _AS.__file__
    _orig_dispatcher_up = _AS._TAKEOFF_DISPATCHER
    _started_up = []

    class _RecordingDispatcher:
        def submit(self, *args):
            _started_up.append(args)

    try:
        _AS.JOBS_FILE = _tmpdir / "multi_upload_jobs.json"
        _AS.JOBS_ARCHIVE_FILE = _tmpdir / "multi_upload_jobs_archive.json"
        _AS.BACKUP_DIR = _tmpdir / "multi_upload_backups"
        _AS.DRAWINGS_DIR = _tmpdir / "drawings"
        _AS.QUOTATIONS_DIR = _tmpdir / "quotations"
        _AS.__file__ = str(_tmpdir / "approval_server.py")
        _AS._TAKEOFF_DISPATCHER = _RecordingDispatcher()
        _client_up = _AS.app.test_client()
        _pdf_a_bytes = _pdf_a.read_bytes()
        _pdf_b_bytes = _pdf_b.read_bytes()

        # Full HTTP lifecycle: a pending AI result cannot be issued as assessor markup;
        # /adjust persists exact snapshot coordinates, then the real download route returns
        # a permanent PDF with the same recoverable geometry and a client-safe filename.
        _marked_route_job = _copy.deepcopy(_marked_job)
        _marked_route_job.update({
            "status":"pending", "decision":None,
            "measurement_state":"MEASURED_UNVERIFIED", "adjusted":None,
            "flags":["assessor: confirm extent + scale"],
        })
        _marked_route_job["result"].update({
            "measurement_state":"MEASURED_UNVERIFIED",
            "flags":["assessor: confirm extent + scale"],
        })
        _AS.save_jobs({_marked_route_job["id"]:_marked_route_job})
        _premature_marked = _client_up.get(
            f"/marked-pdf/{_marked_route_job['id']}.pdf")
        ck("marked-PDF route refuses unreviewed AI geometry",
           _premature_marked.status_code == 409 and
           "after assessor approval or adjustment" in
               (_premature_marked.get_json().get("error") or ""),
           _premature_marked.get_json())
        import gzip as _gzip_marked
        _vector_view = _client_up.get(
            f"/snapshot-vector/{_marked_route_job['id']}.svg",
            headers={"Accept-Encoding":"gzip"})
        _vector_svg = _gzip_marked.decompress(_vector_view.data).decode("utf-8")
        _base_snapshot = _client_up.get(f"/snapshot/{_marked_route_job['id']}")
        ck("true-resolution viewport serves the rotated measured page as scalable SVG",
           _vector_view.status_code == 200 and
           _vector_view.mimetype == "image/svg+xml" and
           _vector_view.headers.get("Content-Encoding") == "gzip" and
           _vector_view.headers.get("X-Vector-Coordinate-Space") ==
               "rotated_pdf_points" and
           _vector_view.headers.get("X-Vector-Page-Width") == "200" and
           _vector_view.headers.get("X-Vector-Page-Height") == "300" and
           '<svg ' in _vector_svg and 'viewBox="0 0 200 300"' in _vector_svg,
           dict(_vector_view.headers))
        ck("vector viewport leaves the exact PNG snapshot coordinate mapping unchanged",
           _base_snapshot.status_code == 200 and
           float(_base_snapshot.headers["X-Snapshot-Scale"]) == 4.0,
           _base_snapshot.headers.get("X-Snapshot-Scale"))
        _marked_adjust = _client_up.post(f"/adjust/{_marked_route_job['id']}", json={
            "regions":[[[40,40],[240,40],[240,240],[40,240]]],
            "region_categories":["external_yard"], "region_scopes":["main"],
            "scale_k":0.1, "snapshot_scale":2.0,
            "cutout_regions":[[[80,80],[120,80],[120,120],[80,120]]],
            "user_channels":[[[40,300],[140,300],[140,380]]],
        })
        _marked_saved = _AS.load_jobs()[_marked_route_job["id"]]
        ck("adjust route persists the exact snapshot transform beside assessor geometry",
           _marked_adjust.status_code == 200 and
           _marked_saved["adjusted"]["snapshot_scale"] == 2.0 and
           _marked_saved["adjusted"]["area_m2"] == 384.0,
           _marked_adjust.get_json())
        # Shape errors used to come back as one message about fifty polygons whatever the
        # actual problem was. Inderjit hit it at Submit Adjustment on 27 Aug — "it was telling
        # me it needs to be fifty polygon, something like that" — with a shape problem, not
        # fifty of anything. Same 400s, distinct strings.
        _shape_cases = [
            ("a 501-vertex outline names the vertex limit",
             {"regions": [[[i, i] for i in range(501)]]}, ["501", "limit is 500"]),
            ("51 outlines name the polygon limit",
             {"regions": [[[0, 0], [10, 0], [10, 10]] for _ in range(51)]}, ["51", "limit is 50"]),
            ("a 2-point outline says an outline needs three points",
             {"regions": [[[0, 0], [10, 0]]]}, ["2 point", "at least 3"]),
            ("a non-finite coordinate says the point is not a finite pair",
             {"regions": [[[0, 0], [10, 0], [1e99, 1]]]}, ["finite"]),
        ]
        for _label, _payload, _expect in _shape_cases:
            _resp = _client_up.post(f"/adjust/{_marked_route_job['id']}",
                                    json=dict(_payload, scale_k=0.1, snapshot_scale=2.0))
            _msg = (_resp.get_json() or {}).get("error", "")
            ck(f"adjust rejection: {_label}",
               _resp.status_code == 400 and all(token in _msg for token in _expect), f"{_resp.status_code} {_msg}")
        _cutout_shape = _client_up.post(f"/adjust/{_marked_route_job['id']}", json={
            "regions": [[[40, 40], [240, 40], [240, 240], [40, 240]]],
            "cutout_regions": [[[0, 0], [10, 0]]], "scale_k": 0.1, "snapshot_scale": 2.0})
        ck("adjust rejection: a broken cut-out is named as a cut-out, not as a region",
           _cutout_shape.status_code == 400
           and "cut-out" in (_cutout_shape.get_json() or {}).get("error", ""),
           (_cutout_shape.get_json() or {}).get("error"))
        _area_shape = _client_up.post(f"/adjust/{_marked_route_job['id']}", json={
            "regions": [[[40, 40], [240, 40], [240, 240], [40, 240]]],
            "area_elements": [{"element_id": "a1", "name": "Footpath", "category": "external_yard",
                               "polygon_pts": [[i, i] for i in range(501)]}],
            "scale_k": 0.1, "snapshot_scale": 2.0})
        ck("adjust rejection: a broken separate area names the element, not a polygon index",
           _area_shape.status_code == 400
           and "separate area 1 (Footpath)" in (_area_shape.get_json() or {}).get("error", ""),
           (_area_shape.get_json() or {}).get("error"))
        # Put the good geometry back so the marked-PDF assertions below still see it.
        _client_up.post(f"/adjust/{_marked_route_job['id']}", json={
            "regions":[[[40,40],[240,40],[240,240],[40,240]]],
            "region_categories":["external_yard"], "region_scopes":["main"],
            "scale_k":0.1, "snapshot_scale":2.0,
            "cutout_regions":[[[80,80],[120,80],[120,120],[80,120]]],
            "user_channels":[[[40,300],[140,300],[140,380]]],
        })
        _marked_route_response = _client_up.get(
            f"/marked-pdf/{_marked_route_job['id']}.pdf")
        with _fitz_marked.open(
                stream=_marked_route_response.data, filetype="pdf") as _route_marked_doc:
            _route_manifest = _json_marked.loads(
                _route_marked_doc.embfile_get(_MARKED_MANIFEST_NAME))
            _route_annots = sum(1 for page in _route_marked_doc
                                for _annot in (page.annots() or []))
        ck("real marked-PDF endpoint returns Bluebeam-openable permanent PDF + manifest",
           _marked_route_response.status_code == 200 and
           _marked_route_response.mimetype == "application/pdf" and
           _marked_route_response.headers.get("X-Fortel-Markup-Schema") ==
               "fortel.markup.v1" and
           "MARKED_QA_01_Rotated_Drawing_REV_03_MARKED.pdf" in
               _marked_route_response.headers.get("Content-Disposition", "") and
           _route_manifest["geometry"]["regions"][0]["area_m2"] == 384.0 and
           _route_annots == 0,
           {"status":_marked_route_response.status_code,
            "content_disposition":_marked_route_response.headers.get("Content-Disposition"),
            "manifest":_route_manifest})

        # P2 lifecycle: adding two independently named areas to an already-adjusted main
        # slab must preserve 384m2 (not turn it into 454m2), and each name must become its
        # own line in the same case workbook. This uses the real HTTP routes.
        _named_area_adjust = _client_up.post(f"/adjust/{_marked_route_job['id']}", json={
            "scale_k":0.1, "snapshot_scale":2.0,
            "area_elements":[
                {"element_id":"footpath-1", "name":"Footpath",
                 "category":"external_yard", "boq_scope":"main",
                 "polygon_pts":[[300,40],[400,40],[400,90],[300,90]]},
                {"element_id":"duct-slab-1", "name":"Duct slab",
                 "category":"dock", "boq_scope":"main",
                 "polygon_pts":[[300,120],[340,120],[340,170],[300,170]]},
            ],
        })
        _named_area_saved = _AS.load_jobs()[_marked_route_job["id"]]
        ck("+Area preserves the main measured total and measures each polygon independently",
           _named_area_adjust.status_code == 200 and
           _named_area_adjust.get_json()["main_area_m2"] == 384.0 and
           _named_area_adjust.get_json()["area_elements_total_m2"] == 70.0 and
           _named_area_saved["area_m2"] == 384.0 and
           _named_area_saved["cutout_regions"] ==
               [[[80,80],[120,80],[120,120],[80,120]]] and
           _named_area_saved["user_channels"] ==
               [[[40,300],[140,300],[140,380]]] and
           [(element["name"], element["area_m2"])
            for element in _named_area_saved["area_elements"]] ==
               [("Footpath",50.0),("Duct slab",20.0)],
           _named_area_adjust.get_json())
        _unclassified_block = _AS._area_element_block_reason({
            "area_elements":[{"name":"Unknown separate slab", "category":"unclassified"}]}) or ""
        ck("unclassified +Area is approval-blocked instead of silently filename-bucketed",
           "no BOQ section" in _unclassified_block, _unclassified_block)
        ck("...and the block names the control the assessor has to use, not the concept",
           "dropdown" in _unclassified_block and "Unknown separate slab" in _unclassified_block
           and "Other / out of scope" in _unclassified_block, _unclassified_block)
        ck("'Other / out of scope' is a real answer, so a non-slab area can clear the gate",
           _AS._area_element_block_reason({
               "area_elements":[{"name":"Landscaping strip", "category":"other"}]}) is None)
        _named_area_approve = _client_up.post(
            f"/approve/{_marked_route_job['id']}", json={"note":"named areas confirmed"})
        ck("classified +Area elements complete the real approval flow",
           _named_area_approve.status_code == 200,
           _named_area_approve.get_json())

        _named_q_response = _client_up.get(
            f"/quotation/{_marked_route_job['id']}.json")
        _named_q = _json_marked.loads(_named_q_response.data)
        _named_rows = [item for item in _named_q["line_items"]
                       if item.get("assessor_named_area")]
        _main_slab_rows = [item for item in _named_q["line_items"]
                           if "slab" in item.get("description", "").lower()
                           and not item.get("assessor_named_area")]
        # Inderjit, 27 Aug: a classified +Area element came back as "just as a one row" with no
        # thickness/mesh/finish. It now gets its own specification block and build-up rows —
        # keyed by element id so it cannot merge into the main slab — while its concrete row
        # keeps EXACTLY the rate the single row carried before.
        _named_specs = [spec for spec in _named_q["specifications"]
                        if any(row.get("description") in {"Footpath", "Duct slab"}
                               for row in spec.get("area_rows") or [])]
        ck("a classified +Area element gets its own specification block, not one bare row",
           sorted((spec["section"], spec["area_m2"]) for spec in _named_specs)
           == [("Dock slabs", 20.0), ("External yard slabs", 50.0)],
           [(spec["section"], spec["area_m2"], len(spec.get("display_lines") or []))
            for spec in _named_specs])
        _named_spec_ids = {spec["id"] for spec in _named_specs}
        _named_element_rows = [item for item in _named_q["line_items"]
                               if item.get("specification_id") in _named_spec_ids]
        ck("...with the assessor's own name on its concrete row and the SAME rate as before",
           [(item["section"], item["qty"], item["rate"]) for item in _named_element_rows
            if item.get("line_role") == "concrete_slab"] == [
               ("External yard slabs", 50.0, 45.07), ("Dock slabs", 20.0, None)]
           and all(name in " | ".join(item["description"] for item in _named_element_rows)
                   for name in ("Footpath", "Duct slab")),
           [(item["description"], item["qty"], item["rate"]) for item in _named_element_rows])
        ck("...and the build-up rows carry NO invented rate: trim and joints stay for the assessor",
           all(item["rate"] is None for item in _named_element_rows
               if item.get("line_role") != "concrete_slab"),
           [(item["description"], item["rate"]) for item in _named_element_rows
            if item.get("line_role") != "concrete_slab"])
        ck("...so the quotation total is exactly what the single row produced before",
           round(sum(item.get("value") or 0 for item in _named_element_rows), 2)
           == round(50.0 * 45.07, 2),
           sum(item.get("value") or 0 for item in _named_element_rows))
        ck("+Area rows do not merge into or inflate the 384m2 main slab row",
           any(item.get("qty") == 384.0 for item in _main_slab_rows) and
           not any(item.get("qty") == 454.0 for item in _named_q["line_items"]),
           _main_slab_rows)

        from openpyxl import load_workbook as _load_named_workbook
        _named_xlsx_response = _client_up.get(
            f"/quotation/{_marked_route_job['id']}.xlsx")
        _named_wb = _load_named_workbook(
            _io3.BytesIO(_named_xlsx_response.data), data_only=False)
        _named_ws = _named_wb.active
        _named_xlsx_rows = [
            (cell.value.split("\n",1)[0], _named_ws.cell(cell.row,2).value,
             _named_ws.cell(cell.row,3).value, _named_ws.cell(cell.row,4).value,
             _named_ws.cell(cell.row,5).value)
            for row in _named_ws.iter_rows() for cell in row if cell.column == 1
            and isinstance(cell.value, str)
            and cell.value.split("\n",1)[0].split(" \u2014 ")[0] in {"Footpath","Duct slab"}
        ]
        # The workbook now carries each named element twice, the way a slab group is carried:
        # the assessor's own quantity row, then its priced concrete row with a live formula.
        _named_qty_rows = [row for row in _named_xlsx_rows if row[0] in {"Footpath", "Duct slab"}]
        _named_priced_rows = [row for row in _named_xlsx_rows if row[0] not in {"Footpath", "Duct slab"}]
        ck("same-workbook xlsx keeps named quantities numeric",
           len(_named_wb.sheetnames) == 1
           and [(row[0], row[1], row[2]) for row in _named_qty_rows] == [
               ("Footpath", 50, "m2"), ("Duct slab", 20, "m2")],
           _named_qty_rows)
        ck("...and each named element's own priced row keeps its rate and a live formula",
           [row[3] for row in _named_priced_rows] == [45.07, None]
           and isinstance(_named_priced_rows[0][4], str)
           and _named_priced_rows[0][4].startswith("=ROUND(B")
           and "*D" in _named_priced_rows[0][4]
           and isinstance(_named_priced_rows[1][4], str)
           and _named_priced_rows[1][4].startswith('=IF(D')
           and "ROUND(B" in _named_priced_rows[1][4],
           _named_priced_rows)

        _named_marked_response = _client_up.get(
            f"/marked-pdf/{_marked_route_job['id']}.pdf")
        with _fitz_marked.open(
                stream=_named_marked_response.data, filetype="pdf") as _named_marked_doc:
            _named_marked_text = "\n".join(
                page.get_text() for page in _named_marked_doc)
            _named_manifest = _json_marked.loads(
                _named_marked_doc.embfile_get(_MARKED_MANIFEST_NAME))
        ck("marked-PDF burns separately named areas and preserves their independent identity",
           "Footpath - 50.00 m2" in _named_marked_text and
           "Duct slab - 20.00 m2" in _named_marked_text and
           sum(bool(region.get("independent_area_element"))
               for region in _named_manifest["geometry"]["regions"]) == 2,
           _named_manifest["geometry"]["regions"])

        _AS.save_jobs({})
        _started_up.clear()
        _multi_resp = _client_up.post("/upload", data={
            "project_ref": "MULTI-001",
            "project_name": "Four Slab Project",
            "client_name": "Fortel QA",
            "pdf": [(_io3.BytesIO(_pdf_a_bytes), "Yard.pdf"),
                    (_io3.BytesIO(_pdf_b_bytes), "Dock.pdf")],
        }, content_type="multipart/form-data")
        _multi_json = _multi_resp.get_json()
        _multi_jobs = _AS.load_jobs()
        ck("multi-file upload returns two job_ids", _multi_resp.status_code == 202 and
           len(_multi_json.get("job_ids", [])) == 2, _multi_json)
        ck("multi-file upload creates one job per drawing under one project",
           len(_multi_jobs) == 2 and
           {j.get("project_ref") for j in _multi_jobs.values()} == {"MULTI-001"} and
           {j.get("project_name") for j in _multi_jobs.values()} == {"Four Slab Project"})
        ck("multi-file upload preserves prefixed, non-overwriting source paths",
           len({j.get("pdf_path") for j in _multi_jobs.values()}) == 2 and
           all(Path(j["pdf_path"]).name.startswith("MULTI-001_") for j in _multi_jobs.values()))
        ck("multi-file upload queues every drawing for bounded takeoff",
           len(_started_up) == 2 and all(len(args) == 4 for args in _started_up))

        # Adding a later drawing through an existing job anchor must use the persisted project
        # identity, even if a stale/tampered browser submits different visible form text.
        _existing_anchor_id = _multi_json["job_ids"][0]
        _existing_add_resp = _client_up.post("/upload", data={
            "existing_project_job_id": _existing_anchor_id,
            "project_ref": "WRONG-REF", "project_name": "Wrong look-alike project",
            "client_name": "Wrong client",
            "pdf": (_io3.BytesIO(_pdf_a_bytes), "Later folder drawing.pdf"),
        }, content_type="multipart/form-data")
        _existing_add_json = _existing_add_resp.get_json()
        _existing_add_jobs = _AS.load_jobs()
        _existing_new_job = _existing_add_jobs.get(_existing_add_json.get("job_id"), {})
        ck("existing-project upload adds one job without replacing prior project drawings",
           _existing_add_resp.status_code == 202 and
           _existing_add_json.get("added_to_project") is True and
           len(_existing_add_jobs) == 3 and
           all(job_id in _existing_add_jobs for job_id in _multi_json["job_ids"]),
           _existing_add_json)
        ck("existing-project anchor is authoritative for ref/name/client grouping",
           _existing_new_job.get("project_ref") == "MULTI-001" and
           _existing_new_job.get("project_name") == "Four Slab Project" and
           _existing_new_job.get("client_name") == "Fortel QA",
           _existing_new_job)
        _missing_existing_add = _client_up.post("/upload", data={
            "existing_project_job_id":"missing-anchor",
            "pdf":(_io3.BytesIO(_pdf_a_bytes),"orphan.pdf"),
        }, content_type="multipart/form-data")
        ck("existing-project upload refuses an unknown anchor without creating an orphan",
           _missing_existing_add.status_code == 404 and
           len(_AS.load_jobs()) == 3,
           _missing_existing_add.get_json())

        _registered_project_paths = _AS._project_pdf_paths(
            "MULTI-001", _existing_new_job["pdf_path"])
        ck("server resolves every readable same-project PDF from persisted job membership",
           len(_registered_project_paths) == 3 and
           set(_registered_project_paths) ==
               {job["pdf_path"] for job in _existing_add_jobs.values()},
           _registered_project_paths)

        # The background worker must hand that registry to the real pipeline. A narrow legacy
        # callable remains compatible because _run_takeoff introspects the optional parameter.
        import sys as _sys_project_files
        _real_project_pipeline = _sys_project_files.modules.get("takeoff_pipeline")
        _project_file_capture = {}
        class _ProjectFilePipeline:
            @staticmethod
            def takeoff(pdf_path, project_name=None, project_ref=None,
                        client_rates_path=None, approval_job_id=None, project_files=None):
                _project_file_capture["paths"] = list(project_files or [])
                return {
                    "file":Path(pdf_path).name, "pdf_path":pdf_path,
                    "project_name":project_name, "project_ref":project_ref,
                    "area_m2":None, "measurement_state":"UNMEASURED",
                    "needs_assessor":True, "flags":["project registry handoff test"],
                }
        try:
            _sys_project_files.modules["takeoff_pipeline"] = _ProjectFilePipeline
            _AS._run_takeoff(
                _existing_anchor_id,
                _existing_add_jobs[_existing_anchor_id]["pdf_path"],
                "Four Slab Project", "MULTI-001")
        finally:
            if _real_project_pipeline is None:
                _sys_project_files.modules.pop("takeoff_pipeline", None)
            else:
                _sys_project_files.modules["takeoff_pipeline"] = _real_project_pipeline
        ck("takeoff worker passes the complete project registry to spec extraction",
           set(_project_file_capture.get("paths") or []) ==
               set(_registered_project_paths), _project_file_capture)

        _AS.save_jobs({})
        _started_up.clear()
        _single_resp = _client_up.post("/upload", data={
            "project_ref": "SINGLE-001", "project_name": "Single Drawing Project",
            "pdf": (_io3.BytesIO(_pdf_a_bytes), "Yard.pdf"),
        }, content_type="multipart/form-data")
        _single_json = _single_resp.get_json()
        ck("single-file upload keeps legacy one-job response shape",
           _single_resp.status_code == 202 and "job_id" in _single_json and
           "job_ids" not in _single_json and len(_AS.load_jobs()) == 1, _single_json)

        _AS.save_jobs({})
        _started_up.clear()
        _zip_resp = _client_up.post("/upload", data={
            "project_ref": "ZIP-001", "project_name": "ZIP Slab Project",
            "pdf": (_io3.BytesIO(_zip_path.read_bytes()), "slabs.zip"),
        }, content_type="multipart/form-data")
        _zip_json = _zip_resp.get_json()
        _zip_jobs = _AS.load_jobs()
        ck("ZIP upload creates a job for every contained PDF",
           _zip_resp.status_code == 202 and len(_zip_json.get("job_ids", [])) == 2 and
           len(_zip_jobs) == 2 and len(_started_up) == 2, _zip_json)
        ck("ZIP jobs share the project ref and record all-drawings provenance",
           {j.get("project_ref") for j in _zip_jobs.values()} == {"ZIP-001"} and
           all(any("every PDF queued" in f for f in j.get("flags", []))
               for j in _zip_jobs.values()))

        _candidate_job_id = "99999999-9999-4999-8999-999999999999"
        _candidate_records = [
            {"candidate_id":"office-p0-level-00-1", "page":0, "level":0,
             "category":"ground_floor", "boq_scope":"ground_floor_core",
             "polygon_pts":[[0,0],[100,0],[100,100],[0,100]]},
            {"candidate_id":"office-p0-level-01-1", "page":0, "level":1,
             "category":"upper_floor", "boq_scope":"main_upper_floor",
             "polygon_pts":[[200,0],[300,0],[300,100],[200,100]]},
        ]
        _AS.save_jobs({_candidate_job_id: {
            "id":_candidate_job_id, "status":"pending", "decision":None,
            "measurement_state":"UNMEASURED", "scale_confirmed":False,
            "candidate_polygons":_candidate_records,
            "result":{"file":"Office-GA.pdf", "area_m2":None,
                      "measurement_state":"UNMEASURED",
                      "candidate_polygons":_candidate_records},
        }})
        _stale_candidate_resp = _client_up.post(f"/adjust/{_candidate_job_id}", json={
            "regions":[[[0,0],[100,0],[100,100],[0,100]]], "scale_k":0.1,
            "candidate_ids":["office-p0-level-99-1"],
        })
        ck("assisted adjustment rejects stale/unknown candidate IDs",
           _stale_candidate_resp.status_code == 409, _stale_candidate_resp.get_json())
        _multi_region_resp = _client_up.post(f"/adjust/{_candidate_job_id}", json={
            "regions":[[[0,0],[100,0],[100,100],[0,100]],
                       [[200,0],[300,0],[300,100],[200,100]]],
            "scale_k":0.1,
            "candidate_ids":["office-p0-level-00-1", "office-p0-level-01-1"],
            "region_categories":["ground_floor", "upper_floor"],
            "region_scopes":["ground_floor_core", "main_upper_floor"],
            "note":"assessor accepted two Office GA regions",
        })
        _multi_region_job = _AS.load_jobs()[_candidate_job_id]
        ck("assessor adjustment measures several Office regions in one atomic decision",
           _multi_region_resp.status_code == 200 and
           _multi_region_resp.get_json()["area_m2"] == 200.0 and
           _multi_region_resp.get_json()["region_count"] == 2 and
           len(_multi_region_job["adjusted"]["regions"]) == 2 and
           _multi_region_job["adjusted"]["candidate_ids"] ==
           ["office-p0-level-00-1", "office-p0-level-01-1"],
           _multi_region_resp.get_json())
        ck("candidate selection alone stayed UNMEASURED; assessor POST performs confirmation",
           _multi_region_job["scale_confirmed"] is True and
           _multi_region_job["measurement_state"] == "MEASURED_VERIFIED")
        ck("assisted Office adjustment persists one real BOQ zone per supplied region category",
           {zone["category"]:zone["area_m2"] for zone in _multi_region_job["zones"]} ==
           {"ground_floor":100.0, "upper_floor":100.0} and
           not _multi_region_job.get("zone_allocation_stale") and
           _multi_region_job["adjusted"]["region_categories"] ==
           ["ground_floor", "upper_floor"] and
           _multi_region_job["adjusted"]["region_scopes"] ==
           ["ground_floor_core", "main_upper_floor"], _multi_region_job.get("zones"))
        _office_adjust_quote = _AS._quotation_for_job(_candidate_job_id)
        ck("real Office region categories override the filename and reach separate BOQ sections",
           [spec["section"] for spec in _office_adjust_quote["specifications"]] ==
           ["Ground floor slabs", "Upper floor slabs"],
           _office_adjust_quote["specifications"])
        _manual_unknown_resp = _client_up.post(f"/adjust/{_candidate_job_id}", json={
            "regions":[[[0,0],[100,0],[100,100],[0,100]]], "scale_k":0.1,
            "region_categories":["unclassified"],
            "note":"manual outline; floor level not resolved",
        })
        _manual_unknown_job = _AS.load_jobs()[_candidate_job_id]
        ck("unresolved manual region is preserved as unclassified and visibly blocks approval",
           _manual_unknown_resp.status_code == 200 and
           _manual_unknown_job["zones"][0]["area_m2"] == 100.0 and
           _manual_unknown_job["zones"][0]["category"] == "unclassified" and
           _manual_unknown_job["zone_classification_required"] and
           "unclassified" in (_AS._approve_block_reason(_manual_unknown_job) or ""),
           _manual_unknown_job.get("zones"))

        _channel_job_id = "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa"
        _channel_zones = [
            {"zone_key":"external_yard", "category":"external_yard", "area_m2":100.0},
            {"zone_key":"dock", "category":"dock", "area_m2":20.0,
             "loading_face_lm":30.0, "loading_face_pts":[[10,10],[10,310]]},
        ]
        _channel_proposals = [
            {"proposal_id":"channel-dock-loading-face", "component":"dock_retaining_wall",
             "proposed_length_lm":30.0, "polyline_pts":[[10,10],[10,310]],
             "assumed":True, "requires_assessor_confirmation":True},
            {"proposal_id":"channel-yard-longest-contained-run",
             "component":"yard_longest_contained_run", "proposed_length_lm":90.0,
             "polyline_pts":[[20,20],[920,20]], "assumed":True,
             "requires_assessor_confirmation":True},
        ]
        _channel_costing = _copy.deepcopy(_demo_result["costing"])
        _AS.save_jobs({_channel_job_id: {
            "id":_channel_job_id, "status":"pending", "decision":None,
            "measurement_state":"MEASURED_VERIFIED", "scale_confirmed":False,
            "zones":_copy.deepcopy(_channel_zones),
            "channel_proposals":_copy.deepcopy(_channel_proposals),
            "channel_proposal_decisions":{}, "costing":_copy.deepcopy(_channel_costing),
            "result":{"file":"Raw External.pdf", "area_m2":120.0,
                      "scale_k":0.1,
                      "measurement_state":"MEASURED_VERIFIED",
                      "zones":_copy.deepcopy(_channel_zones),
                      "channel_proposals":_copy.deepcopy(_channel_proposals),
                      "channel_proposal_decisions":{},
                      "costing":_copy.deepcopy(_channel_costing)},
        }})
        ck("unreviewed assumed channel proposals block approval",
           "accept/edit/remove" in (_AS._approve_block_reason(
               _AS.load_jobs()[_channel_job_id]) or ""))
        _channel_first_resp = _client_up.post(
            f"/channel-proposals/{_channel_job_id}", json={"decisions":[{
                "proposal_id":"channel-dock-loading-face", "action":"accept",
                "length_lm":31.25,
            }]})
        _channel_first_job = _AS.load_jobs()[_channel_job_id]
        ck("assessor can edit and accept a proposed channel length without measuring it",
           _channel_first_resp.status_code == 200 and
           not _channel_first_resp.get_json()["review_complete"] and
           _channel_first_job["channel_proposal_decisions"][
               "channel-dock-loading-face"]["length_lm"] == 31.25 and
           _channel_first_job["channel_proposal_decisions"][
               "channel-dock-loading-face"]["edited"] is True)
        ck("one pending channel proposal continues to block approval",
           _AS._approve_block_reason(_channel_first_job) is not None)
        _channel_geometry_resp = _client_up.post(
            f"/channel-proposals/{_channel_job_id}", json={"decisions":[{
                "proposal_id":"channel-dock-loading-face", "action":"accept",
                # Deliberately contradictory browser length: server must derive 35m from
                # the two PDF-point endpoints and stored 0.1m/pt scale.
                "length_lm":1.0, "polyline_pts":[[10,10],[360,10]],
            }]})
        _channel_geometry_job = _AS.load_jobs()[_channel_job_id]
        _channel_geometry_decision = _channel_geometry_job["channel_proposal_decisions"][
            "channel-dock-loading-face"]
        ck("assessor can drag a channel endpoint; server persists axis-aligned geometry",
           _channel_geometry_resp.status_code == 200 and
           _channel_geometry_decision["polyline_pts"] == [[10.0,10.0],[360.0,10.0]] and
           _channel_geometry_decision["geometry_edited"] is True,
           _channel_geometry_decision)
        ck("edited channel length is derived from geometry and scale, not browser arithmetic",
           _channel_geometry_decision["length_lm"] == 35.0,
           _channel_geometry_decision)
        _channel_diagonal_resp = _client_up.post(
            f"/channel-proposals/{_channel_job_id}", json={"decisions":[{
                "proposal_id":"channel-yard-longest-contained-run", "action":"accept",
                "length_lm":90.0, "polyline_pts":[[20,20],[920,120]],
            }]})
        ck("channel edit API refuses diagonal geometry rather than storing it",
           _channel_diagonal_resp.status_code == 400 and
           "non-diagonal" in _channel_diagonal_resp.get_json()["error"] and
           "channel-yard-longest-contained-run" not in
               _AS.load_jobs()[_channel_job_id]["channel_proposal_decisions"],
           _channel_diagonal_resp.get_json())
        _channel_second_resp = _client_up.post(
            f"/channel-proposals/{_channel_job_id}", json={"decisions":[{
                "proposal_id":"channel-yard-longest-contained-run", "action":"remove",
            }]})
        _channel_reviewed_job = _AS.load_jobs()[_channel_job_id]
        ck("assessor remove completes channel-proposal review and releases its gate",
           _channel_second_resp.status_code == 200 and
           _channel_second_resp.get_json()["review_complete"] and
           _AS._approve_block_reason(_channel_reviewed_job) is None)
        ck("channel decisions never enter measured zones or an approvable costing total",
           _channel_reviewed_job["zones"] == _channel_zones and
           _channel_reviewed_job["result"]["zones"] == _channel_zones and
           not any(z.get("category") == "channel" for z in _channel_reviewed_job["zones"]) and
           _channel_reviewed_job["costing"] == _channel_costing and
           _channel_reviewed_job["result"]["costing"] == _channel_costing)
        _channel_approve = _client_up.post(f"/approve/{_channel_job_id}", json={
            "note":"channel assumptions reviewed",
        })
        _channel_approved_job = _AS.load_jobs()[_channel_job_id]
        _channel_quote_paths = _channel_approved_job.get("quotation_paths") or {}
        _channel_saved_quote = (__import__("json").loads(
            Path(_channel_quote_paths["json"]).read_text())
            if _channel_quote_paths.get("json") and
            Path(_channel_quote_paths["json"]).exists() else {})
        _channel_saved_rows = [item for item in _channel_saved_quote.get("line_items", [])
                               if item.get("description") == FORTEL_CHANNEL_ROW]
        ck("approved server quotation carries accepted channel Lm with blank assessor rate",
           _channel_approve.status_code == 200 and len(_channel_saved_rows) == 1 and
           _channel_saved_rows[0]["qty"] == 35.0 and
           _channel_saved_rows[0]["rate"] is None and
           _channel_saved_rows[0]["value"] is None and
           _channel_saved_rows[0]["provisional"],
           _channel_saved_rows)

        _transition_job_id = "bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb"
        _transition_candidates = [
            {"candidate_id":"transition-yard-region-1", "region_id":"yard-region-1",
             "category":"transition", "proposed_length_lm":14.0,
             "polyline_pts":[[10,10],[150,10]], "assumed":True,
             "basis":"macadam-to-concrete Yard entrance boundary"},
            {"candidate_id":"transition-yard-region-2", "region_id":"yard-region-2",
             "category":"transition", "proposed_length_lm":12.0,
             "polyline_pts":[[20,20],[140,20]], "assumed":True,
             "basis":"macadam-to-concrete Yard entrance boundary"},
        ]
        _AS.save_jobs({_transition_job_id: {
            "id":_transition_job_id, "status":"pending", "decision":None,
            "measurement_state":"MEASURED_VERIFIED", "scale_confirmed":False,
            "zones":_copy.deepcopy(_channel_zones),
            "transition_candidates":_copy.deepcopy(_transition_candidates),
            "transition_candidate_decisions":{},
            "accepted_transition_quantities":[],
            "costing":_copy.deepcopy(_channel_costing),
            "result":{"file":"Raw External transitions.pdf", "area_m2":120.0,
                      "scale_k":0.1, "measurement_state":"MEASURED_VERIFIED",
                      "zones":_copy.deepcopy(_channel_zones),
                      "transition_candidates":_copy.deepcopy(_transition_candidates),
                      "transition_candidate_decisions":{},
                      "accepted_transition_quantities":[],
                      "costing":_copy.deepcopy(_channel_costing)},
        }})
        ck("unreviewed assumed Transition candidates block approval",
           "accept/edit/remove" in (_AS._approve_block_reason(
               _AS.load_jobs()[_transition_job_id]) or ""))
        _transition_accept_resp = _client_up.post(
            f"/transition-candidates/{_transition_job_id}", json={"decisions":[{
                "candidate_id":"transition-yard-region-1", "action":"accept",
                "length_lm":15.25,
            }]})
        _transition_accept_job = _AS.load_jobs()[_transition_job_id]
        ck("Transition endpoint persists assessor-edited accepted quantity",
           _transition_accept_resp.status_code == 200 and
           not _transition_accept_resp.get_json()["review_complete"] and
           _transition_accept_job["transition_candidate_decisions"][
               "transition-yard-region-1"]["length_lm"] == 15.25 and
           _transition_accept_job["transition_candidate_decisions"][
               "transition-yard-region-1"]["edited"] is True and
           _transition_accept_job["accepted_transition_quantities"] == [{
               "candidate_id":"transition-yard-region-1",
               "region_id":"yard-region-1", "category":"transition",
               "measurement_kind":"length", "length_lm":15.25, "unit":"Lm",
               "assumed":True, "provisional":True,
               "basis":"macadam-to-concrete Yard entrance boundary",
               "source":None, "assessor_edited":True,
           }], _transition_accept_resp.get_json())
        ck("one pending Transition candidate continues to block approval",
           _AS._approve_block_reason(_transition_accept_job) is not None)
        _transition_remove_resp = _client_up.post(
            f"/transition-candidates/{_transition_job_id}", json={"decisions":[{
                "candidate_id":"transition-yard-region-2", "action":"remove",
            }]})
        _transition_reviewed_job = _AS.load_jobs()[_transition_job_id]
        ck("Transition remove completes review and releases the approval gate",
           _transition_remove_resp.status_code == 200 and
           _transition_remove_resp.get_json()["review_complete"] and
           _AS._approve_block_reason(_transition_reviewed_job) is None)
        ck("accepted Transition remains provisional state, never measured zone or costing",
           not any(zone.get("category") == "transition"
                   for zone in _transition_reviewed_job["zones"]) and
           _transition_reviewed_job["costing"] == _channel_costing and
           _transition_reviewed_job["accepted_transition_quantities"][0][
               "provisional"] is True)
        _transition_quote = _AS._quotation_for_job(_transition_job_id)
        _transition_rows = [item for item in _transition_quote["line_items"]
                            if item.get("description") == FORTEL_TRANSITION_ROW]
        ck("persisted accepted Transition reaches server quotation with blank rate",
           len(_transition_rows) == 1 and _transition_rows[0]["qty"] == 15.25 and
           _transition_rows[0]["rate"] is None and
           _transition_rows[0]["value"] is None and
           _transition_rows[0]["provisional"], _transition_rows)

        # Critical assisted-loop regression: begin with a real /upload-created record, then
        # install the completed zoned MEASURED_UNVERIFIED takeoff that the background worker
        # would save. Confirming the existing scale/extent must preserve both zones, release
        # the gate, approve, and write a real quotation — no destructive /adjust in between.
        _AS.save_jobs({})
        _started_up.clear()
        _loop_upload = _client_up.post("/upload", data={
            "project_ref":"E2E-ZONE-001", "project_name":"Zoned confirmation loop",
            "client_name":"Fortel QA",
            "pdf":(_io3.BytesIO(_pdf_a_bytes), "External Unit-1.pdf"),
        }, content_type="multipart/form-data")
        _loop_job_id = _loop_upload.get_json()["job_id"]
        _loop_zones = [
            {"zone_key":"external_yard", "category":"external_yard", "area_m2":100.0},
            {"zone_key":"dock", "category":"dock", "area_m2":20.0},
        ]
        _loop_jobs = _AS.load_jobs()
        _loop_jobs[_loop_job_id].update({
            "status":"pending", "measurement_state":"MEASURED_UNVERIFIED",
            "scale_confirmed":False, "zone_allocation_stale":True,
            "zones":_copy.deepcopy(_loop_zones),
            "flags":["ZONE ALLOCATION STALE: retained zones await assessor confirmation"],
            "result":{
                "file":"E2E-ZONE-001_External Unit-1.pdf", "area_m2":120.0,
                "scale_k":0.1, "measurement_state":"MEASURED_UNVERIFIED",
                "zones":_copy.deepcopy(_loop_zones), "zone_allocation_stale":True,
                "flags":["ZONE ALLOCATION STALE: retained zones await assessor confirmation"],
            },
        })
        _AS.save_jobs(_loop_jobs)
        _loop_confirm = _client_up.post(f"/confirm-measurement/{_loop_job_id}", json={
            "confirm_scale_extent":True, "note":"scale and coloured extents checked",
        })
        _loop_confirmed_job = _AS.load_jobs()[_loop_job_id]
        ck("existing scale+extent confirmation preserves zoned measurement and clears stale gate",
           _loop_confirm.status_code == 200 and
           _loop_confirm.get_json()["zone_count"] == 2 and
           _loop_confirmed_job["zones"] == _loop_zones and
           _loop_confirmed_job["result"]["zones"] == _loop_zones and
           _loop_confirmed_job["scale_confirmed"] is True and
           not _loop_confirmed_job["zone_allocation_stale"] and
           _AS._approve_block_reason(_loop_confirmed_job) is None,
           _loop_confirm.get_json())
        _loop_approve = _client_up.post(f"/approve/{_loop_job_id}", json={
            "note":"approved after non-destructive confirmation",
        })
        _loop_approved_job = _AS.load_jobs()[_loop_job_id]
        _loop_paths = _loop_approved_job.get("quotation_paths") or {}
        _loop_quote = (__import__("json").loads(Path(_loop_paths["json"]).read_text())
                       if _loop_paths.get("json") and Path(_loop_paths["json"]).exists() else {})
        ck("E2E upload -> confirm -> approve terminates in an APPROVED zoned quotation",
           _loop_approve.status_code == 200 and
           _loop_approved_job["status"] == "approved" and
           Path(_loop_paths.get("xlsx", "missing")).exists() and
           [spec["section"] for spec in _loop_quote.get("specifications", [])] ==
           ["External yard slabs", "Dock slabs"],
            {"confirm":_loop_confirm.get_json(), "approve":_loop_approve.get_json(),
             "status":_loop_approved_job.get("status"), "quotation_paths":_loop_paths})

        # Money-path regression: assessor geometry, cut-outs, and scale must replace the AI
        # inputs everywhere that can be costed or quoted. These exercise the real Flask routes
        # and then inspect the generated quotation, rather than trusting only /adjust's reply.
        def _assessor_truth_job(job_id, project_ref, original_area=6816.0):
            original_costing = _copy.deepcopy(_demo_result["costing"])
            original_costing["area_m2"] = original_area
            zones = [{"zone_key":"external_yard", "category":"external_yard",
                      "area_m2":original_area}]
            result = {
                "file":f"{project_ref}_External.pdf", "type":"UNMARKED vector",
                "source_discipline":"architect", "area_m2":original_area,
                "scale_k":0.17639, "measurement_state":"MEASURED_UNVERIFIED",
                "zones":_copy.deepcopy(zones), "flags":[],
                "costing":_copy.deepcopy(original_costing),
            }
            return {
                "id":job_id, "status":"pending", "decision":None,
                "project_ref":project_ref, "project_name":"Assessor truth regression",
                "client_name":"Fortel QA", "created_at":"2026-08-12T00:00:00",
                "measurement_state":"MEASURED_UNVERIFIED", "scale_confirmed":False,
                "zones":_copy.deepcopy(zones), "costing":original_costing,
                "result":result,
            }

        def _confirm_and_adjust(job_id, *, width=200, cutouts=None, channels=None):
            blocked_before = _AS._approve_block_reason(_AS.load_jobs()[job_id]) is not None
            confirmed_response = _client_up.post(
                f"/confirm-measurement/{job_id}", json={
                    "confirm_scale_extent":True, "note":"assessor confirmed extent",
                })
            adjusted_response = _client_up.post(f"/adjust/{job_id}", json={
                "regions":[[[100,100],[100 + width,100],[100 + width,200],[100,200]]],
                "region_categories":["external_yard"], "region_scopes":["main"],
                "scale_k":0.1, "cutout_regions":cutouts or [],
                "user_channels":channels or [], "note":"categorized assessor trace",
            })
            return blocked_before, confirmed_response, adjusted_response

        _money_cutout_id = "91000000-0000-4000-8000-000000000001"
        _AS.save_jobs({_money_cutout_id:_assessor_truth_job(
            _money_cutout_id, "MONEY-CUTOUT")})
        _cutout_gate, _cutout_confirm, _cutout_adjust = _confirm_and_adjust(
            _money_cutout_id,
            cutouts=[[[150,120],[200,120],[200,140],[150,140]]])
        _cutout_job = _AS.load_jobs()[_money_cutout_id]
        _cutout_quote = _AS._quotation_for_job(_money_cutout_id)
        _cutout_slab = next(item for item in _cutout_quote["line_items"]
                            if item.get("line_role") == "concrete_slab")
        _cutout_json = __import__("json").loads(quotation_json(_cutout_quote))
        _cutout_json_slab = next(item for item in _cutout_json["line_items"]
                                 if item.get("line_role") == "concrete_slab")
        _cutout_text = quotation_text(_cutout_quote)
        _cutout_html = quotation_html(_cutout_quote)
        _cutout_text_slab = next(line for line in _cutout_text.splitlines()
                                  if "Concrete Slabs" in line)
        _cutout_wb = _load_workbook(_BytesIO(quotation_xlsx(_cutout_quote)), data_only=False)
        _cutout_ws = _cutout_wb["REV_01"]
        _cutout_xlsx_source_values = [
            _cutout_ws.cell(row, 2).value for row in range(1, _cutout_ws.max_row + 1)
            if "External.pdf" in str(_cutout_ws.cell(row, 1).value or "")
        ]
        ck("assessor cut-out net area reaches zones and every quotation format",
           _cutout_gate and _cutout_confirm.status_code == 200 and
           _cutout_adjust.get_json()["area_m2"] == 190.0 and
           _cutout_job["zones"][0]["area_m2"] == 190.0 and
           _cutout_slab["qty"] == 190.0 and _cutout_json_slab["qty"] == 190.0 and
           "190" in _cutout_text_slab and ">190 m²</td>" in _cutout_html and
           190.0 in _cutout_xlsx_source_values and
           200.0 not in _cutout_xlsx_source_values,
           {"adjust":_cutout_adjust.get_json(), "zones":_cutout_job["zones"],
            "slab":_cutout_slab, "xlsx_sources":_cutout_xlsx_source_values})
        ck("assessor adjustment preserves four-state gate and explicit verification",
           _cutout_job["measurement_state"] == "MEASURED_VERIFIED" and
           _cutout_job["scale_confirmed"] is True and
           _AS._approve_block_reason(_cutout_job) is None)

        _money_approve_id = "91000000-0000-4000-8000-000000000002"
        _AS.save_jobs({_money_approve_id:_assessor_truth_job(
            _money_approve_id, "MONEY-APPROVE")})
        _confirm_and_adjust(_money_approve_id, width=190)
        _approve_response = _client_up.post(f"/approve/{_money_approve_id}", json={
            "note":"approve assessor correction",
        })
        _approved_truth_job = _AS.load_jobs()[_money_approve_id]
        _approved_truth_quote = _AS._quotation_for_job(_money_approve_id)
        _approved_truth_slab = next(item for item in _approved_truth_quote["line_items"]
                                    if item.get("line_role") == "concrete_slab")
        ck("approval costs and quotes assessor-adjusted area instead of original AI area",
           _approve_response.status_code == 200 and
           _approve_response.get_json()["costing"]["area_m2"] == 190.0 and
           _approved_truth_job["costing"]["area_m2"] == 190.0 and
           _approved_truth_slab["qty"] == 190.0,
           {"approve":_approve_response.get_json(), "slab":_approved_truth_slab})

        _money_channel_id = "91000000-0000-4000-8000-000000000003"
        _AS.save_jobs({_money_channel_id:_assessor_truth_job(
            _money_channel_id, "MONEY-CHANNEL")})
        _confirm_and_adjust(
            _money_channel_id, channels=[[[100,250],[200,250]]])
        _channel_truth_quote = _AS._quotation_for_job(_money_channel_id)
        _channel_truth_row = next(item for item in _channel_truth_quote["line_items"]
                                  if item["description"] == FORTEL_CHANNEL_ROW)
        ck("assessor channel quotation length uses assessor scale travelling with geometry",
           _channel_truth_row["qty"] == 10.0 and
           _AS._quotation_result_for_job(
               _AS.load_jobs()[_money_channel_id])["scale_k"] == 0.1,
           _channel_truth_row)

        _money_bend_id = "91000000-0000-4000-8000-000000000005"
        _AS.save_jobs({_money_bend_id:_assessor_truth_job(
            _money_bend_id, "MONEY-BENDING-CHANNEL")})
        _, _bend_confirm, _bend_adjust = _confirm_and_adjust(
            _money_bend_id,
            # Two segments: 30 px + 40 px at assessor k=0.1 m/px => 7.00 Lm.
            channels=[[[100,250],[130,250],[130,290]]],
        )
        _bend_job = _AS.load_jobs()[_money_bend_id]
        _bend_quote = _AS._quotation_for_job(_money_bend_id)
        _bend_row = next(item for item in _bend_quote["line_items"]
                         if item["description"] == FORTEL_CHANNEL_ROW)
        ck("assessor can persist a bending channel polyline instead of a forced straight chord",
           _bend_confirm.status_code == 200 and _bend_adjust.status_code == 200 and
           _bend_job["adjusted"]["user_channels"] ==
               [[[100,250],[130,250],[130,290]]],
           _bend_adjust.get_json())
        ck("quotation sums every segment of the assessor channel polyline",
           _bend_row["qty"] == 7.0, _bend_row)

        _money_plain_id = "91000000-0000-4000-8000-000000000004"
        _plain_job = _assessor_truth_job(_money_plain_id, "MONEY-PLAIN", original_area=125.0)
        _plain_job["measurement_state"] = "MEASURED_VERIFIED"
        _plain_job["scale_confirmed"] = True
        _plain_job["result"]["measurement_state"] = "MEASURED_VERIFIED"
        _AS.save_jobs({_money_plain_id:_plain_job})
        _plain_approve = _client_up.post(f"/approve/{_money_plain_id}", json={
            "note":"unchanged normal path",
        })
        _plain_saved = _AS.load_jobs()[_money_plain_id]
        _plain_quote = _AS._quotation_for_job(_money_plain_id)
        _plain_slab = next(item for item in _plain_quote["line_items"]
                           if item.get("line_role") == "concrete_slab")
        ck("unadjusted verified job keeps its original pricing path unchanged",
           _plain_approve.status_code == 200 and
           _plain_approve.get_json()["costing"]["area_m2"] == 125.0 and
           _plain_saved["costing"]["area_m2"] == 125.0 and
           _plain_slab["qty"] == 125.0,
           {"approve":_plain_approve.get_json(), "slab":_plain_slab})

        _route_costing_a = _copy.deepcopy(_demo_result["costing"])
        _route_costing_a.update({"area_m2": 100, "assumed": True})
        _route_costing_b = _copy.deepcopy(_demo_result["costing"])
        _route_costing_b.update({"area_m2": 150, "assumed": True})
        _route_jobs = {
            "11111111-1111-4111-8111-111111111111": {
                "id": "11111111-1111-4111-8111-111111111111", "decision": "approved",
                "status": "approved", "project_ref": "QUOTE-MULTI-001",
                "project_name": "Two Yard Units", "client_name": "Fortel QA",
                "created_at": "2026-07-15T10:00:00",
                "costing": _route_costing_a,
                "result": {"file": "Yard-A.pdf", "quotation_section": "External yard slabs",
                           "area_m2": 100, "costing": _route_costing_a, "flags": []},
            },
            "22222222-2222-4222-8222-222222222222": {
                "id": "22222222-2222-4222-8222-222222222222", "decision": "adjusted",
                "status": "adjusted", "project_ref": "QUOTE-MULTI-001",
                "project_name": "Two Yard Units", "client_name": "Fortel QA",
                "created_at": "2026-07-15T10:01:00",
                "costing": _route_costing_b,
                "result": {"file": "Yard-B.pdf", "quotation_section": "External yard slabs",
                           "area_m2": 150, "costing": _route_costing_b, "flags": []},
            },
        }
        _AS.save_jobs(_route_jobs)
        _xlsx_route_resp = _client_up.get(
            "/quotation/11111111-1111-4111-8111-111111111111.xlsx")
        _xlsx_route_wb = _load_workbook(_BytesIO(_xlsx_route_resp.data), data_only=False)
        _xlsx_route_ws = _xlsx_route_wb["REV_01"]
        _xlsx_route_slab_row = next(
            row for row in range(1, _xlsx_route_ws.max_row + 1)
            if "Concrete Slabs" in str(_xlsx_route_ws.cell(row, 1).value or ""))
        ck("xlsx download route returns a valid attachment",
           _xlsx_route_resp.status_code == 200 and
           _xlsx_route_resp.mimetype ==
           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" and
           "QUOTE-MULTI-001.xlsx" in _xlsx_route_resp.headers.get("Content-Disposition", ""))
        ck("xlsx route aggregates approved sibling units sharing project_ref",
           _xlsx_route_ws.cell(_xlsx_route_slab_row, 2).data_type == "f" and
           any(_xlsx_route_ws.cell(row, 1).value == "Total Area Take Off:" and
               _xlsx_route_ws.cell(row, 2).data_type == "f"
               for row in range(1, _xlsx_route_ws.max_row + 1)) and
           sorted(float(_xlsx_route_ws.cell(row, 2).value)
                  for row in range(1, _xlsx_route_ws.max_row + 1)
                  if _xlsx_route_ws.cell(row, 1).value in ("Yard-A.pdf", "Yard-B.pdf"))
           == [100.0, 150.0])

        # Aryan field report 17 Jul: uploading a fresh case produced a SEPARATE xlsx per
        # document (pending siblings were excluded from aggregation), and unmeasured office
        # GA plans (line/hatch -> assessor trace) vanished from the output entirely. The
        # case quotation must be ONE workbook: pending-but-measured siblings included
        # (marked provisional), unmeasured documents listed as awaiting trace — never absent.
        _pend_costing_a = _copy.deepcopy(_demo_result["costing"]); _pend_costing_a.update({"area_m2": 100, "assumed": True})
        _pend_costing_b = _copy.deepcopy(_demo_result["costing"]); _pend_costing_b.update({"area_m2": 150, "assumed": True})
        _pend_jobs = {
            "aaaaaaaa-1111-4111-8111-111111111111": {
                "id": "aaaaaaaa-1111-4111-8111-111111111111", "decision": None,
                "status": "pending", "project_ref": "CASE-PEND-001",
                "project_name": "Fresh Case", "client_name": "Fortel QA",
                "created_at": "2026-07-17T10:00:00", "costing": _pend_costing_a,
                "result": {"file": "Yard-A.pdf", "quotation_section": "External yard slabs",
                           "area_m2": 100, "costing": _pend_costing_a, "flags": []},
            },
            "bbbbbbbb-2222-4222-8222-222222222222": {
                "id": "bbbbbbbb-2222-4222-8222-222222222222", "decision": None,
                "status": "pending", "project_ref": "CASE-PEND-001",
                "project_name": "Fresh Case", "client_name": "Fortel QA",
                "created_at": "2026-07-17T10:01:00", "costing": _pend_costing_b,
                "result": {"file": "Yard-B.pdf", "quotation_section": "External yard slabs",
                           "area_m2": 150, "costing": _pend_costing_b, "flags": []},
            },
            "cccccccc-3333-4333-8333-333333333333": {
                "id": "cccccccc-3333-4333-8333-333333333333", "decision": None,
                "status": "pending", "project_ref": "CASE-PEND-001",
                "project_name": "Fresh Case", "client_name": "Fortel QA",
                "created_at": "2026-07-17T10:02:00",
                "result": {"file": "Office-Floors-U1.pdf", "area_m2": None,
                           "measurement_state": "UNMEASURED",
                           "flags": ["NON-COLOUR-CODED (line/hatch) drawing — assessor trace"]},
            },
        }
        _AS.save_jobs(_pend_jobs)
        _pend_resp = _client_up.get("/quotation/aaaaaaaa-1111-4111-8111-111111111111.xlsx")
        ck("case xlsx succeeds for a fresh (all-pending) case", _pend_resp.status_code == 200,
           _pend_resp.status_code)
        _pend_ws = _load_workbook(_BytesIO(_pend_resp.data), data_only=False)["REV_01"]
        _pend_cells = [str(_pend_ws.cell(r, 1).value or "") for r in range(1, _pend_ws.max_row + 1)]
        ck("pending-but-measured siblings aggregate into ONE case workbook",
           any("Yard-A.pdf" in c for c in _pend_cells) and any("Yard-B.pdf" in c for c in _pend_cells))
        ck("unmeasured document is LISTED in the case workbook (never silently absent)",
           any("Office-Floors-U1.pdf" in c and "NOT YET MEASURED" in c for c in _pend_cells),
           [c for c in _pend_cells if "Office" in c])
        ck("pending quantities are marked provisional pending approval",
           any("not yet" in c.lower() and "approved" in c.lower() for c in _pend_cells))
        _pend_json = _client_up.get("/quotation/aaaaaaaa-1111-4111-8111-111111111111.json").get_json()
        ck("case quotation JSON carries the unmeasured document list",
           any(d.get("file") == "Office-Floors-U1.pdf" for d in _pend_json.get("unmeasured", [])))

        # Aryan follow-up 18 Jul: a case containing ONLY unmeasurable docs (e.g. office GA
        # plans, all awaiting assessor trace) previously 400'd on download — the office
        # drawings "still skipped". An office-only case must still yield the case workbook.
        _office_only = {
            "dddddddd-4444-4444-8444-444444444444": {
                "id": "dddddddd-4444-4444-8444-444444444444", "decision": None,
                "status": "pending", "project_ref": "CASE-OFFICE-ONLY",
                "project_name": "Office Only", "client_name": "Fortel QA",
                "created_at": "2026-07-18T09:00:00",
                "result": {"file": "Office-GA-L00.pdf", "area_m2": None,
                           "measurement_state": "UNMEASURED", "flags": ["line/hatch"]},
            },
            "eeeeeeee-5555-4555-8555-555555555555": {
                "id": "eeeeeeee-5555-4555-8555-555555555555", "decision": None,
                "status": "pending", "project_ref": "CASE-OFFICE-ONLY",
                "project_name": "Office Only", "client_name": "Fortel QA",
                "created_at": "2026-07-18T09:01:00",
                "result": {"file": "Office-GA-L01.pdf", "area_m2": None,
                           "measurement_state": "UNMEASURED", "flags": ["line/hatch"]},
            },
        }
        _AS.save_jobs(_office_only)
        _oo_resp = _client_up.get("/quotation/dddddddd-4444-4444-8444-444444444444.xlsx")
        ck("office-only case still yields the case workbook (no 400)",
           _oo_resp.status_code == 200, _oo_resp.status_code)
        _oo_ws = _load_workbook(_BytesIO(_oo_resp.data), data_only=False)["REV_01"]
        _oo_cells = [str(_oo_ws.cell(r, 1).value or "") for r in range(1, _oo_ws.max_row + 1)]
        ck("BOTH office docs listed as NOT YET MEASURED in the office-only workbook",
           all(any(f in c and "NOT YET MEASURED" in c for c in _oo_cells)
               for f in ("Office-GA-L00.pdf", "Office-GA-L01.pdf")), _oo_cells[-6:])
        # Unit labels derived from Fortel filename conventions (per-unit BOQ rows)
        from quotation import _unit_label_from_filename as _ulf
        ck("unit label: 'External Markup Unit-1.pdf' + D-ref context -> 'Unit 1'",
           _ulf("External Markup Unit-1.pdf") == "Unit 1")
        ck("unit label: D-ref included when present",
           _ulf("Unit_3 D410 Hard Landscaping.pdf") == "Unit 3 (D410)")
        ck("unit label: no unit pattern -> None (caller keeps filename)",
           _ulf("Proposed_Site_Plan.pdf") is None)

        # /upload stores enquiry identification (subject/body) on every job in the batch
        _AS.save_jobs({})
        _id_resp = _client_up.post("/upload", data={
            "project_ref": "IDENT-1", "project_name": "Ident Case",
            "email_subject": "RE: Winwick tender enquiry",
            "email_body": "Please price the attached drawings.",
            "pdf": (_io3.BytesIO(_pdf_a_bytes), "yard.pdf"),
        }, content_type="multipart/form-data")
        _id_jobs = _AS.load_jobs()
        ck("upload stores email_subject/email_body for enquiry identification",
           _id_resp.status_code in (201, 202) and
           all(j.get("email_subject") == "RE: Winwick tender enquiry" and
               "attached drawings" in j.get("email_body", "") for j in _id_jobs.values()),
           list(_id_jobs.values())[:1])

        # count_manholes_marked: unreadable file -> None (couldn't check), never a silent 0
        from robust_takeoff import count_manholes_marked as _cmm
        ck("count_manholes_marked returns None (not 0) for an unreadable file",
           _cmm("/nonexistent/nope.pdf") is None)

        _AS.save_jobs(_route_jobs)   # restore the store for the spec-capture tests below

        # Fortel's supplied Brief_Spec is a blank checklist. Capture applicable fields
        # atomically without touching the job's four-state/measurement record; a partial
        # pricing spec remains assumed even though the unchanged calculation can re-price.
        _spec_job_id = "11111111-1111-4111-8111-111111111111"
        _spec_jobs_before = _AS.load_jobs()
        _spec_jobs_before[_spec_job_id]["measurement_state"] = "MEASURED_VERIFIED"
        _spec_jobs_before[_spec_job_id]["adjusted"] = {
            "area_m2": 100, "scale_k": 0.2,
            "polygon_pts": [[0, 0], [1, 0], [1, 1]],
        }
        _initial_spec_costing = _copy.deepcopy(_spec_jobs_before[_spec_job_id]["costing"])
        _initial_spec_paths = _AS._save_quotation(
            _spec_job_id,
            _AS._quotation_result_for_job(_spec_jobs_before[_spec_job_id]),
            _initial_spec_costing,
            file_stem="QUOTE-MULTI-001-REV_01",
        )
        _spec_jobs_before[_spec_job_id].update({
            "quotation_paths": _initial_spec_paths,
            "quotation_revision": 1,
            "quotation_status": "ready",
            "quotation_history": [{
                "revision": 1, "label": "REV_01",
                "issued_at": "2026-07-15T10:02:00",
                "reason": "initial approval", "paths": dict(_initial_spec_paths),
            }],
        })
        _initial_spec_bytes = {
            fmt: Path(path).read_bytes() for fmt, path in _initial_spec_paths.items()
        }
        _AS.save_jobs(_spec_jobs_before)
        _spec_resp = _client_up.post(f"/spec-override/{_spec_job_id}", json={
            "slab_type": "external_yard",
            "fields": {"depth_mm": 200, "conc_mix": None, "mesh": None, "layers": None,
                       "bay_sizes": "5m x 5m", "joint_details": None},
        })
        _spec_json = _spec_resp.get_json()
        _spec_saved_job = _AS.load_jobs()[_spec_job_id]
        ck("partial Brief_Spec capture succeeds but keeps costing provisional",
           _spec_resp.status_code == 200 and _spec_json["repriced"] is True and
           _spec_json["costing"]["assumed"] is True and
           not _spec_json["brief_spec"]["fields"]["depth_mm"]["provisional"] and
           _spec_json["brief_spec"]["fields"]["mesh"]["provisional"], _spec_json)
        ck("approved specification correction preserves four-state, geometry and assessor approval",
           _spec_saved_job["brief_spec"]["fields"]["bay_sizes"]["value"] == "5m x 5m" and
           _spec_saved_job["measurement_state"] == "MEASURED_VERIFIED" and
           _spec_saved_job["status"] == "approved" and
           _spec_saved_job["decision"] == "approved" and
           _spec_saved_job["adjusted"] == _spec_jobs_before[_spec_job_id]["adjusted"])
        ck("approved specification correction re-costs the assessor-approved area",
           _spec_json["post_approval_correction"] is True and
           _spec_saved_job["costing"]["area_m2"] == 100 and
           _spec_saved_job["costing"]["spec"]["depth_mm"] == 200 and
           _spec_saved_job["costing"]["total_gbp"] !=
               _initial_spec_costing["total_gbp"],
           {"before": _initial_spec_costing, "after": _spec_saved_job["costing"]})
        _revised_spec_paths = _spec_saved_job.get("quotation_paths") or {}
        ck("post-approval correction creates visible REV_02 files and immutable revision history",
           _spec_json["quotation_revision"] == 2 and
           _spec_saved_job["quotation_revision"] == 2 and
           [entry["label"] for entry in _spec_saved_job["quotation_history"]] ==
               ["REV_01", "REV_02"] and
           all("REV_02" in Path(path).stem for path in _revised_spec_paths.values()) and
           all(Path(_initial_spec_paths[fmt]).read_bytes() == original
               for fmt, original in _initial_spec_bytes.items()),
           _spec_saved_job.get("quotation_history"))
        _revised_spec_text = Path(_revised_spec_paths["txt"]).read_text()
        ck("corrected quotation visibly declares correction-after-approval and revision",
           "CORRECTION AFTER APPROVAL" in _revised_spec_text and
           "REV_02" in _revised_spec_text, _revised_spec_text[:500])
        _sibling_revised_text = quotation_text(_AS._quotation_for_job(
            "22222222-2222-4222-8222-222222222222"))
        ck("case correction remains REV_02 with a caveat when downloaded from a sibling drawing",
           "CORRECTION AFTER APPROVAL" in _sibling_revised_text and
           "REV_02" in _sibling_revised_text, _sibling_revised_text[:500])

        # Normal pipeline jobs carry per-zone checklists. Their approved correction route must
        # re-cost the corrected zone (and only that zone), otherwise the portal says REV_02 while
        # the mixed-zone quotation still contains the pre-correction blank/stale rate.
        from slab_spec import empty_brief_spec as _empty_zone_brief
        _zone_spec_job_id = "11111111-1111-4111-8111-111111111119"
        _zone_spec_zones = [
            {"category": "external_yard", "area_m2": 100.0, "measurement_kind": "area"},
            {"category": "dock", "area_m2": 20.0, "measurement_kind": "area"},
        ]
        _zone_briefs = {
            "external_yard": _empty_zone_brief("external_yard"),
            "dock": _empty_zone_brief("dock"),
        }
        _zone_result = {
            "file": "Approved Mixed Zones.pdf", "area_m2": 120.0,
            "measurement_state": "MEASURED_VERIFIED", "scale_k": 0.1,
            "zones": _zone_spec_zones, "brief_specs": _zone_briefs,
            "costing": _copy.deepcopy(_initial_spec_costing),
        }
        _zone_job = {
            "id": _zone_spec_job_id, "project_ref": "ZONE-SPEC-CORR-1",
            "project_name": "Zone correction QA", "status": "approved",
            "decision": "approved", "decided_at": "2026-07-15T11:00:00",
            "measurement_state": "MEASURED_VERIFIED", "scale_confirmed": True,
            "result": _zone_result, "zones": _zone_spec_zones,
            "brief_specs": _zone_briefs, "costing": _copy.deepcopy(_initial_spec_costing),
        }
        _zone_jobs = _AS.load_jobs()
        _zone_jobs[_zone_spec_job_id] = _zone_job
        _AS.save_jobs(_zone_jobs)
        _zone_rev1_paths = _AS._save_quotation(
            _zone_spec_job_id, _AS._quotation_result_for_job(_zone_job),
            _zone_job["costing"], file_stem="ZONE-SPEC-CORR-1-REV_01")
        _zone_jobs = _AS.load_jobs()
        _zone_jobs[_zone_spec_job_id].update({
            "quotation_paths": _zone_rev1_paths, "quotation_revision": 1,
            "quotation_status": "ready", "quotation_history": [{
                "revision": 1, "label": "REV_01", "issued_at": "2026-07-15T11:00:00",
                "reason": "initial approval", "paths": dict(_zone_rev1_paths),
            }],
        })
        _AS.save_jobs(_zone_jobs)
        _zone_spec_resp = _client_up.post(f"/spec-override/{_zone_spec_job_id}", json={
            "zone_category": "external_yard", "fields": {
                "depth_mm": 200, "conc_mix": "C32/40", "mesh": "A252", "layers": 1,
                "bay_sizes": None, "joint_details": None,
            },
        })
        _zone_spec_json = _zone_spec_resp.get_json()
        _zone_spec_saved = _AS.load_jobs()[_zone_spec_job_id]
        _zone_costing = (_zone_spec_saved.get("zone_costings") or {}).get("external_yard")
        ck("approved per-zone specification correction re-costs the exact assessor-scoped zone",
           _zone_spec_resp.status_code == 200 and _zone_spec_json["repriced"] is True and
           _zone_costing and _zone_costing["area_m2"] == 100.0 and
           _zone_costing["spec"]["depth_mm"] == 200 and
           "dock" not in (_zone_spec_saved.get("zone_costings") or {}),
           {"response": _zone_spec_json, "zone_costings": _zone_spec_saved.get("zone_costings")})
        _zone_quote = _AS._quotation_for_job(_zone_spec_job_id)
        _zone_yard_lines = [item for item in _zone_quote["line_items"]
                            if item["section"] == "External yard slabs"]
        _zone_dock_lines = [item for item in _zone_quote["line_items"]
                            if item["section"] == "Dock slabs"]
        ck("zone correction prices Yard in REV_02 while unsupplied Dock stays blank",
           _zone_spec_json["quotation_revision"] == 2 and
           _zone_yard_lines and isinstance(_zone_yard_lines[0].get("rate"), (int, float)) and
           _zone_dock_lines and _zone_dock_lines[0].get("rate") is None,
           {"yard": _zone_yard_lines, "dock": _zone_dock_lines})
        ck("per-zone correction never overwrites the legacy aggregate costing",
           _zone_spec_saved["costing"] == _zone_job["costing"],
           _zone_spec_saved["costing"])

        _bad_spec_resp = _client_up.post(f"/spec-override/{_spec_job_id}", json={
            "slab_type": "upper_floor", "fields": {"bay_sizes": "5m x 5m"},
        })
        ck("non-applicable Brief_Spec field is rejected cleanly",
           _bad_spec_resp.status_code == 400 and "does not apply" in
           (_bad_spec_resp.get_json().get("error") or ""), _bad_spec_resp.get_json())
        _unsupported_resp = _client_up.post(f"/spec-override/{_spec_job_id}", json={
            "slab_type": "external_yard",
            "fields": {"depth_mm": 200, "conc_mix": "Client mix", "mesh": "CLIENT-MESH",
                       "layers": 1, "bay_sizes": None, "joint_details": None},
        })
        _unsupported_json = _unsupported_resp.get_json()
        _unsupported_saved = _AS.load_jobs()[_spec_job_id]
        ck("unsupported open-text client spec is saved without inventing a rate",
           _unsupported_resp.status_code == 200 and
           _unsupported_json["repriced"] is False and
           bool(_unsupported_json["pricing_warning"]) and
           _unsupported_saved["brief_spec"]["fields"]["mesh"]["value"] == "CLIENT-MESH")
        ck("unsupported client spec hard-blocks approval for human pricing review",
           _AS._approve_block_reason(_unsupported_saved) is not None)
        _blocked_quote_resp = _client_up.get(f"/quotation/{_spec_job_id}.json")
        ck("unsupported client spec blocks stale quotation downloads even after prior decision",
           _blocked_quote_resp.status_code == 409 and "human pricing review" in
           (_blocked_quote_resp.get_json().get("error") or ""),
           _blocked_quote_resp.get_json())

        _unmeasured_spec_id = "33333333-3333-4333-8333-333333333333"
        _unmeasured_jobs = _AS.load_jobs()
        _unmeasured_jobs[_unmeasured_spec_id] = {
            "id": _unmeasured_spec_id, "status": "error", "decision": None,
            "measurement_state": "UNMEASURED",
            "result": {"file": "Dock-Unmeasured.pdf", "measurement_state": "UNMEASURED"},
        }
        _AS.save_jobs(_unmeasured_jobs)
        _unmeasured_spec_resp = _client_up.post(
            f"/spec-override/{_unmeasured_spec_id}", json={
                "slab_type": "dock", "fields": {"depth_mm": 225, "conc_mix": None,
                                                    "mesh": None, "layers": None,
                                                    "bay_sizes": None, "joint_details": None},
            })
        ck("Brief_Spec capture works before a drawing has a measurable area",
           _unmeasured_spec_resp.status_code == 200 and
           _unmeasured_spec_resp.get_json()["repriced"] is False and
           _AS.load_jobs()[_unmeasured_spec_id]["brief_spec"]["fields"]["depth_mm"]["value"] == 225,
           _unmeasured_spec_resp.get_json())

        _zone_job_id = "44444444-4444-4444-8444-444444444444"
        _zone_jobs = _AS.load_jobs()
        _zone_jobs[_zone_job_id] = {
            "id": _zone_job_id, "status": "pending", "decision": None,
            "measurement_state": "MEASURED_VERIFIED", "zone_classification_required": True,
            "zones": [{"zone_key":"unclassified:fdns", "category":"unclassified",
                       "subjects":["FDNS"], "measurement_kind":"unparsed",
                       "area_m2":None, "length_lm":None, "annotation_count":4}],
            "markup_annotations": [{"subject":"FDNS", "type":"Polygon"}],
            "flags": ["assessor: classify zone 'FDNS'"],
            "result": {"file":"External Markup Unit-1.pdf", "area_m2":3185.8,
                       "measurement_state":"MEASURED_VERIFIED",
                       "zone_classification_required":True,
                       "zones":[{"zone_key":"unclassified:fdns", "category":"unclassified",
                                 "subjects":["FDNS"], "measurement_kind":"unparsed",
                                 "area_m2":None, "length_lm":None, "annotation_count":4}],
                       "markup_annotations":[{"subject":"FDNS", "type":"Polygon"}],
                       "flags":["assessor: classify zone 'FDNS'"]},
        }
        _AS.save_jobs(_zone_jobs)
        ck("unclassified marked zone hard-blocks approval",
           _AS._approve_block_reason(_zone_jobs[_zone_job_id]) is not None)
        _classify_resp = _client_up.post(f"/zones/{_zone_job_id}", json={
            "classifications":[{"zone_key":"unclassified:fdns", "category":"other"}],
        })
        _classified_job = _AS.load_jobs()[_zone_job_id]
        ck("assessor can classify out-of-scope FDNS without changing its measurement",
           _classify_resp.status_code == 200 and
           _classified_job["zones"][0]["category"] == "other" and
           not _classified_job["zone_classification_required"] and
           _classified_job["markup_annotations"] == [{"subject":"FDNS", "type":"Polygon"}],
           _classify_resp.get_json())
        _ack_jobs = _AS.load_jobs()
        _ack_jobs[_zone_job_id]["zone_reference_mismatch"] = True
        _ack_jobs[_zone_job_id]["result"]["zone_reference_mismatch"] = True
        _AS.save_jobs(_ack_jobs)
        _ack_resp = _client_up.post(f"/zones/{_zone_job_id}", json={
            "acknowledge_reference_mismatch": True,
        })
        _ack_job = _AS.load_jobs()[_zone_job_id]
        ck("assessor can explicitly acknowledge a BOQ mismatch before approval",
           _ack_resp.status_code == 200 and not _ack_job["zone_reference_mismatch"] and
           _ack_job["result"].get("zone_reference_reviewed_at") and
           _AS._approve_block_reason(_ack_job) is None, _ack_resp.get_json())

        _mixed_zone_id = "55555555-5555-4555-8555-555555555555"
        _mixed_jobs = _AS.load_jobs()
        _mixed_costing = _copy.deepcopy(_demo_result["costing"])
        _mixed_jobs[_mixed_zone_id] = {
            "id": _mixed_zone_id, "status":"pending", "decision":None,
            "measurement_state":"MEASURED_VERIFIED", "costing":_mixed_costing,
            "zones":[{"zone_key":"external_yard", "category":"external_yard", "area_m2":100},
                     {"zone_key":"dock", "category":"dock", "area_m2":20}],
            "brief_specs":{"external_yard":_empty_brief_spec("external_yard"),
                           "dock":_empty_brief_spec("dock")},
            "result":{"file":"External Markup Unit-9.pdf", "area_m2":120,
                      "measurement_state":"MEASURED_VERIFIED", "costing":_mixed_costing,
                      "zones":[{"zone_key":"external_yard", "category":"external_yard", "area_m2":100},
                               {"zone_key":"dock", "category":"dock", "area_m2":20}],
                      "brief_specs":{"external_yard":_empty_brief_spec("external_yard"),
                                     "dock":_empty_brief_spec("dock")}},
        }
        _AS.save_jobs(_mixed_jobs)
        _zone_spec_resp = _client_up.post(f"/spec-override/{_mixed_zone_id}", json={
            "zone_category":"dock", "slab_type":"dock",
            "fields":{"depth_mm":250, "conc_mix":None, "mesh":None, "layers":None,
                      "bay_sizes":None, "joint_details":None},
        })
        _zone_spec_job = _AS.load_jobs()[_mixed_zone_id]
        ck("per-zone slab checklist re-prices only its zone without overwriting aggregate rate",
           _zone_spec_resp.status_code == 200 and _zone_spec_resp.get_json()["repriced"] and
           _zone_spec_job["brief_specs"]["dock"]["fields"]["depth_mm"]["value"] == 250 and
           _zone_spec_job["zone_costings"]["dock"]["area_m2"] == 20 and
           _zone_spec_job["zone_costings"]["dock"]["spec"]["depth_mm"] == 250 and
           _zone_spec_job["costing"] == _mixed_costing,
           _zone_spec_resp.get_json())
        _zone_adjust_resp = _client_up.post(f"/adjust/{_mixed_zone_id}", json={
            "assessed_area_m2":125, "note":"aggregate correction",
        })
        _zone_adjusted_job = _AS.load_jobs()[_mixed_zone_id]
        ck("aggregate adjustment clears stale split and re-blocks zone approval",
           _zone_adjust_resp.status_code == 200 and _zone_adjusted_job["zones"] == [] and
           _zone_adjusted_job["zone_allocation_stale"] and
           _AS._approve_block_reason(_zone_adjusted_job) is not None)
        _empty_stale_confirm = _client_up.post(
            f"/confirm-measurement/{_mixed_zone_id}", json={"confirm_scale_extent":True})
        ck("confirmation cannot resurrect zones erased by a true aggregate replacement",
           _empty_stale_confirm.status_code == 409 and
           _AS.load_jobs()[_mixed_zone_id]["zone_allocation_stale"],
           _empty_stale_confirm.get_json())

        _portal_html_up = (Path(_orig_server_file_up).parent / "assessor_portal.html").read_text()
        ck("portal file input allows multiple PDFs and ZIPs",
           'accept=".pdf,.zip" multiple' in _portal_html_up)
        ck("portal submits every selected file under the backward-compatible pdf field",
           "files.forEach(file => fd.append('pdf', file))" in _portal_html_up)
        ck("portal groups repeated project refs under collapsible project headers",
           "projectCounts.get(ref)" in _portal_html_up and
           'class="project-group-header' in _portal_html_up and
           "toggleProjectGroup(this)" in _portal_html_up)
        ck("portal exposes editable xlsx quotation download",
           'id="linkXlsx"' in _portal_html_up and
           "quotation/${job.id}.xlsx" in _portal_html_up)
        ck("portal exposes marked-up PDF independently of quotation pricing state",
           all(marker in _portal_html_up for marker in (
               'id="markedPdfLinks"', 'id="linkMarkedPdf"',
               "marked-pdf/${job.id}.pdf", "Permanent Bluebeam-ready markup",
               "recoverable Fortel geometry embedded")))
        ck("portal exposes renameable +Area geometry separately from main regions",
           all(marker in _portal_html_up for marker in (
               'id="btnNewArea"', "startNewAreaElement",
               "Separate area name", "area_elements: namedAreaEntries.map",
               # the summary line now names what it is summing: the primary surface on a
               # single-zone drawing, every measured zone on a multi-surface one
               "'main area'", "'all measured zones'")))
        ck("portal accumulates folder selections and can target an existing project anchor",
           all(marker in _portal_html_up for marker in (
               "selectedUploadFiles.push(...Array.from(fileList || []))",
               "document.getElementById('upFile').value = ''",
               'id="upFileList"', 'id="btnAddDrawings"',
               "beginAddDrawingsById", "existing_project_job_id")))
        ck("portal switches to a scalable PDF visual layer on zoom without changing snapScale",
           all(marker in _portal_html_up for marker in (
               "ensureVectorSurface", "snapshot-vector/${encodeURIComponent(requestedJobId)}.svg",
               "if (!vectorSurfaceReady && img)",
               "if (vectorSurface) vectorSurface.style.transform = transform",
               "res.scale_k ? res.scale_k / snapScale : null")))
        ck("portal exposes exact Brief_Spec fields without silent fallback form values",
           all(label in _portal_html_up for label in (
               "External/Service Yard Slabs", "Dock Slabs", "Ground Floor Slabs(Core Areas)",
               "Upper Floors", "Bay sizes if joint layout available", "Nr of mesh layers")) and
           "${spec.depth_mm||190}" not in _portal_html_up and
           "${esc(spec.mesh||'A252')}" not in _portal_html_up and
           "ASSUMED / no details provided" in _portal_html_up and
           "projectPricingBlocked" in _portal_html_up)
        ck("portal displays drawing-file and page citations only when extraction evidence exists",
           all(marker in _portal_html_up for marker in (
               "function specSourceCitation", "evidence.file", "evidence.page",
               "Source: ${esc(citation)}")))
        ck("portal permits an approved spec correction and labels its immutable quote revision",
           all(marker in _portal_html_up for marker in (
               "Correct approved specification", "Save as new quotation revision",
               "post_approval_correction", "CORRECTED AFTER APPROVAL", "quotation_revision")))
        # Both are still shown; the labels are now plain English. Inderjit, 4 Sep call: "I'm not
        # sure what scale K is" — the panel called it "Current k (m/px)".
        ck("portal presents the measurement scale as both metres-per-pixel and conventional 1:N",
           all(marker in _portal_html_up for marker in (
               'id="scaleDisplay"', 'id="scaleRatioDisplay"', "Metres per pixel", "Drawing scale",
               "(mpp * snapScale) * 72 / 0.0254")))
        ck("...and it no longer calls the scale 'k' at the assessor",
           "Current k (m/px)" not in _portal_html_up)
        # Inderjit traced a footpath with + Region and it landed in the main slab total.
        ck("the two trace buttons say what they do to the total",
           "its area is ADDED to the main measured total" in _portal_html_up
           and "kept out of the main total" in _portal_html_up
           and "use + Area instead" in _portal_html_up)
        ck("portal renders and captures per-zone quantities/classifications/specs",
           all(marker in _portal_html_up for marker in (
               "Measured zones", "ZONE REVIEW REQUIRED", "classifyZone(",
               "acknowledgeZoneReferenceMismatch", "zone_category", "effectiveBriefSpecs")))
        ck("portal shows correct banner for yard-region review vs zone classification",
           all(marker in _portal_html_up for marker in (
               "yardRegionReview", "YARD REGION REVIEW REQUIRED",
               "Multiple same-tint Yard regions detected",
               "yard_region_review_required")))
        ck("portal exposes assisted Office candidates without auto-submitting them",
           all(marker in _portal_html_up for marker in (
               "ASSISTED TRACE CANDIDATES", "candidate_polygons", "loadTraceCandidate(",
               "Add to trace", "btnNewRegion", "traceRegions", 'id="traceScope"',
               "ground_floor_core", "main_upper_floor", "plant_deck",
               "pod_first_floor", "region_scopes: regionEntries.map")) and
           "function loadTraceCandidate" in _portal_html_up and
           "const proposed = candidatePolygons.find" in _portal_html_up and
           "loadTraceCandidate(proposed.candidate_id)" in _portal_html_up and
           "regions: regionPayload" in _portal_html_up and
           "region_categories: regionEntries.map" in _portal_html_up)
        ck("portal offers non-destructive existing scale+extent confirmation",
           all(marker in _portal_html_up for marker in (
               "btnConfirmExisting", "confirmExistingMeasurement",
               "/confirm-measurement/", "Confirm scale + extent")))
        ck("portal explains candidate confidence and keeps unresolved levels visible",
           all(marker in _portal_html_up for marker in (
               "confidence_reasons", "confidence_score", "outline_status",
               "Trace manually", "candidate.regions")))
        ck("portal labels channel proposals as assumptions and provides review controls",
           all(marker in _portal_html_up for marker in (
               "ASSUMED CHANNEL PROPOSALS - NOT MEASURED OR PRICED",
               "channel_proposals", "reviewChannelProposal(",
               "/channel-proposals/", "Accept / save edit", "Remove",
               "Dock-level retaining-wall/loading face", "Full Yard width (no Dock level)")))
        ck("portal gives Transition candidates the full accept/edit/remove lifecycle",
           all(marker in _portal_html_up for marker in (
               "ASSUMED TRANSITION CANDIDATES - NOT PRICED UNTIL REVIEWED",
               "transition_candidate_decisions", "editTransitionLength(",
               "reviewTransitionCandidate(", "/transition-candidates/",
               "Accept / save edit", "Remove", "blank assessor rate")))
        ck("portal visibly carries exclusion checks and construction-joint classification",
           all(marker in _portal_html_up for marker in (
               "SLAB EXCLUSIONS", "CHECK EXCLUSION", "EXCLUDED ·",
               "Construction joint (CJ)", "construction_joint")))
        ck("portal supports axis-locked endpoint drag plus numeric channel geometry edits",
           all(marker in _portal_html_up for marker in (
               "channelDrag", "function editChannelLength", "syncChannelLengthInput",
               "line remains straight/non-diagonal", "polyline_pts:polylinePts")))
        ck("portal separately supports assessor-drawn bending channel polylines",
           all(marker in _portal_html_up for marker in (
               "click each bend", "active.points.push(p)", "active.points.length",
               "ch.points.length >= 2", "all bends reach the server/quotation")))
        ck("portal surfaces every retained Yard region and posts explicit keep/exclude decisions",
           all(marker in _portal_html_up for marker in (
               "effectiveYardRegions", "yard-region-toggle", "saveYardRegionReview",
               "/yard-regions/", "Save kept/excluded regions", "bbox_pdf_pts")))
        ck("portal gives assumed/proposed/provisional values one unmissable provenance style",
           all(marker in _portal_html_up for marker in (
               'class="assumption-badge" data-provenance="assumed"',
               "assumption-item", "assumption-legend", "assumption-basis",
               "assumptionBadge('PROPOSED'", "assumptionBadge('PROVISIONAL'",
               "assumptionBadge('ASSUMED'", "assumptionBadge('ESTIMATED'")))
        _zone_assumption_fn = _portal_html_up.split("function zoneAssumption", 1)[1].split(
            "function provenanceFlag", 1)[0]
        ck("measured zone quantities never gain an assumption badge without explicit provenance",
           all(marker in _zone_assumption_fn for marker in (
               "zone.assumed", "zone.proposed", "zone.provisional", "zone.estimate")) and
           all(marker not in _zone_assumption_fn for marker in (
               "measurement_state", "measurement_kind", "area_m2", "length_lm")))
        _candidate_fn = _portal_html_up.split("function loadTraceCandidate", 1)[1].split(
            "function calcArea", 1)[0]
        ck("one-click candidate load is non-mutating until Submit Adjustment",
           "fetch(" not in _candidate_fn and "poly = regions[0]" in _candidate_fn)
    finally:
        _AS._TAKEOFF_DISPATCHER = _orig_dispatcher_up
        _AS.__file__ = _orig_server_file_up
        _AS.JOBS_FILE = _orig_jobs_file_up
        _AS.JOBS_ARCHIVE_FILE = _orig_jobs_archive_file_up
        _AS.BACKUP_DIR = _orig_backup_dir_up
        _AS.DRAWINGS_DIR = _orig_drawings_dir_up
        _AS.QUOTATIONS_DIR = _orig_quotations_dir_up

    # approve hard-block mirrors the >£200k escalation guard mechanism (fb5b92b)
    ck("UNMEASURED job blocks approve",
       _AS._approve_block_reason({"measurement_state": "UNMEASURED", "scale_confirmed": False}) is not None)
    ck("MEASURED_UNVERIFIED job blocks approve",
       _AS._approve_block_reason({"measurement_state": "MEASURED_UNVERIFIED", "scale_confirmed": False}) is not None)
    ck("MEASURED_VERIFIED job does not block approve",
       _AS._approve_block_reason({"measurement_state": "MEASURED_VERIFIED", "scale_confirmed": False}) is None)
    ck("assessor-confirmed UNMEASURED job no longer blocks approve",
       _AS._approve_block_reason({"measurement_state": "UNMEASURED", "scale_confirmed": True}) is None)
    ck("REJECTED job blocks approve", _AS._approve_block_reason({"measurement_state": "REJECTED"}) is not None)

    # A restart must RESUME what it only delayed. On 4 Sep 2026 an assessor uploaded 92 files;
    # with two workers at ~90s a drawing that batch drains for over an hour, and a deploy landed
    # ten minutes before his review call. Every still-queued job was swept to UNMEASURED with
    # "PIPELINE INTERRUPTED" — his exact words were "all that runs remain unmeasured... I haven't
    # got any response at all". A queued job has lost nothing: the PDF is on the volume.
    _resume_pdf = Path(_tmpdir) / "resume_me.pdf"
    _resume_doc = _fitz_fast_refusal.open(); _resume_doc.new_page(width=300, height=300)
    _resume_doc.save(str(_resume_pdf)); _resume_doc.close()
    _submitted = []
    _real_submit = _AS._TAKEOFF_DISPATCHER.submit
    _AS._TAKEOFF_DISPATCHER.submit = lambda *args: _submitted.append(args)
    try:
        _AS.save_jobs({
            "queued-job": {"id": "queued-job", "status": "processing", "takeoff_phase": "queued",
                           "pdf_path": str(_resume_pdf), "project_name": "P", "project_ref": "R",
                           "flags": [], "result": {}},
            "measuring-job": {"id": "measuring-job", "status": "processing",
                              "takeoff_phase": "measuring", "pdf_path": str(_resume_pdf),
                              "flags": [], "result": {}},
            "queued-but-file-gone": {"id": "queued-but-file-gone", "status": "processing",
                                     "takeoff_phase": "queued", "pdf_path": "/nonexistent.pdf",
                                     "flags": [], "result": {}},
        })
        _AS._sweep_stranded_processing_jobs()
        _swept = _AS.load_jobs()
    finally:
        _AS._TAKEOFF_DISPATCHER.submit = _real_submit
    ck("a job still QUEUED at a restart is put back on the queue, not marked unmeasurable",
       [args[0] for args in _submitted] == ["queued-job"]
       and _swept["queued-job"]["status"] == "processing"
       and _swept["queued-job"]["measurement_state"] != "UNMEASURED"
       if "measurement_state" in _swept["queued-job"] else True,
       {"submitted": [args[0] for args in _submitted],
        "state": _swept["queued-job"].get("measurement_state")})
    ck("...and it says so, rather than looking like nothing happened",
       any("QUEUED WORK RESUMED" in flag for flag in _swept["queued-job"]["flags"]),
       _swept["queued-job"]["flags"])
    ck("a job that was MID-MEASUREMENT still routes to the assessor — that work is really gone",
       _swept["measuring-job"]["measurement_state"] == "UNMEASURED"
       and any("PIPELINE INTERRUPTED" in flag for flag in _swept["measuring-job"]["flags"]),
       _swept["measuring-job"].get("flags"))
    ck("a queued job whose file has vanished is NOT silently re-queued forever",
       _swept["queued-but-file-gone"]["measurement_state"] == "UNMEASURED"
       and "queued-but-file-gone" not in [args[0] for args in _submitted],
       _swept["queued-but-file-gone"].get("measurement_state"))
    _AS.save_jobs({})

    # /status must say what a deploy decision needs: not just how many jobs exist, but how many
    # are IN FLIGHT and would be interrupted by the restart a push causes.
    _AS.save_jobs({"a": {"status": "processing", "takeoff_phase": "queued"},
                   "b": {"status": "processing", "takeoff_phase": "measuring"},
                   "c": {"status": "done"}})
    _inflight = _client_up.get("/status").get_json()
    ck("/status reports what a deploy would interrupt: in-flight and queued counts",
       _inflight["job_count"] == 3 and _inflight["processing_count"] == 2
       and _inflight["queued_count"] == 1, _inflight)
    _AS.save_jobs({})

    # The >£200k approve lock is CONFIG, not code. It was written for an automatic pipeline;
    # every upload is done by hand today, so it blocked the assessor who was already the human
    # review it demanded (Aryan, 4 Sep). Default: no block, and the portal shows a notice
    # instead. One env var restores the hard block when the pipeline goes automatic.
    _esc_env = os.environ.get("ESCALATION_LOCK_GBP")
    try:
        os.environ.pop("ESCALATION_LOCK_GBP", None)
        ck("escalation lock is OFF by default, so a manual assessor is never blocked by it",
           _AS._escalation_lock_gbp() is None)
        os.environ["ESCALATION_LOCK_GBP"] = "200000"
        ck("...and one env var brings the £200k lock back, with no code change",
           _AS._escalation_lock_gbp() == 200000.0, _AS._escalation_lock_gbp())
        os.environ["ESCALATION_LOCK_GBP"] = "not-a-number"
        ck("...a malformed value disables the lock rather than crashing the server",
           _AS._escalation_lock_gbp() is None)
        os.environ["ESCALATION_LOCK_GBP"] = "0"
        ck("...and zero means off, not block-everything",
           _AS._escalation_lock_gbp() is None)
    finally:
        os.environ.pop("ESCALATION_LOCK_GBP", None)
        if _esc_env is not None:
            os.environ["ESCALATION_LOCK_GBP"] = _esc_env
    _esc_status = _client_up.get("/status").get_json()
    ck("the portal is told the lock value by /status, so client and server cannot disagree",
       "escalation_lock_gbp" in _esc_status and _esc_status["escalation_lock_gbp"] is None,
       _esc_status)
    _portal_html_esc = _client_up.get("/portal").data.decode("utf-8", "replace")
    ck("the portal blocks approve ONLY when a lock is configured, and never on a hard-coded 200000",
       "escalationLockGbp !== null && c && c.assumed" in _portal_html_esc
       and "c.total_gbp > 200000)" not in _portal_html_esc)
    # Aryan, 4 Sep: "even after confirming the second zone the total measurement area remains
    # calculated for the first primary region only". The pipeline's top-level area_m2 is the
    # PRIMARY surface on purpose (so a road/dock can never inherit the Yard rate), so the fix is
    # in the display: sum the measured zones and show the split.
    ck("the portal headline sums every measured zone instead of showing the primary one",
       "function zoneAreaBreakdown(job)" in _portal_html_esc
       and "const aiArea = zoneBreakdown ? zoneBreakdown.total" in _portal_html_esc)
    ck("...and it names the split, so a two-surface drawing is legible at a glance",
       "AI measurement · ${zoneBreakdown.parts.map" in _portal_html_esc
       and "all measured zones" in _portal_html_esc)
    ck("...while a single-zone drawing is untouched (breakdown needs 2+ measured zones)",
       "if (zones.length < 2) return null;" in _portal_html_esc)
    ck("...and a large assumed-spec quotation still shows a visible notice with the number",
       "Large quotation on an assumed build-up" in _portal_html_esc
       and "confirm the commercial basis" in _portal_html_esc)

    shutil.rmtree(_tmpdir, ignore_errors=True)
except ImportError as _e:
    print(f"  [SKIP] approval_server upload/approve tests — missing dependency: {_e}")

print("approval_server: bounded 26-PDF queue excludes wait time from watchdog budget")
try:
    import approval_server as _AS_queue
    import fitz as _fitz_queue, io as _io_queue, tempfile as _tempfile_queue
    import sys as _sys_queue, time as _time_queue, threading as _threading_queue

    _queue_tmp = Path(_tempfile_queue.mkdtemp(prefix="ci_takeoff_queue_"))
    _queue_originals = {
        "jobs": _AS_queue.JOBS_FILE,
        "archive": _AS_queue.JOBS_ARCHIVE_FILE,
        "backup": _AS_queue.BACKUP_DIR,
        "drawings": _AS_queue.DRAWINGS_DIR,
        "dispatcher": _AS_queue._TAKEOFF_DISPATCHER,
        "timeout": _AS_queue.TAKEOFF_TIMEOUT_S,
    }
    _real_pipeline_queue = _sys_queue.modules.get("takeoff_pipeline")
    _active_queue = 0
    _max_active_queue = 0
    _active_lock_queue = _threading_queue.Lock()

    ck("takeoff worker default is CPU-sized and deliberately capped at two",
       _AS_queue._takeoff_worker_count(raw_value="", cpu_count=8) == 2 and
       _AS_queue._takeoff_worker_count(raw_value="", cpu_count=1) == 1)
    ck("TAKEOFF_WORKERS explicitly overrides the CPU-sized default",
       _AS_queue._takeoff_worker_count(raw_value="3", cpu_count=1) == 3)

    class _QueuePipeline:
        @staticmethod
        def takeoff(pdf_path, project_name=None, project_ref=None,
                    client_rates_path=None, approval_job_id=None):
            global _active_queue, _max_active_queue
            with _active_lock_queue:
                _active_queue += 1
                _max_active_queue = max(_max_active_queue, _active_queue)
            try:
                _time_queue.sleep(0.03)
                return {
                    "file": Path(pdf_path).name,
                    "project_name": project_name,
                    "project_ref": project_ref,
                    "area_m2": 250.0,
                    "measurement_state": "MEASURED_VERIFIED",
                    "needs_assessor": False,
                    "scale_verified": True,
                    "flags": ["queue regression stub completed"],
                }
            finally:
                with _active_lock_queue:
                    _active_queue -= 1

    try:
        _AS_queue.JOBS_FILE = _queue_tmp / "approval_jobs.json"
        _AS_queue.JOBS_ARCHIVE_FILE = _queue_tmp / "approval_jobs_archive.json"
        _AS_queue.BACKUP_DIR = _queue_tmp / "backups"
        _AS_queue.DRAWINGS_DIR = _queue_tmp / "drawings"
        _AS_queue.TAKEOFF_TIMEOUT_S = 0.15
        _AS_queue._TAKEOFF_DISPATCHER = _AS_queue._TakeoffDispatcher(2)
        _sys_queue.modules["takeoff_pipeline"] = _QueuePipeline
        _AS_queue.save_jobs({})

        _queue_doc = _fitz_queue.open()
        _queue_doc.new_page(width=300, height=200)
        _queue_bytes = _queue_doc.tobytes()
        _queue_doc.close()
        _queue_files = [
            (_io_queue.BytesIO(_queue_bytes), f"Tender_Drawing_{index:02d}.pdf")
            for index in range(1, 27)
        ]
        _queue_started = _time_queue.monotonic()
        _queue_response = _AS_queue.app.test_client().post("/upload", data={
            "project_ref": "QUEUE-26",
            "project_name": "26 Drawing Tender Pack",
            "pdf": _queue_files,
        }, content_type="multipart/form-data")
        _AS_queue._TAKEOFF_DISPATCHER.wait_for_idle()
        _queue_elapsed = _time_queue.monotonic() - _queue_started
        _queue_jobs = _AS_queue.load_jobs()
        _queue_states = [job.get("measurement_state") for job in _queue_jobs.values()]
        _queue_timeout_flags = [
            flag for job in _queue_jobs.values() for flag in (job.get("flags") or [])
            if "PIPELINE TIMEOUT" in flag
        ]

        ck("26-PDF upload returns one queued job per drawing",
           _queue_response.status_code == 202 and len(_queue_jobs) == 26 and
           len(_queue_response.get_json().get("job_ids", [])) == 26,
           {"http": _queue_response.status_code,
            "jobs": len(_queue_jobs), "body": _queue_response.get_json()})
        ck("bounded takeoff queue never exceeds configured worker concurrency",
           _max_active_queue == 2, _max_active_queue)
        ck("every queued drawing reaches a legitimate four-state outcome",
           all(state in {"MEASURED_VERIFIED", "MEASURED_UNVERIFIED", "UNMEASURED", "REJECTED"}
               for state in _queue_states) and
           all(job.get("takeoff_phase") == "completed" for job in _queue_jobs.values()),
           {"states": _queue_states,
            "phases": [job.get("takeoff_phase") for job in _queue_jobs.values()]})
        ck("queue wait longer than one watchdog budget causes zero watchdog failures",
           _queue_elapsed > _AS_queue.TAKEOFF_TIMEOUT_S and not _queue_timeout_flags and
           not any(job.get("status") == "error" for job in _queue_jobs.values()),
           {"elapsed_s": round(_queue_elapsed, 3),
            "watchdog_s": _AS_queue.TAKEOFF_TIMEOUT_S,
            "timeout_flags": _queue_timeout_flags})
        ck("26-file bounded batch finishes promptly",
           _queue_elapsed < 5.0, round(_queue_elapsed, 3))
        print(f"  [EVIDENCE] 26 files, workers=2, watchdog=0.15s, "
              f"batch={_queue_elapsed:.3f}s, max_active={_max_active_queue}, "
              f"watchdog_kills={len(_queue_timeout_flags)}")
    finally:
        _AS_queue.JOBS_FILE = _queue_originals["jobs"]
        _AS_queue.JOBS_ARCHIVE_FILE = _queue_originals["archive"]
        _AS_queue.BACKUP_DIR = _queue_originals["backup"]
        _AS_queue.DRAWINGS_DIR = _queue_originals["drawings"]
        _AS_queue._TAKEOFF_DISPATCHER = _queue_originals["dispatcher"]
        _AS_queue.TAKEOFF_TIMEOUT_S = _queue_originals["timeout"]
        if _real_pipeline_queue is None:
            _sys_queue.modules.pop("takeoff_pipeline", None)
        else:
            _sys_queue.modules["takeoff_pipeline"] = _real_pipeline_queue
        shutil.rmtree(_queue_tmp, ignore_errors=True)
except ImportError as _e:
    print(f"  [SKIP] bounded takeoff queue tests — missing dependency: {_e}")

print("approval_server: /snapshot status codes for all four measurement states "
      "(Aryan field report — 'session which renders screenshots is not working properly')")
try:
    import approval_server as _AS2
    import fitz as _fitz4, uuid as _uuid2, tempfile as _tempfile2

    _client = _AS2.app.test_client()
    _tmpdir2 = Path(_tempfile2.mkdtemp(prefix="ci_snapshot_"))

    # Save/restore the real jobs file around this block — snapshot() reads via load_jobs()
    # which is a real file read, not mockable without a live Flask app context.
    _jobs_backup = _AS2.JOBS_FILE.read_text() if _AS2.JOBS_FILE.exists() else None

    def _mk_pdf(path, w=600, h=400, n_pages=1):
        d = _fitz4.open()
        for _ in range(n_pages):
            d.new_page(width=w, height=h)
        d.save(str(path))
        return path

    try:
        _jobs = _AS2.load_jobs()

        # 1. REJECTED job (no pdf_path at all) -> 404, not 500
        _jid_rej = str(_uuid2.uuid4())
        _jobs[_jid_rej] = {"id": _jid_rej, "status": "rejected", "measurement_state": "REJECTED",
                           "pdf_path": None, "result": {"measurement_state": "REJECTED"}}

        # 2. UNMEASURED job with a real PDF on disk -> 200 (assessor still needs to see it to trace)
        _pdf_unm = _mk_pdf(_tmpdir2 / "unmeasured.pdf")
        _jid_unm = str(_uuid2.uuid4())
        _jobs[_jid_unm] = {"id": _jid_unm, "status": "error", "measurement_state": "UNMEASURED",
                           "pdf_path": str(_pdf_unm),
                           "result": {"pdf_path": str(_pdf_unm), "page": 0, "measurement_state": "UNMEASURED"}}

        # 3. UNMEASURED job whose PDF is missing from disk (temp dir cleaned up) -> 404, not 500
        _jid_gone = str(_uuid2.uuid4())
        _jobs[_jid_gone] = {"id": _jid_gone, "status": "error", "measurement_state": "UNMEASURED",
                            "pdf_path": str(_tmpdir2 / "does_not_exist.pdf"),
                            "result": {"pdf_path": str(_tmpdir2 / "does_not_exist.pdf"),
                                      "measurement_state": "UNMEASURED"}}

        # 4. MEASURED_VERIFIED multi-page job whose result["page"] != 0 -> snapshot must render
        # THAT page (this was the root cause of "AI polygon not shown": /snapshot always
        # rendered page 0 regardless of which page the pipeline actually measured).
        _pdf_multi = _mk_pdf(_tmpdir2 / "multi.pdf", n_pages=3)
        _jid_page = str(_uuid2.uuid4())
        _jobs[_jid_page] = {"id": _jid_page, "status": "pending", "measurement_state": "MEASURED_VERIFIED",
                            "pdf_path": str(_pdf_multi),
                            "result": {"pdf_path": str(_pdf_multi), "page": 2,
                                      "polygon_pts": [[10, 10], [100, 10], [100, 100], [10, 100]],
                                      "measurement_state": "MEASURED_VERIFIED"}}

        # 5. Out-of-range page index (stale data) -> must fall back to page 0, never 500
        _jid_badpage = str(_uuid2.uuid4())
        _jobs[_jid_badpage] = {"id": _jid_badpage, "status": "pending", "measurement_state": "MEASURED_VERIFIED",
                               "pdf_path": str(_pdf_multi),
                               "result": {"pdf_path": str(_pdf_multi), "page": 99,
                                         "measurement_state": "MEASURED_VERIFIED"}}

        _AS2.save_jobs(_jobs)

        _r_rej = _client.get(f"/snapshot/{_jid_rej}")
        ck("REJECTED job snapshot -> 404 (not 500)", _r_rej.status_code == 404, _r_rej.status_code)

        _r_unm = _client.get(f"/snapshot/{_jid_unm}")
        ck("UNMEASURED job with PDF on disk -> 200 (assessor can still trace)",
           _r_unm.status_code == 200, _r_unm.status_code)

        _r_gone = _client.get(f"/snapshot/{_jid_gone}")
        ck("UNMEASURED job with missing PDF -> 404 (not 500)", _r_gone.status_code == 404, _r_gone.status_code)

        _r_page = _client.get(f"/snapshot/{_jid_page}")
        ck("multi-page job snapshot -> 200", _r_page.status_code == 200, _r_page.status_code)
        # Verify it actually rendered page 2's dimensions, not page 0's (both pages here are
        # the same size so we check indirectly: render page 2 directly and diff against the
        # response bytes' pixel dimensions via the PNG header — same width guaranteed by
        # construction, so the meaningful assertion is the X-Snapshot-Scale header matches
        # snapshot_scale() computed for page 2 specifically.
        from approval_email import snapshot_scale as _snap_scale_fn
        _expected_scale = f"{_snap_scale_fn(str(_pdf_multi), page=2):.6f}"
        ck("multi-page snapshot X-Snapshot-Scale computed for the MEASURED page (not page 0)",
           _r_page.headers.get("X-Snapshot-Scale") == _expected_scale,
           (_r_page.headers.get("X-Snapshot-Scale"), _expected_scale))

        _r_badpage = _client.get(f"/snapshot/{_jid_badpage}")
        ck("out-of-range page index falls back to page 0 (not 500)",
           _r_badpage.status_code == 200, _r_badpage.status_code)

        _r_404job = _client.get(f"/snapshot/{_uuid2.uuid4()}")
        ck("nonexistent job -> 404", _r_404job.status_code == 404, _r_404job.status_code)

    finally:
        if _jobs_backup is not None:
            _AS2.JOBS_FILE.write_text(_jobs_backup)
        shutil.rmtree(_tmpdir2, ignore_errors=True)
except ImportError as _e:
    print(f"  [SKIP] approval_server snapshot tests — missing dependency: {_e}")

print("approval_server: watchdog-vs-completion race (Aryan field report — 'server is unstable')")
try:
    import approval_server as _AS3
    import sys as _sys3, time as _time3, uuid as _uuid3
    from unittest import mock as _mock3

    _jobs_backup3 = _AS3.JOBS_FILE.read_text() if _AS3.JOBS_FILE.exists() else None
    try:
        _jid_wd = str(_uuid3.uuid4())
        _jobs3 = _AS3.load_jobs()
        _jobs3[_jid_wd] = {"id": _jid_wd, "status": "processing", "flags": []}
        _AS3.save_jobs(_jobs3)

        # _mark_job_unmeasured with watchdog_fired=True sets the sentinel used to detect the race
        _AS3._mark_job_unmeasured(_jid_wd, "PIPELINE TIMEOUT: took too long", watchdog_fired=True)
        _j_after_wd = _AS3.load_jobs()[_jid_wd]
        ck("watchdog fire sets _watchdog_fired sentinel", _j_after_wd.get("_watchdog_fired") is True)
        ck("watchdog fire flips job to UNMEASURED", _j_after_wd.get("measurement_state") == "UNMEASURED")
        ck("watchdog fire records a PIPELINE TIMEOUT flag",
           any("PIPELINE TIMEOUT" in f for f in _j_after_wd.get("flags", [])))

        # Now simulate the pipeline finishing LATE (after the watchdog already fired) by
        # driving the real _run_takeoff() with a stubbed takeoff_pipeline module whose
        # takeoff() sleeps past a 1s watchdog timeout — exercises the actual production
        # code path, not a re-implementation of its logic.
        _orig_timeout = _AS3.TAKEOFF_TIMEOUT_S
        _AS3.TAKEOFF_TIMEOUT_S = 1
        _jid_wd2 = str(_uuid3.uuid4())
        _jobs3 = _AS3.load_jobs()
        _jobs3[_jid_wd2] = {"id": _jid_wd2, "status": "processing", "flags": []}
        _AS3.save_jobs(_jobs3)

        _fake_pipeline = _mock3.MagicMock()
        def _slow_takeoff(pdf_path, project_name=None, project_ref=None):
            _time3.sleep(2.2)
            return {"measurement_state": "MEASURED_VERIFIED", "area_m2": 3159.0,
                    "flags": ["completed ok"], "project_name": project_name, "project_ref": project_ref,
                    "candidate_polygons":[{"candidate_id":"office-p0-level-01-1",
                                           "polygon_pts":[[0,0],[1,0],[1,1]]}]}
        _fake_pipeline.takeoff = _slow_takeoff
        _real_module = _sys3.modules.get("takeoff_pipeline")
        _sys3.modules["takeoff_pipeline"] = _fake_pipeline
        try:
            _AS3._run_takeoff(_jid_wd2, "drawings/_int_d77.pdf", "QA WD race", "QA-PORTAL-CI-WDRACE")
        finally:
            if _real_module is not None:
                _sys3.modules["takeoff_pipeline"] = _real_module
            else:
                _sys3.modules.pop("takeoff_pipeline", None)
            _AS3.TAKEOFF_TIMEOUT_S = _orig_timeout

        _j_final = _AS3.load_jobs()[_jid_wd2]
        ck("late pipeline completion overwrites watchdog UNMEASURED with the real result",
           _j_final.get("measurement_state") == "MEASURED_VERIFIED", _j_final.get("measurement_state"))
        ck("stale 'PIPELINE TIMEOUT' flag stripped once the pipeline actually completes",
           not any("PIPELINE TIMEOUT" in f for f in _j_final.get("flags", [])), _j_final.get("flags"))
        ck("_watchdog_fired sentinel cleared after the race resolves",
           "_watchdog_fired" not in _j_final)
        ck("background takeoff mirrors assisted candidates at job and result level",
           _j_final.get("candidate_polygons") ==
           _j_final.get("result", {}).get("candidate_polygons") and
           _j_final.get("candidate_polygons", [])[0]["candidate_id"] ==
           "office-p0-level-01-1")

        _jobs3 = _AS3.load_jobs()
        _jobs3.pop(_jid_wd, None); _jobs3.pop(_jid_wd2, None)
        _AS3.save_jobs(_jobs3)
    finally:
        if _jobs_backup3 is not None:
            _AS3.JOBS_FILE.write_text(_jobs_backup3)
except ImportError as _e:
    print(f"  [SKIP] approval_server watchdog-race tests — missing dependency: {_e}")

print("approval_server: approval_jobs.json concurrent read/write does not raise "
      "(Aryan field report — 'the server is unstable')")
try:
    import approval_server as _AS4
    import threading as _threading4, tempfile as _tempfile4

    _tmp_jobs_file = Path(_tempfile4.mkdtemp(prefix="ci_atomic_")) / "jobs.json"
    _orig_jobs_file = _AS4.JOBS_FILE
    _AS4.JOBS_FILE = _tmp_jobs_file
    try:
        _big = {str(_i): {"x": "y" * 500} for _i in range(500)}
        _AS4.save_jobs(_big)

        _errors4 = []
        def _reader4():
            for _ in range(150):
                try:
                    _d = _AS4.load_jobs()
                    if not isinstance(_d, dict):
                        _errors4.append("load_jobs did not return a dict")
                except Exception as _e:
                    _errors4.append(str(_e))

        def _writer4():
            for _ in range(150):
                _AS4.save_jobs(_big)

        _t1 = _threading4.Thread(target=_reader4)
        _t2 = _threading4.Thread(target=_writer4)
        _t1.start(); _t2.start(); _t1.join(); _t2.join()

        ck("concurrent load_jobs()/save_jobs() never raises or returns a torn read",
           len(_errors4) == 0, _errors4[:3])
        ck("no leftover .tmp files after concurrent saves",
           list(_tmp_jobs_file.parent.glob("*.tmp*")) == [])
    finally:
        _AS4.JOBS_FILE = _orig_jobs_file
        shutil.rmtree(_tmp_jobs_file.parent, ignore_errors=True)
except ImportError as _e:
    print(f"  [SKIP] approval_server atomic-write tests — missing dependency: {_e}")

