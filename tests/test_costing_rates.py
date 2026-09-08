#!/usr/bin/env python3
"""price_with_defaults, the client rate-override layer, supplier fields.

Sections in this module (printed in this order):
  - pipeline price_with_defaults
  - client-editable rate override layer — defaults untouched, versioned/audited, quotation-stamped
  - spec extractor — supplier fields
  - supplier inquiry generator

Imported by ci_tests.py, which owns the running total. Do not reorder.
"""
import shutil
from pathlib import Path
from tests import ck
from io import BytesIO as _BytesIO
from openpyxl import load_workbook as _load_workbook
from quotation import generate_quotation, quotation_html, quotation_json, quotation_text, quotation_xlsx

print("pipeline price_with_defaults")
import contextlib, io as _io
with contextlib.redirect_stdout(_io.StringIO()):
    from takeoff_pipeline import price_with_defaults, _needs_approval
_c = price_with_defaults(26080)
ck("26,080 m² at defaults -> £1,175,425.60", _c["total_gbp"] == 1175425.60)
ck("price_with_defaults assumed=True (no spec)", _c["assumed"] is True)
_c2 = price_with_defaults(3172, {"depth_mm": 200, "mesh": "A393", "layers": 1,
                                  "conc_mix": "C32/40", "conc_rate": 128})
ck("all four client construction fields supplied -> assumed=False", _c2["assumed"] is False)
_c3 = price_with_defaults(3172, {"depth_mm": 200})
ck("partial client construction spec stays assumed/provisional", _c3["assumed"] is True)
ck("approval trigger on assessor flag",
   _needs_approval({"type":"UNMARKED vector","confidence":"medium",
                    "flags":["assessor: confirm extent + scale"]}))
ck("no approval trigger on clean marked",
   not _needs_approval({"type":"MARKED vector","confidence":"high","flags":[]}))

print("client-editable rate override layer — defaults untouched, versioned/audited, quotation-stamped")
try:
    import hashlib as _hashlib_rates
    import tempfile as _tempfile_rates
    import json as _json_rates
    from client_rates import (apply_client_rates as _apply_client_rates_test,
                              load_rate_store as _load_rate_store_test,
                              save_client_rates as _save_client_rates_test)
    from defaults import DEFAULT_SPEC as _DEFAULT_SPEC_RATES
    from takeoff_pipeline import MANHOLE_EO_RATE as _MANHOLE_RATE_DEFAULT

    _rates_tmpdir = Path(_tempfile_rates.mkdtemp(prefix="ci_client_rates_"))
    _rates_path = _rates_tmpdir / "client_rates.json"
    _rate_defaults = {
        key: _DEFAULT_SPEC_RATES[key]
        for key in ("conc_rate", "steel_rate_t", "margin", "labour", "dpm", "curing",
                    "trim", "conc_wastage", "steel_wastage", "lap_acc")
    }
    _rate_defaults["manhole_eo_rate"] = _MANHOLE_RATE_DEFAULT
    try:
        _legacy_costing = price_with_defaults(26080, client_rates_path=_rates_path)
        _legacy_bytes = _json_rates.dumps(
            _legacy_costing, sort_keys=True, separators=(",", ":"), ensure_ascii=False
        ).encode()
        ck("absent client_rates.json keeps the complete legacy costing byte-identical",
           _hashlib_rates.sha256(_legacy_bytes).hexdigest() ==
           "38194af48023162689f095a4372d6293e7d4d0db87a71415b703ed576d0bab50" and
           "rates_version" not in _legacy_costing and
           "client_rates_applied" not in _legacy_costing)

        _edited_concrete = _rate_defaults["conc_rate"] * 1.01
        _edited_manhole = _rate_defaults["manhole_eo_rate"] * 1.01
        _saved_rates, _rate_changes = _save_client_rates_test(
            {"conc_rate": _edited_concrete, "manhole_eo_rate": _edited_manhole},
            _rate_defaults, path=_rates_path, who="assessor-token-authenticated",
            when="2026-07-18T10:00:00+00:00",
        )
        ck("first client-rates save creates version 1 and one audit row per changed field",
           _saved_rates["version"] == 1 and len(_saved_rates["audit"]) == 2 and
           {entry["field"] for entry in _saved_rates["audit"]} ==
           {"conc_rate", "manhole_eo_rate"})
        _audit_concrete = next(entry for entry in _saved_rates["audit"]
                               if entry["field"] == "conc_rate")
        ck("client-rates audit records authenticated assessor + exact old -> new",
           _audit_concrete["who"] == "assessor-token-authenticated" and
           _audit_concrete["old"] == _rate_defaults["conc_rate"] and
           _audit_concrete["new"] == _edited_concrete)

        _overridden_costing = price_with_defaults(
            26080, manhole_count=2, client_rates_path=_rates_path)
        ck("override changes only fresh pricing and carries rates provenance",
           _overridden_costing["rate"] != _legacy_costing["rate"] and
           _overridden_costing["total_gbp"] != _legacy_costing["total_gbp"] and
           _overridden_costing["rates_version"] == 1 and
           _overridden_costing["client_rates_applied"] is True)
        ck("manhole E/O uses the client override without changing its built-in rate",
           _overridden_costing["extras"][0]["rate"] == _edited_manhole and
           _MANHOLE_RATE_DEFAULT == _rate_defaults["manhole_eo_rate"])

        _rates_quote_result = {
            "file": "Fresh-Rates.pdf", "area_m2": 26080,
            "quotation_section": "External yard slabs",
            "costing": _overridden_costing, "flags": [],
        }
        _rates_quote = generate_quotation(
            _rates_quote_result, project="Rates Test", client="Fortel", ref="RATE-001")
        ck("fresh quotation records rates_version + CLIENT-EDITED provenance declaration",
           _rates_quote.get("rates_version") == 1 and
           _rates_quote.get("client_rates_applied") is True and
           any("CLIENT-EDITED RATES" in note and "version 1" in note
               for note in _rates_quote["declarations"]))
        ck("client-rate provenance is visible in text, HTML and JSON quotation outputs",
           all("CLIENT-EDITED RATES" in output for output in (
               quotation_text(_rates_quote), quotation_html(_rates_quote),
               quotation_json(_rates_quote))))
        _rates_book = _load_workbook(
            _BytesIO(quotation_xlsx(_rates_quote)), data_only=False)["REV_01"]
        ck("xlsx header records the applied client-rates version",
           _rates_book["D1"].value == "Client-edited rates version: 1")

        _rates_path.unlink()
        _restored_costing = price_with_defaults(26080, client_rates_path=_rates_path)
        _restored_bytes = _json_rates.dumps(
            _restored_costing, sort_keys=True, separators=(",", ":"), ensure_ascii=False
        ).encode()
        ck("removing client_rates.json restores the exact default output",
           _restored_bytes == _legacy_bytes and not _rates_path.exists())
    finally:
        shutil.rmtree(_rates_tmpdir, ignore_errors=True)
except (ImportError, OSError, ValueError) as _e:
    ck("client-editable rate override layer imports and runs", False, _e)

print("spec extractor — supplier fields")
from spec_extractor import extract_spec_from_text
_s5 = extract_spec_from_text("20mm crushed aggregate, 0.45 w/c ratio, S3 slump, air-entrained")
ck("aggregate 20mm extracted",  _s5.get("aggregate_mm") == 20)
ck("wc_ratio 0.45 extracted",   abs(_s5.get("wc_ratio", 0) - 0.45) < 0.001)
ck("slump S3 extracted",        _s5.get("slump_class") == "S3")
ck("air_entrained extracted",   _s5.get("air_entrained") is True)
_s6 = extract_spec_from_text("12mm aggregate 0.50 w/c S4 slump class CEM I")
ck("aggregate 12mm extracted",  _s6.get("aggregate_mm") == 12)
ck("slump S4 extracted",        _s6.get("slump_class") == "S4")

print("supplier inquiry generator")
from supplier_inquiry import generate_inquiry, format_cubes
ck("cubes calc 26080m² 190mm 3%",
   format_cubes(26080, 190, 0.03) == round(26080 * 0.190 * 1.03, 1))
_demo = {
    "area_m2": 3172, "project_name": "Test Project", "project_ref": "2132",
    "costing": {"spec": {"depth_mm": 190, "conc_mix": "C32/40", "cement_type": "CEM I",
                          "air_entrained": True, "aggregate_mm": 20, "wc_ratio": 0.45,
                          "slump_class": "S3", "conc_wastage": 0.03}}
}
_inq = generate_inquiry(_demo)
ck("inquiry has subject",        bool(_inq["subject"]))
ck("inquiry subject has mix",    "C32/40" in _inq["subject"])
ck("inquiry subject has cubes",  "m³" in _inq["subject"])
ck("inquiry text has slump",     "S3" in _inq["text"])
ck("inquiry text has aggregate", "20 mm" in _inq["text"])
ck("inquiry text has wc",        "0.45" in _inq["text"])
ck("inquiry html starts DOCTYPE","<!DOCTYPE" in _inq["html"])
ck("inquiry cubes > 0",          _inq["cubes_m3"] > 0)
_inq_no_proj = generate_inquiry({"area_m2": 1000, "costing": {}})
ck("inquiry works with no project info", bool(_inq_no_proj["subject"]))

