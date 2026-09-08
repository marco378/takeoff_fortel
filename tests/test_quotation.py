#!/usr/bin/env python3
"""quotation generation, BOQ rows and the xlsx/html/json/text exports.

Sections in this module (printed in this order):
  - quotation generator

Imported by ci_tests.py, which owns the running total. Do not reorder.
"""
from tests import ck
from slab_spec import build_brief_spec as _build_brief_spec
from tests.test_spec_extractor import _assumed_brief, _confirmed_brief, _e7, _effective_spec

print("quotation generator")
from quotation import (generate_quotation, quotation_text, quotation_html, quotation_json,
                       quotation_xlsx, SECTION_ORDER, PROVISIONAL_LABEL, PROVISIONAL_COL,
                       FORTEL_MH_ROW, FORTEL_CHANNEL_ROW, FORTEL_TRANSITION_ROW)
from geometry import polygon_perimeter_lm
from openpyxl import load_workbook as _load_workbook
from io import BytesIO as _BytesIO
import copy as _copy
_demo_result = {
    "file": "D77.pdf", "type": "UNMARKED vector", "confidence": "medium",
    "source_discipline": "architect",
    "costing": {
        "area_m2": 3172, "rate": 44.89, "total_gbp": 142391.08, "assumed": True,
        "spec": {"depth_mm": 190, "mesh": "A252", "conc_mix": "C32/40", "layers": 1, "conc_rate": 128},
        "breakdown": {"concrete": 25.05, "steel": 4.30, "dpm": 0.46,
                      "curing": 0.23, "labour": 10.00, "trim": 0.40, "nett": 40.44, "margin%": 11},
    },
    "flags": ["BUILD-UP ASSUMED: 190mm / A252 / C32/40"],
}
_q = generate_quotation(_demo_result, project="Test", client="Client", ref="TST-001")
ck("quotation total > 0",     _q["total_gbp"] > 0)
ck("quotation assumed=True",  _q["assumed"] is True)
ck("has declaration",         len(_q["declarations"]) >= 1)
ck("slab line item present",  any("slab" in li["description"].lower() for li in _q["line_items"]))
ck("text contains total",     "TOTAL NETT" in quotation_text(_q))
ck("html contains total",     "TOTAL NETT" in quotation_html(_q) or "Total" in quotation_html(_q))
ck("html is valid-ish",       quotation_html(_q).startswith("<!DOCTYPE html>"))

# Client-call quotation requirements: one editable tab, canonical section order, aggregate
# identical unit specs, retain different specs, mark assumptions provisional, and expose an
# informational perimeter without pricing it. Synthetic geometry only — no drawings fixture.
def _quotation_unit(filename, section, area, *, mesh="A252", assumed=False):
    unit = _copy.deepcopy(_demo_result)
    unit.update({"file": filename, "quotation_section": section, "area_m2": area,
                 "source_discipline": "engineer", "flags": []})
    unit["costing"].update({"area_m2": area, "assumed": assumed})
    unit["costing"]["spec"] = dict(unit["costing"]["spec"], mesh=mesh)
    return unit

_quote_units = [
    _quotation_unit("Upper.pdf", "Upper floor slabs", 40),
    _quotation_unit("Yard-A.pdf", "External yard slabs", 100, assumed=True),
    _quotation_unit("Footpath.pdf", "Footpath slabs", 15),
    _quotation_unit("Dock.pdf", "Dock slabs", 30),
    _quotation_unit("Ground.pdf", "Ground floor slabs", 20),
    _quotation_unit("Yard-B.pdf", "External yard slabs", 150, assumed=True),
]
_q_multi = generate_quotation(
    _quote_units, project="Multi-unit", client="Fortel", ref="TST-MULTI",
    extras=[{"section": "Prelims", "description": "Existing prelim item",
             "qty": 1, "unit": "Item", "rate": _demo_result["costing"]["rate"]}],
)
_actual_sections = list(dict.fromkeys(li["section"] for li in _q_multi["line_items"]))
ck("quotation sections follow client order", _actual_sections == list(SECTION_ORDER), _actual_sections)
_yard_slabs = [li for li in _q_multi["line_items"]
               if li["section"] == "External yard slabs"
               and li.get("line_role") == "concrete_slab"]
ck("matching-spec units aggregate into one slab row", len(_yard_slabs) == 1 and
   _yard_slabs[0]["qty"] == 250, _yard_slabs)
_q_diff_spec = generate_quotation([
    _quotation_unit("Yard-A.pdf", "External yard slabs", 100),
    _quotation_unit("Yard-C.pdf", "External yard slabs", 50, mesh="A393"),
], ref="TST-DIFF-SPEC")
ck("different unit specs remain separate on one quotation",
   len([li for li in _q_diff_spec["line_items"]
        if li.get("line_role") == "concrete_slab"]) == 2)
_unit1_ground = _quotation_unit(
    "Unit-1 Ground Floor Core.pdf", "Ground floor slabs", 100)
_unit5_ground = _quotation_unit(
    "Unit-5 Ground Floor Core.pdf", "Ground floor slabs", 80)
for _unit, _depth in ((_unit1_ground, 193), (_unit5_ground, 150)):
    _unit["costing"]["spec"].update({"depth_mm":_depth, "mesh":"A252", "layers":1})
    _unit["brief_spec"] = _build_brief_spec(
        "ground_floor", effective_spec=_unit["costing"]["spec"],
        confirmed={"depth_mm":_depth, "mesh":"A252", "layers":1},
        source="engineer_drawing")
_q_ground_spec_split = generate_quotation(
    [_unit1_ground, _unit5_ground], ref="TST-GROUND-SPEC-SPLIT")
_ground_spec_rows = [
    item for item in _q_ground_spec_split["line_items"]
    if item["section"] == "Ground floor slabs"
    and item.get("line_role") == "concrete_slab"
]
ck("193mm A252 and 150mm A252 ground slabs in one case are never merged",
   len(_ground_spec_rows) == 2 and
   {item["qty"] for item in _ground_spec_rows} == {100.0, 80.0} and
   {spec["fields"]["depth_mm"]["value"]
    for spec in _q_ground_spec_split["specifications"]} == {193, 150},
   {"items":_ground_spec_rows, "specs":_q_ground_spec_split["specifications"]})
_confirmed_unit_a = _quotation_unit("Confirmed-A.pdf", "External yard slabs", 60)
_confirmed_unit_b = _quotation_unit("Confirmed-B.pdf", "External yard slabs", 40)
_confirmed_unit_a["brief_spec"] = _confirmed_brief
_confirmed_unit_b["brief_spec"] = _copy.deepcopy(_confirmed_brief)
_q_confirmed_aggregate = generate_quotation([_confirmed_unit_a, _confirmed_unit_b])
ck("equally confirmed Brief_Spec units aggregate",
   len([li for li in _q_confirmed_aggregate["line_items"]
        if li.get("line_role") == "concrete_slab"]) == 1)
_assumed_unit = _quotation_unit("Assumed.pdf", "External yard slabs", 40, assumed=True)
_assumed_unit["brief_spec"] = _assumed_brief
_q_provenance_split = generate_quotation([_confirmed_unit_a, _assumed_unit])
ck("assumed and confirmed equal-value Brief_Spec units remain separate",
   len([li for li in _q_provenance_split["line_items"]
        if li.get("line_role") == "concrete_slab"]) == 2)

_rect = [[0, 0], [40, 0], [40, 20], [0, 20]]
ck("perimeter_lm rectangle: 20m × 10m -> 60.0m",
   polygon_perimeter_lm(_rect, 0.5) == 60.0)
_perimeter_result = _quotation_unit("Dock-Perimeter.pdf", "Dock slabs", 200)
_perimeter_result.update({"polygon_pts": _rect, "scale_k": 0.5})
_q_perimeter = generate_quotation(_perimeter_result, ref="TST-PERIMETER")
ck("quotation surfaces perimeter as informational, unpriced quantity",
   _q_perimeter["perimeter_lm"] == 60.0 and
   any(m["description"] == "Slab perimeter" and m["qty"] == 60.0
       for m in _q_perimeter["measurements"]) and
   all("perimeter" not in li["description"].lower() for li in _q_perimeter["line_items"]))

# A real case can mix pipeline-measured siblings (top-level perimeter_lm) with an
# assessor-traced sibling whose perimeter lives on a zone.  Both sources intentionally
# share the same measurement key and must aggregate without either provenance path assuming
# it was the one that created the record.
_mixed_top_perimeter = _quotation_unit(
    "Yard Unit-1.pdf", "External yard slabs", 100)
_mixed_top_perimeter["perimeter_lm"] = 40.0
_mixed_zone_perimeter = _quotation_unit(
    "Yard Unit-2.pdf", "External yard slabs", 120)
_mixed_zone_perimeter["zones"] = [{
    "category": "external_yard", "area_m2": 120.0, "perimeter_lm": 50.0,
}]
_q_mixed_perimeter = generate_quotation(
    [_mixed_top_perimeter, _mixed_zone_perimeter],
    project="Mixed perimeter case", ref="TST-MIXED-PERIMETER")
_mixed_outputs = (
    quotation_text(_q_mixed_perimeter), quotation_html(_q_mixed_perimeter),
    quotation_json(_q_mixed_perimeter), quotation_xlsx(_q_mixed_perimeter),
)
_mixed_ws = _load_workbook(
    _BytesIO(_mixed_outputs[3]), data_only=False)["REV_01"]
_mixed_measurement = next(
    measurement for measurement in _q_mixed_perimeter["measurements"]
    if measurement["section"] == "External yard slabs"
    and measurement["description"] == "Slab perimeter")
ck("case quotation renders txt/html/json/xlsx with mixed perimeter provenance",
   _mixed_measurement["qty"] == 90.0 and
   len(_mixed_measurement["quantity_rows"]) == 2 and
   all(output for output in _mixed_outputs[:3]) and
   _mixed_ws.max_row > 7,
   _mixed_measurement)

_q_text = quotation_text(_q)
_q_html = quotation_html(_q)
_q_json = quotation_json(_q)
ck("assumed build-up is provisional in text/html/json",
   all(PROVISIONAL_LABEL in output for output in (_q_text, _q_html, _q_json)))
ck("unknown client spec fields are visible in text/html/json",
   all("ASSUMED / no details provided" in output for output in (_q_text, _q_html, _q_json)))
_xss_brief = _build_brief_spec(
    "external_yard", effective_spec=_effective_spec,
    confirmed=dict(_effective_spec, joint_details="<script>alert(1)</script>"),
)
_xss_unit = _quotation_unit("Safe.pdf", "External yard slabs", 10)
_xss_unit["brief_spec"] = _xss_brief
_xss_html = quotation_html(generate_quotation(_xss_unit))
ck("assessor-entered Brief_Spec text is HTML-escaped in served quotation",
   "<script>alert(1)</script>" not in _xss_html and
   "&lt;script&gt;alert(1)&lt;/script&gt;" in _xss_html)

_cited_brief = _build_brief_spec(
    "external_yard", effective_spec=_effective_spec,
    confirmed={"depth_mm": 190, "mesh": "A252"}, source="engineer_drawing",
    evidence={
        "depth_mm": {"file": "64426-112-External-Works-Details-Rev.T2.pdf", "page": 2,
                     "text": "minimum 190mm developer specification"},
        "mesh": {"file": "64426-112-External-Works-Details-Rev.T2.pdf", "page": 2,
                 "text": "A252 mesh fabric reinforcement"},
    },
)
_cited_unit = _quotation_unit("64426-111-Yard.pdf", "External yard slabs", 10)
_cited_unit["brief_spec"] = _cited_brief
_cited_quote = generate_quotation(_cited_unit, ref="TST-SPEC-CITATION")
_cited_outputs = (
    quotation_text(_cited_quote), quotation_html(_cited_quote),
    quotation_json(_cited_quote), quotation_xlsx(_cited_quote),
)
_cited_wb = _load_workbook(_BytesIO(_cited_outputs[3]), data_only=False)
_cited_xlsx_text = "\n".join(
    str(cell.value or "") for row in _cited_wb.active.iter_rows() for cell in row)
_source_citation = "64426-112-External-Works-Details-Rev.T2.pdf, page 2"
ck("every quotation format cites the drawing file and page for extracted specification fields",
   all(_source_citation in output for output in _cited_outputs[:3]) and
   _source_citation in _cited_xlsx_text,
   {"text": _cited_outputs[0], "xlsx": _cited_xlsx_text})

_xlsx_bytes = quotation_xlsx(_q_multi)
_xlsx_wb = _load_workbook(_BytesIO(_xlsx_bytes), data_only=False)
_xlsx_ws = _xlsx_wb["REV_01"]
ck("xlsx export reopens as exactly one editable quotation tab", _xlsx_wb.sheetnames == ["REV_01"])
ck("xlsx uses client BOQ column labels and order",
   tuple(_xlsx_ws.cell(7, col).value for col in range(1, 6)) ==
   ("DESCRIPTION", "QTY", "UNIT", "RATE", "VALUE"))
ck("xlsx header matches real BOQ project/client/date/rev/drawing layout",
   _xlsx_ws["A1"].value == "Project: Multi-unit" and
   _xlsx_ws["A2"].value == "Client: Fortel" and
   str(_xlsx_ws["A3"].value).startswith("Date: ") and
   _xlsx_ws["A4"].value == "Rev: TST-MULTI" and
   {str(cell_range) for cell_range in _xlsx_ws.merged_cells.ranges} == {"D4:E4"})
ck("xlsx drawing register is multiline in the real BOQ's A5 cell",
   _xlsx_ws["A5"].value.startswith("Drawing ref available at tender:\n") and
   all(name in _xlsx_ws["A5"].value for name in ("Yard-A.pdf", "Yard-B.pdf")))
_expected_xlsx_sections = (
    "External Yard Slabs- Provisional Cost (No Details)",
    "Footpath Slabs- Provisional Cost (No Details)",
    "Dock Slabs- Provisional Cost (No Details)",
    "Ground Floor Slabs- Provisional Cost (No Details)",
    "Upper Floors- Provisional Cost (No Details)",
    "Prelims",
)
_xlsx_section_rows = {
    _xlsx_ws.cell(row, 1).value: row for row in range(1, _xlsx_ws.max_row + 1)
    if _xlsx_ws.cell(row, 1).value in _expected_xlsx_sections
}
ck("xlsx section headers follow client order",
   tuple(_xlsx_section_rows) == _expected_xlsx_sections, _xlsx_section_rows)
_xlsx_item_row = next(row for row in range(1, _xlsx_ws.max_row + 1)
                      if "Concrete Slabs" in str(_xlsx_ws.cell(row, 1).value or ""))
_xlsx_source_rows = [row for row in range(1, _xlsx_ws.max_row + 1)
                     if _xlsx_ws.cell(row, 1).value in ("Yard-A.pdf", "Yard-B.pdf")]
_xlsx_area_total_row = next(row for row in range(1, _xlsx_ws.max_row + 1)
                            if _xlsx_ws.cell(row, 1).value == "Total Area Take Off:")
ck("xlsx keeps editable numeric per-unit source quantities and formula aggregate",
   [float(_xlsx_ws.cell(row, 2).value) for row in _xlsx_source_rows] == [100.0, 150.0] and
   _xlsx_ws.cell(_xlsx_area_total_row, 2).value ==
   f"=SUM(B{_xlsx_source_rows[0]}:B{_xlsx_source_rows[-1]})")
ck("xlsx priced qty references aggregate, rate is numeric, and value is rounded qty*rate",
   _xlsx_ws.cell(_xlsx_item_row, 2).value == f"=B{_xlsx_area_total_row}" and
   isinstance(_xlsx_ws.cell(_xlsx_item_row, 4).value, (int, float)) and
   _xlsx_ws.cell(_xlsx_item_row, 5).data_type == "f" and
   _xlsx_ws.cell(_xlsx_item_row, 5).value ==
   f"=ROUND(B{_xlsx_item_row}*D{_xlsx_item_row},2)")
_yard_section_start = _xlsx_section_rows[_expected_xlsx_sections[0]]
_footpath_section_start = _xlsx_section_rows[_expected_xlsx_sections[1]]
_yard_rows_by_label = {
    str(_xlsx_ws.cell(row, 1).value or "").splitlines()[0]: row
    for row in range(_yard_section_start + 1, _footpath_section_start)
    if str(_xlsx_ws.cell(row, 1).value or "").splitlines()
}
_included_labels = (
    "A252 Mesh Fabric x Single Layer", "Curing Agent",
    "DPM 1200G (Excl. tapes and seals to laps)", "Brush Finish",
)
ck("xlsx uses Fortel's Incl. convention without splitting or inventing component rates",
   all(label in _yard_rows_by_label and
       _xlsx_ws.cell(_yard_rows_by_label[label], 2).value == f"=B{_xlsx_area_total_row}" and
       _xlsx_ws.cell(_yard_rows_by_label[label], 4).value is None and
       _xlsx_ws.cell(_yard_rows_by_label[label], 5).value == "Incl."
       for label in _included_labels),
   {label:(_xlsx_ws.cell(_yard_rows_by_label.get(label, 1), 4).value,
           _xlsx_ws.cell(_yard_rows_by_label.get(label, 1), 5).value)
    for label in _included_labels})
ck("xlsx quantity/unit display matches client template without forced .00",
   _xlsx_ws.cell(_xlsx_item_row, 2).number_format == "#,##0.##" and
   _xlsx_ws.cell(_xlsx_item_row, 3).value == "m2")
_xlsx_total_row = next(row for row in range(1, _xlsx_ws.max_row + 1)
                       if _xlsx_ws.cell(row, 1).value == "TOTAL NETT")
ck("xlsx follows real BOQ: no section subtotals and one nett formula",
   not any(str(_xlsx_ws.cell(row, 1).value or "").startswith("Subtotal —")
           for row in range(1, _xlsx_ws.max_row + 1)) and
   _xlsx_ws.cell(_xlsx_total_row, 5).value == f"=SUM(E7:E{_xlsx_total_row - 1})")
ck("xlsx matches real BOQ widths, accounting display, and portrait layout",
   abs(_xlsx_ws.column_dimensions["A"].width - 82.43) < .01 and
   "£" in _xlsx_ws.cell(_xlsx_item_row, 5).number_format and
   _xlsx_ws.page_setup.orientation == "portrait" and _xlsx_ws.freeze_panes is None and
   _xlsx_ws.auto_filter.ref is None)
ck("xlsx visibly marks assumed quantity provisional",
   # Marker relocated out of DESCRIPTION into its own column; check either so the assertion is
   # about the marker being VISIBLE, not about which cell holds it.
   any(PROVISIONAL_LABEL in str(_xlsx_ws.cell(row, col).value or "")
       for row in range(1, _xlsx_ws.max_row + 1) for col in (1, PROVISIONAL_COL)))
ck("xlsx visibly carries every unknown client checklist field",
   any("Bay sizes if joint layout available: ASSUMED / no details provided" in
       str(_xlsx_ws.cell(row, 1).value or "") for row in range(1, _xlsx_ws.max_row + 1)))
_dock_foundation_row = next(
    row for row in range(1, _xlsx_ws.max_row + 1)
    if _xlsx_ws.cell(row, 1).value ==
       "Foundation thickenings directly underneath Dock Slab region")
ck("xlsx carries the Dock foundation-thickening subgroup without inventing a quantity or rate",
   all(_xlsx_ws.cell(_dock_foundation_row, col).value is None for col in range(2, 6)))

_bay_unit = _quotation_unit("Unit-1 Yard Joint Layout.pdf", "External yard slabs", 100)
_bay_unit["brief_spec"] = _build_brief_spec(
    "external_yard", effective_spec=_bay_unit["costing"]["spec"],
    confirmed={"bay_sizes":_e7["bay_sizes"]}, source="engineer_drawing",
    evidence={"bay_sizes":(_e7.get("_evidence") or {}).get("bay_sizes")},
)
_bay_quote = generate_quotation(_bay_unit, ref="TST-JOINT-BAYS")
_bay_joint_row = next(
    item for item in _bay_quote["line_items"]
    if item["description"].startswith("Joints (Excl. Mastic)"))
ck("extracted bay dimensions flow into the Fortel Joints row wording",
   _bay_joint_row["description"] ==
   "Joints (Excl. Mastic) - Based on 4.9m wide x 6.5m long bays",
   _bay_joint_row)
_q_status = generate_quotation(
    _quotation_unit("Status.pdf", "External yard slabs", 10),
    extras=[{"section": "Prelims", "description": "Commercial option", "qty": 1,
             "unit": "Item", "rate": 12.34, "value_status": "RATE ONLY"}],
)
_status_ws = _load_workbook(_BytesIO(quotation_xlsx(_q_status)), data_only=False)["REV_01"]
_status_row = next(row for row in range(1, _status_ws.max_row + 1)
                   if _status_ws.cell(row, 1).value == "Commercial option")
ck("xlsx supports the real BOQ's explicit RATE ONLY value token without inferring it",
   _status_ws.cell(_status_row, 4).value == 12.34 and
   _status_ws.cell(_status_row, 5).value == "RATE ONLY")

_channel_quote_unit = _quotation_unit("External Unit-1.pdf", "External yard slabs", 100)
_channel_quote_unit["channel_proposals"] = [{
    "proposal_id":"channel-dock-loading-face", "component":"dock_retaining_wall",
    "proposed_length_lm":90.0,
    "basis":"two straight runs where channels are not drawn",
}]
_channel_quote_unit["channel_proposal_decisions"] = {
    "channel-dock-loading-face": {
        "decision":"accepted", "length_lm":96.7, "edited":True,
    }
}
_q_channel = generate_quotation(_channel_quote_unit, ref="CHANNEL-QUOTE-001")
_channel_rows = [item for item in _q_channel["line_items"]
                 if item["description"] == FORTEL_CHANNEL_ROW]
ck("accepted/edited channel proposal becomes a provisional blank-rate Lm quote line",
   len(_channel_rows) == 1 and _channel_rows[0]["qty"] == 96.7 and
   _channel_rows[0]["unit"] == "Lm" and _channel_rows[0]["rate"] is None and
   _channel_rows[0]["value"] is None and _channel_rows[0]["assessor_rate_required"] and
   _channel_rows[0]["provisional"], _channel_rows)
_channel_ws = _load_workbook(_BytesIO(quotation_xlsx(_q_channel)), data_only=False)["REV_01"]
_channel_row = next(row for row in range(1, _channel_ws.max_row + 1)
                    if str(_channel_ws.cell(row, 1).value or "").startswith(
                        FORTEL_CHANNEL_ROW))
ck("accepted channel quantity is exact in text/HTML/XLSX and is never auto-priced",
   "96.7" in quotation_text(_q_channel) and "96.7 Lm" in quotation_html(_q_channel) and
   _channel_ws.cell(_channel_row, 2).value == 96.7 and
   _channel_ws.cell(_channel_row, 4).value is None and
   _channel_ws.cell(_channel_row, 5).data_type == "f" and
   # The provisional marker moved out of DESCRIPTION into its own column (right of VALUE,
   # mirroring Fortel's REMEASURE caveat column) so rows stop wrapping onto two lines. The row
   # must still be MARKED provisional — only where the marker sits changed.
   PROVISIONAL_LABEL in str(_channel_ws.cell(_channel_row, PROVISIONAL_COL).value))
_pending_channel_unit = _copy.deepcopy(_channel_quote_unit)
_pending_channel_unit["channel_proposal_decisions"] = {}
_q_pending_channel = generate_quotation(_pending_channel_unit, ref="CHANNEL-PENDING-001")
ck("unactioned channel proposal is declared explicitly and never becomes a quote quantity",
   not any(item["description"] == FORTEL_CHANNEL_ROW
           for item in _q_pending_channel["line_items"]) and
   any("has not been actioned" in note and "no channel quantity or price is included" in note
       for note in _q_pending_channel["declarations"]))

_transition_quote_unit = _quotation_unit("External Unit-2.pdf", "External yard slabs", 100)
_transition_quote_unit["transition_candidates"] = [{
    "candidate_id":"transition-yard-region-1", "region_id":"yard-region-1",
    "proposed_length_lm":15.0,
    "basis":"macadam-to-concrete boundary at the Yard entrance",
}]
_transition_quote_unit["transition_candidate_decisions"] = {
    "transition-yard-region-1": {
        "decision":"accepted", "length_lm":17.5, "edited":True,
    }
}
_q_transition_candidate = generate_quotation(
    _transition_quote_unit, ref="TRANSITION-QUOTE-001")
_transition_candidate_rows = [
    item for item in _q_transition_candidate["line_items"]
    if item["description"] == FORTEL_TRANSITION_ROW
]
ck("accepted/edited Transition candidate becomes provisional blank-rate Lm quote line",
   len(_transition_candidate_rows) == 1 and
   _transition_candidate_rows[0]["qty"] == 17.5 and
   _transition_candidate_rows[0]["unit"] == "Lm" and
   _transition_candidate_rows[0]["rate"] is None and
   _transition_candidate_rows[0]["value"] is None and
   _transition_candidate_rows[0]["assessor_rate_required"] and
   _transition_candidate_rows[0]["provisional"] and
   "macadam-to-concrete" in _transition_candidate_rows[0]["assumption_basis"],
   _transition_candidate_rows)
_pending_transition_unit = _copy.deepcopy(_transition_quote_unit)
_pending_transition_unit["transition_candidate_decisions"] = {}
_q_pending_transition = generate_quotation(
    _pending_transition_unit, ref="TRANSITION-PENDING-001")
ck("unactioned Transition candidate stays outside totals and is declared explicitly",
   not any(item["description"] == FORTEL_TRANSITION_ROW
           for item in _q_pending_transition["line_items"]) and
   any("has not been actioned" in note and
       "no Transition quantity or price is included" in note
       for note in _q_pending_transition["declarations"]),
   _q_pending_transition["declarations"])

# Fortel's standard costing sheet has one visible extra-over row for each quantity class.
# Exercise all three together, including two accepted channel runs, so the regression test
# proves the final workbook shape rather than three isolated generator branches.
_fortel_eo_unit = _quotation_unit("External Unit-3.pdf", "External yard slabs", 100)
_fortel_eo_unit["perimeter_lm"] = 80.0
_fortel_eo_unit["channel_proposals"] = [
    {"proposal_id":"channel-dock", "component":"dock_retaining_wall",
     "basis":"assessor-reviewed dock run"},
    {"proposal_id":"channel-yard", "component":"yard_wall_adjacent_run",
     "basis":"assessor-reviewed Yard run"},
]
_fortel_eo_unit["channel_proposal_decisions"] = {
    "channel-dock":{"decision":"accepted", "length_lm":12.5},
    "channel-yard":{"decision":"accepted", "length_lm":20.0},
}
_fortel_eo_unit["transition_candidates"] = [{
    "candidate_id":"transition-yard-region-1", "region_id":"yard-region-1",
    "basis":"assessor-reviewed Yard entrance",
}]
_fortel_eo_unit["transition_candidate_decisions"] = {
    "transition-yard-region-1":{"decision":"accepted", "length_lm":8.25},
}
_q_fortel_eo = generate_quotation(
    _fortel_eo_unit, ref="FORTEL-EO-ROWS",
    extras=[{"section":"External yard slabs", "description":FORTEL_MH_ROW,
             "qty":2, "unit":"Nr", "rate":None}],
)
_fortel_eo_lines = {
    item["description"]:item for item in _q_fortel_eo["line_items"]
    if item["description"] in {FORTEL_MH_ROW, FORTEL_CHANNEL_ROW, FORTEL_TRANSITION_ROW}
}
ck("channel, Transition and manhole quantities use exactly Fortel's three row labels",
   set(_fortel_eo_lines) == {FORTEL_MH_ROW, FORTEL_CHANNEL_ROW, FORTEL_TRANSITION_ROW} and
   _fortel_eo_lines[FORTEL_MH_ROW]["unit"] == "Nr" and
   _fortel_eo_lines[FORTEL_CHANNEL_ROW]["unit"] == "Lm" and
   _fortel_eo_lines[FORTEL_TRANSITION_ROW]["unit"] == "Lm" and
   _fortel_eo_lines[FORTEL_CHANNEL_ROW]["qty"] == 32.5,
   _fortel_eo_lines)
_fortel_eo_ws = _load_workbook(
    _BytesIO(quotation_xlsx(_q_fortel_eo)), data_only=False)["REV_01"]
_fortel_eo_row_numbers = {
    str(_fortel_eo_ws.cell(row, 1).value or "").splitlines()[0]:row
    for row in range(1, _fortel_eo_ws.max_row + 1)
    if str(_fortel_eo_ws.cell(row, 1).value or "").splitlines() and
       str(_fortel_eo_ws.cell(row, 1).value or "").splitlines()[0] in
       {"Slab perimeter", FORTEL_MH_ROW, FORTEL_CHANNEL_ROW, FORTEL_TRANSITION_ROW}
}
ck("xlsx writes perimeter then MH/channel/Transition with editable blank rates and formulas",
   list(_fortel_eo_row_numbers) == [
       "Slab perimeter", FORTEL_MH_ROW, FORTEL_CHANNEL_ROW, FORTEL_TRANSITION_ROW] and
   [_fortel_eo_ws.cell(_fortel_eo_row_numbers[label], 3).value
    for label in (FORTEL_MH_ROW, FORTEL_CHANNEL_ROW, FORTEL_TRANSITION_ROW)] ==
       ["Nr", "Lm", "Lm"] and
   all(_fortel_eo_ws.cell(_fortel_eo_row_numbers[label], 4).value is None and
       _fortel_eo_ws.cell(_fortel_eo_row_numbers[label], 5).value ==
       f'=IF(D{_fortel_eo_row_numbers[label]}="","",ROUND('
       f'B{_fortel_eo_row_numbers[label]}*D{_fortel_eo_row_numbers[label]},2))'
       for label in (FORTEL_MH_ROW, FORTEL_CHANNEL_ROW, FORTEL_TRANSITION_ROW)),
   _fortel_eo_row_numbers)

