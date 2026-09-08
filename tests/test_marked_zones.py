#!/usr/bin/env python3
"""marked zone-aware measurement, BOQ allocation, Yard/Dock split, channels.

Sections in this module (printed in this order):
  - marked zone-aware measurement + multi-unit BOQ allocation
  - raw external zone measurement — Yard/Dock split + low-confidence state cap
  - raw multi-region Yard fixture + assessor keep/exclude loop
  - assessor Yard-region exclusion endpoint
  - channel proposal geometry — retaining-wall adjacent and never diagonal

Imported by ci_tests.py, which owns the running total. Do not reorder.
"""
import shutil
from geometry import measure_regions
from pathlib import Path
from tests import _FixtureNotPresent, _require_fixture, ck
from io import BytesIO as _BytesIO
from openpyxl import load_workbook as _load_workbook
from quotation import FORTEL_CHANNEL_ROW, FORTEL_TRANSITION_ROW, generate_quotation, quotation_xlsx
from slab_spec import empty_brief_spec as _empty_brief_spec
from tests.test_quotation import _expected_xlsx_sections, _quotation_unit

print("marked zone-aware measurement + multi-unit BOQ allocation")
import fitz as _fitz_zones
from robust_takeoff import read_marked as _read_marked_legacy, read_marked_zones as _read_marked_zones
from markup import parse_area_m2 as _parse_area_m2
ck("prefix-free Bluebeam measurement lines parse without accepting incidental area prose",
   _parse_area_m2("Unit-1&2\r235.37 sq m") == 235.37 and
   _parse_area_m2("Unit-3\r133.79 sq m") == 133.79 and
   _parse_area_m2("Rate note: allow 12.50 sq m at this location") is None and
   _parse_area_m2("Pricing note A = 12.50 sq m at this rate") is None and
   _parse_area_m2("This note contains 12.50 sq m for context only") is None)
ck("area truth parser still rejects linear, thickness and scale labels",
   all(_parse_area_m2(label) is None for label in ("12.30 m", "150 mm", "1:200")))
_zone_pdf = "/tmp/ci_marked_zones.pdf"
_zone_doc = _fitz_zones.open()
_zone_page = _zone_doc.new_page(width=600, height=600)
for _subject, _value, _rect in (
        ("Yard", 100, (20, 20, 220, 220)),
        ("Dock ", 30, (250, 20, 400, 120)),
        ("Mystery slab", 5, (250, 150, 350, 250))):
    _annot = _zone_page.add_polygon_annot([
        (_rect[0], _rect[1]), (_rect[2], _rect[1]),
        (_rect[2], _rect[3]), (_rect[0], _rect[3]),
    ])
    _annot.set_info(title="Fortel QA", subject=_subject, content=f"Area\n{_value:.2f} sq m")
    _annot.update()
_channel_annot = _zone_page.add_polyline_annot([(20, 300), (120, 300), (160, 330)])
_channel_annot.set_info(title="Fortel QA", subject="Channel", content="Channel\n12.50 m")
_channel_annot.update()
_zone_doc.save(_zone_pdf)
_zone_doc.close()
_zone_read = _read_marked_zones(_zone_pdf)
_zone_by_category = {zone["category"]: zone for zone in _zone_read["zones"]}
ck("marked zone reader preserves legacy aggregate across every labelled polygon",
   _zone_read["area_m2"] == 135.0 and _zone_read["regions"] == 3 and
   _read_marked_legacy(_zone_pdf) == (135.0, 3), _zone_read)
ck("Bluebeam subjects separate Yard, Dock and Channel without colour fallback",
   _zone_by_category["external_yard"]["area_m2"] == 100.0 and
   _zone_by_category["dock"]["area_m2"] == 30.0 and
   _zone_by_category["channel"]["length_lm"] == 12.5, _zone_by_category)
ck("unknown measurable subject is unclassified and visibly flagged",
   _zone_by_category["unclassified"]["area_m2"] == 5.0 and
   any("assessor: classify zone 'Mystery slab'" in flag for flag in _zone_read["flags"]),
   _zone_read["flags"])
ck("zone reader retains per-annotation subject/author/colour evidence",
   len(_zone_read["markup_annotations"]) == 4 and
   all("subject" in record and "author" in record and "stroke_color" in record
       for record in _zone_read["markup_annotations"]))

_exclusion_pdf = "/tmp/ci_marked_slab_exclusions.pdf"
_exclusion_doc = _fitz_zones.open()
_exclusion_page = _exclusion_doc.new_page(width=600, height=600)
for _subject, _value, _x in (
        ("Ground floor core", 500.0, 20),
        ("Lift shaft", 12.0, 260),
        ("Data riser", 4.0, 380),
        ("Precast staircase foundation", 8.0, 480)):
    _annot = _exclusion_page.add_polygon_annot([
        (_x, 20), (_x + 80, 20), (_x + 80, 120), (_x, 120),
    ])
    _annot.set_info(title="Fortel QA", subject=_subject,
                    content=f"Area\n{_value:.2f} sq m")
    _annot.update()
_exclusion_doc.save(_exclusion_pdf)
_exclusion_doc.close()
_exclusion_read = _read_marked_zones(_exclusion_pdf)
ck("explicit lift/riser/stair-foundation markups are excluded, recorded and never summed",
   _exclusion_read["area_m2"] == 500.0 and
   _exclusion_read["regions"] == 1 and
   len(_exclusion_read["zones"]) == 1 and
   {item["exclusion_id"] for item in _exclusion_read["exclusions"]} == {
       "lift_void", "service_data_riser", "precast_stair_foundation"} and
   all(record["excluded_from_slab"]
       for record in _exclusion_read["markup_annotations"][1:]),
   _exclusion_read)
from measurement_rules import (
    classify_exclusion as _classify_client_exclusion,
    exclusion_review_prompts as _exclusion_review_prompts,
)
_yard_exclusion_prompts = _exclusion_review_prompts(
    ["external_yard"], "Proposed Gatehouse and Hub Office")
ck("raw labels create visible unresolved Gatehouse/Hub-office prompts, never fake geometry",
   {prompt["exclusion_id"] for prompt in _yard_exclusion_prompts
    if prompt["status"] == "outline_unresolved"} == {"gatehouse", "hub_office"} and
   all(prompt.get("requires_assessor_confirmation")
       for prompt in _yard_exclusion_prompts), _yard_exclusion_prompts)
ck("a bare pit is never guessed to be a lift void or precast stair foundation",
   _classify_client_exclusion("Pit", "300 345 600 mm") is None)
ck("a lift lobby is slab circulation space, not guessed to be a lift void",
   _classify_client_exclusion("Lift lobby", "") is None)
_explicit_client_exclusions = {
    subject: _classify_client_exclusion(subject, content)
    for subject, content in (
        ("Gatehouse", ""),
        ("Hub office", ""),
        ("Area Measurement", "Lift bit"),
        ("Riser data", ""),
        ("Pre-cast concrete staircase foundation", ""),
    )
}
ck("5-Aug exclusion wording is recognised from explicit subject/content evidence",
   {subject: record and record.get("exclusion_id")
    for subject, record in _explicit_client_exclusions.items()} == {
       "Gatehouse":"gatehouse", "Hub office":"hub_office",
       "Area Measurement":"lift_void", "Riser data":"service_data_riser",
       "Pre-cast concrete staircase foundation":"precast_stair_foundation",
   }, _explicit_client_exclusions)
_ground_exclusion_checklist = _exclusion_review_prompts(["ground_floor"], "")
ck("every ground-floor trace keeps the lift/riser/stair checklist visible without a text hit",
   {prompt["exclusion_id"] for prompt in _ground_exclusion_checklist} == {
       "lift_void", "service_data_riser", "precast_stair_foundation"} and
   all(prompt.get("status") == "assessor_check" and
       prompt.get("requires_assessor_confirmation") and prompt.get("assumed") and
       prompt.get("basis") for prompt in _ground_exclusion_checklist),
   _ground_exclusion_checklist)
_labelled_ground_exclusions = _exclusion_review_prompts(
    ["ground_floor"], "Lift pit and data riser")
ck("drawing-text evidence gates unresolved lift/data outlines instead of missing slash labels",
   {prompt["exclusion_id"] for prompt in _labelled_ground_exclusions
    if prompt["status"] == "outline_unresolved"} == {
       "lift_void", "service_data_riser"}, _labelled_ground_exclusions)

_scope_pdf = "/tmp/ci_internal_warehouse_scope.pdf"
_scope_doc = _fitz_zones.open()
_scope_page = _scope_doc.new_page(width=600, height=300)
for _subject, _value, _x in (("Ground floor core", 120.0, 20),
                             ("Internal warehouse slab", 900.0, 260)):
    _annot = _scope_page.add_polygon_annot([
        (_x, 20), (_x + 180, 20), (_x + 180, 180), (_x, 180),
    ])
    _annot.set_info(title="Fortel QA", subject=_subject,
                    content=f"Area\n{_value:.2f} sq m")
    _annot.update()
_scope_doc.save(_scope_pdf)
_scope_doc.close()
_scope_read = _read_marked_zones(_scope_pdf)
ck("explicit internal warehouse slabs stay visible but never enter our area or BOQ zones",
   _scope_read["area_m2"] == 120.0 and _scope_read["regions"] == 1 and
   len(_scope_read["zones"]) == 1 and
   _scope_read["zones"][0]["category"] == "ground_floor" and
   any(item["exclusion_id"] == "internal_warehouse_scope" and
       item["area_m2"] == 900.0 for item in _scope_read["exclusions"]) and
   any("Internal warehouse slab" in flag for flag in _scope_read["flags"]),
   _scope_read)

_transition_pdf = "/tmp/ci_unit_yard_transitions.pdf"
_transition_doc = _fitz_zones.open()
_transition_page = _transition_doc.new_page(width=600, height=300)
for _subject, _value, _y in (("Unit 1 Transition", 12.5, 50),
                             ("Unit 2 Transition", 8.25, 150)):
    _annot = _transition_page.add_polyline_annot([(20, _y), (220, _y)])
    _annot.set_info(title="Fortel QA", subject=_subject,
                    content=f"{_value:.2f} m")
    _annot.update()
_transition_doc.save(_transition_pdf)
_transition_doc.close()
_transition_read = _read_marked_zones(_transition_pdf)
_transition_zones = [zone for zone in _transition_read["zones"]
                     if zone["category"] == "transition"]
ck("yard-entrance transitions remain one measured Lm zone per unit, never one anonymous lump",
   len(_transition_zones) == 2 and
   [zone["unit_label"] for zone in _transition_zones] == ["unit 1", "unit 2"] and
   [zone["length_lm"] for zone in _transition_zones] == [12.5, 8.25] and
   all("tarmac-to-concrete" in zone["basis"] for zone in _transition_zones),
   _transition_zones)
_transition_quote = generate_quotation({
    "file":"Site yard transitions.pdf", "source_discipline":"engineer",
    "zones":_transition_zones, "flags":[], "area_m2":None,
}, project="Transition QA")
_transition_measurement = next(
    item for item in _transition_quote["measurements"]
    if item["description"] == FORTEL_TRANSITION_ROW)
ck("measured Yard transitions reach the quotation as blank-rate Lm source rows",
   _transition_measurement["qty"] == 20.75 and
   _transition_measurement["assessor_rate_required"] and
   [row["description"] for row in _transition_measurement["quantity_rows"]] ==
   ["unit 1", "unit 2"], _transition_measurement)

from takeoff_unmarked import detect_raw_construction_joint as _detect_raw_cj
_raw_cj_pdf = "/tmp/ci_raw_construction_joint.pdf"
_raw_cj_doc = _fitz_zones.open()
_raw_cj_page = _raw_cj_doc.new_page(width=600, height=500)
_raw_cj_page.insert_text((180, 80), "CJ indicates internal construction joint")
_raw_cj_legend = _raw_cj_page.new_shape()
_raw_cj_legend.draw_line((100, 76), (160, 76))
_raw_cj_legend.finish(color=(0.0, 0.75, 0.0)); _raw_cj_legend.commit()
_raw_cj_body = _raw_cj_page.new_shape()
_raw_cj_body.draw_line((50, 200), (150, 200))
_raw_cj_body.draw_line((200, 300), (300, 300))
_raw_cj_body.finish(color=(0.0, 0.75, 0.0)); _raw_cj_body.commit()
_raw_cj_other = _raw_cj_page.new_shape()
_raw_cj_other.draw_line((50, 400), (350, 400))
_raw_cj_other.finish(color=(0.0, 0.0, 0.0)); _raw_cj_other.commit()
_raw_cj_doc.save(_raw_cj_pdf); _raw_cj_doc.close()
_raw_cj_detected = _detect_raw_cj(_raw_cj_pdf, 0.1)
ck("raw CJ reads its green legend swatch and measures only matching body vectors",
   _raw_cj_detected.get("zone", {}).get("category") == "construction_joint" and
   _raw_cj_detected["zone"]["length_lm"] == 20.0 and
   len(_raw_cj_detected["zone"]["polyline_segments"]) == 2 and
   "600 centres" in _raw_cj_detected["zone"]["joint_detail"],
   _raw_cj_detected)

_marked_cj_pdf = "/tmp/Ground Floor CJ QA.pdf"
_marked_cj_doc = _fitz_zones.open()
_marked_cj_page = _marked_cj_doc.new_page(width=500, height=300)
_marked_cj = _marked_cj_page.add_polyline_annot([(20, 50), (145, 50)])
_marked_cj.set_info(title="Fortel QA", subject="CJ internal", content="12.50 m")
_marked_cj.update()
_roller = _marked_cj_page.add_polyline_annot([(20, 150), (300, 150)])
_roller.set_info(title="Fortel QA", subject="Roller shutter detail", content="99.00 m")
_roller.update()
_marked_cj_doc.save(_marked_cj_pdf); _marked_cj_doc.close()
_marked_cj_read = _read_marked_zones(_marked_cj_pdf)
_marked_cj_zone = next(zone for zone in _marked_cj_read["zones"]
                       if zone["category"] == "construction_joint")
ck("explicit CJ is a measured Lm zone while roller-shutter detail stays out of scope",
   _marked_cj_zone["length_lm"] == 12.5 and
   _marked_cj_zone.get("slab_category") == "ground_floor" and
   any(zone["category"] == "other" for zone in _marked_cj_read["zones"]),
   _marked_cj_read["zones"])
_marked_cj_quote = generate_quotation({
    "file":"Ground Floor CJ QA.pdf", "source_discipline":"engineer",
    "zones":_marked_cj_read["zones"], "flags":[], "area_m2":None,
}, project="CJ QA")
_cj_measurement = next(item for item in _marked_cj_quote["measurements"]
                       if item["description"] == "Internal construction joint (CJ)")
ck("CJ reaches Ground-floor quotation as quantity-only Lm with blank assessor rate",
   _cj_measurement["section"] == "Ground floor slabs" and
   _cj_measurement["qty"] == 12.5 and _cj_measurement["assessor_rate_required"] and
   any("150x150x8 mm" in note for note in _marked_cj_quote["declarations"]),
   {"measurement":_cj_measurement,
    "declarations":_marked_cj_quote["declarations"]})

_upper_scope_pdf = "/tmp/ci_upper_floor_scopes.pdf"
_upper_scope_doc = _fitz_zones.open()
_upper_scope_page = _upper_scope_doc.new_page(width=700, height=300)
for _subject, _value, _x in (
        ("Unit 5 First Floor", 100.0, 20),
        ("Unit 5 Plant Deck", 30.0, 260),
        ("Unit 5 POD First Floor", 40.0, 500)):
    _annot = _upper_scope_page.add_polygon_annot([
        (_x, 20), (_x + 120, 20), (_x + 120, 150), (_x, 150),
    ])
    _annot.set_info(title="Fortel QA", subject=_subject,
                    content=f"Area\n{_value:.2f} sq m")
    _annot.update()
_upper_scope_doc.save(_upper_scope_pdf)
_upper_scope_doc.close()
_upper_scope_read = _read_marked_zones(_upper_scope_pdf)
ck("main upper floor, Plant deck and Unit-5 POD stay three explicit measured scopes",
   len(_upper_scope_read["zones"]) == 3 and
   {zone["boq_scope"] for zone in _upper_scope_read["zones"]} == {
       "main_upper_floor", "plant_deck", "pod_first_floor"} and
   all(zone["category"] == "upper_floor" and
       zone["boundary_rule"] == "measure to the edge of the metal decking"
       for zone in _upper_scope_read["zones"]), _upper_scope_read["zones"])
_upper_scope_quote = generate_quotation({
    "file":"Unit-5 Upper Floors.pdf", "source_discipline":"engineer",
    "zones":_upper_scope_read["zones"], "flags":[], "area_m2":170.0,
    "costing":{"area_m2":170.0, "rate":None, "total_gbp":None,
               "assumed":True, "spec":{}, "breakdown":{}, "extras":[]},
}, project="Upper Scope QA")
_upper_scope_items = [item for item in _upper_scope_quote["line_items"]
                      if item["section"] == "Upper floor slabs" and
                      item.get("line_role") == "concrete_slab"]
ck("Plant deck and POD are separate Upper-floor BOQ rows, never merged with main floor",
   len(_upper_scope_items) == 3 and
   any(item["description"].startswith("Plant deck —") for item in _upper_scope_items) and
   any(item["description"].startswith("POD first floor —") for item in _upper_scope_items),
   _upper_scope_items)

_unit4_pdf = "/tmp/Ground Floor Unit 4 Subunits.pdf"
_unit4_doc = _fitz_zones.open()
_unit4_page = _unit4_doc.new_page(width=700, height=300)
for _index, _subject in enumerate(("Unit 4A", "Unit 4B", "Unit 4C", "Unit 4D")):
    _x = 20 + _index * 160
    _annot = _unit4_page.add_polygon_annot([
        (_x, 20), (_x + 100, 20), (_x + 100, 120), (_x, 120),
    ])
    _annot.set_info(title="Fortel QA", subject=_subject, content="Area\n25.00 sq m")
    _annot.update()
_unit4_doc.save(_unit4_pdf)
_unit4_doc.close()
_unit4_read = _read_marked_zones(_unit4_pdf)
ck("complete Unit-4A/4B/4C/4D markup is one combined ground-floor slab zone",
   len(_unit4_read["zones"]) == 1 and
   _unit4_read["zones"][0]["category"] == "ground_floor" and
   _unit4_read["zones"][0]["area_m2"] == 100.0 and
   _unit4_read["zones"][0]["annotation_count"] == 4 and
   _unit4_read["zones"][0]["unit_label"] == "Unit 4 (4A-4D combined)" and
   not _unit4_read["unit_group_review_required"], _unit4_read)

_unit_zone_pdf = "/tmp/Yard Markup Unit Labels.pdf"
_unit_zone_doc = _fitz_zones.open()
_unit_zone_page = _unit_zone_doc.new_page(width=600, height=600)
for _subject, _value, _x in (("Unit-1&2", 235.37, 20), ("Unit-3", 133.79, 300)):
    _annot = _unit_zone_page.add_polygon_annot([
        (_x, 20), (_x + 200, 20), (_x + 200, 220), (_x, 220),
    ])
    _annot.set_info(title="Fortel QA", subject=_subject,
                    content=f"Area\n{_value:.2f} sq m")
    _annot.update()
_unit_zone_doc.save(_unit_zone_pdf)
_unit_zone_doc.close()
_unit_zone_read = _read_marked_zones(_unit_zone_pdf)
_unit_yard_zones = [zone for zone in _unit_zone_read["zones"]
                    if zone["category"] == "external_yard"]
ck("unit-labelled regions inherit a strong Yard filename context and preserve each area",
   len(_unit_yard_zones) == 2 and
   sorted(zone["area_m2"] for zone in _unit_yard_zones) == [133.79, 235.37] and
   sorted(zone["unit_label"] for zone in _unit_yard_zones) == ["Unit-1&2", "Unit-3"] and
   _unit_zone_read["area_m2"] == 369.2,
   _unit_zone_read)
_weak_unit_zone_pdf = "/tmp/General Markup Unit Labels.pdf"
shutil.copyfile(_unit_zone_pdf, _weak_unit_zone_pdf)
_weak_unit_read = _read_marked_zones(_weak_unit_zone_pdf)
ck("unit labels with weak context stay unclassified while every measured area is preserved",
   all(zone["category"] == "unclassified" and zone["needs_assessor"]
       for zone in _weak_unit_read["zones"]) and
   abs(sum(zone["area_m2"] for zone in _weak_unit_read["zones"]) - 369.16) < 0.001 and
   _weak_unit_read["area_m2"] == 369.2 and
   len(_weak_unit_read["flags"]) == 2,
   _weak_unit_read)

_zone_quote_results = []
for _unit_n in range(1, 5):
    _unit = _quotation_unit(f"Castle Unit-{_unit_n}.pdf", "External yard slabs", 1)
    _unit["zones"] = [
        {"category":"external_yard", "area_m2":100 * _unit_n, "perimeter_lm":10 * _unit_n},
        {"category":"dock", "area_m2":10 * _unit_n, "perimeter_lm":5 * _unit_n},
        {"category":"ground_floor", "area_m2":5 * _unit_n, "perimeter_lm":3 * _unit_n},
        {"category":"upper_floor", "area_m2":20 * _unit_n, "perimeter_lm":4 * _unit_n},
        {"category":"channel", "length_lm":7 * _unit_n},
        {"category":"transition", "length_lm":2 * _unit_n},
    ]
    _unit["brief_specs"] = {
        category: _empty_brief_spec(category)
        for category in ("external_yard", "dock", "ground_floor", "upper_floor")
    }
    _zone_quote_results.append(_unit)
_q_zones = generate_quotation(_zone_quote_results, project="Castle", client="Winvic",
                              ref="ZONE-001")
_zone_section_order = [
    "External yard slabs", "Dock slabs", "Ground floor slabs", "Upper floor slabs",
]
ck("mixed marked files allocate into all four BOQ sections",
   [spec["section"] for spec in _q_zones["specifications"]] == _zone_section_order)
ck("each BOQ section keeps four numeric Unit-N source rows",
   all([row["description"] for row in spec["area_rows"]] ==
       ["Unit-1", "Unit-2", "Unit-3", "Unit-4"]
       for spec in _q_zones["specifications"]), _q_zones["specifications"])
ck("aggregate job rate is never copied onto mixed zones",
   all(item.get("rate") is None and item.get("value") is None
       for item in _q_zones["line_items"]
       if item.get("line_role") == "concrete_slab"))
ck("channel, transition and zone perimeters remain unpriced Lm source quantities",
   any(m["description"] == FORTEL_CHANNEL_ROW and m["qty"] == 70
       for m in _q_zones["measurements"]) and
   any(m["description"] == FORTEL_TRANSITION_ROW and m["qty"] == 20
       for m in _q_zones["measurements"]) and
   all(m.get("assessor_rate_required") for m in _q_zones["measurements"]))
_zone_ws = _load_workbook(_BytesIO(quotation_xlsx(_q_zones)), data_only=False)["REV_01"]
_zone_xlsx_sections = [
    _expected_xlsx_sections[0], _expected_xlsx_sections[2],
    _expected_xlsx_sections[3], _expected_xlsx_sections[4],
]
_zone_section_labels = [
    _zone_ws.cell(row, 1).value for row in range(1, _zone_ws.max_row + 1)
    if _zone_ws.cell(row, 1).value in _zone_xlsx_sections
]
_channel_row = next(row for row in range(1, _zone_ws.max_row + 1)
                    if _zone_ws.cell(row, 1).value == FORTEL_CHANNEL_ROW)
ck("zone XLSX preserves section order with editable blank assessor rates",
   _zone_section_labels == _zone_xlsx_sections and
   _zone_ws.cell(_channel_row, 4).value is None and
   _zone_ws.cell(_channel_row, 5).value ==
   f'=IF(D{_channel_row}="","",ROUND(B{_channel_row}*D{_channel_row},2))')

_castle_dir = Path("drawings/castle_donington")
_castle_names = [
    *(f"External Markup Unit-{number}.pdf" for number in range(1, 5)),
    *(f"Office Floors Unit-{number}.pdf" for number in range(1, 5)),
]
try:
    for _castle_name in _castle_names:
        _require_fixture(_castle_dir / _castle_name, "Castle Donington client zone-gold checks")
    import json as _json_zone_gold
    from takeoff_pipeline import _zone_reference_flags as _zone_reference_flags_test
    _zone_gold = _json_zone_gold.loads(Path("gold.json").read_text())
    _castle_reads = {}
    for _castle_name in _castle_names:
        _castle_path = _castle_dir / _castle_name
        _castle_marked = _read_marked_zones(str(_castle_path))
        _castle_reads[_castle_name] = _castle_marked
        _entry = _zone_gold[str(_castle_path)]
        _actual = {z["category"]: z.get("area_m2") for z in _castle_marked["zones"]
                   if z.get("area_m2") is not None}
        _aggregate_delta = abs(_castle_marked["area_m2"] - _entry["net_m2"]) / _entry["net_m2"] * 100
        _zones_pass = all(
            category in _actual and abs(_actual[category] - expected) / expected * 100 <=
            _entry["zone_tol_pct"]
            for category, expected in _entry["zones_m2"].items()
        )
        ck(f"Castle zone gold: {_castle_name} aggregate + BOQ sections",
           _aggregate_delta <= _entry["tol_pct"] and _zones_pass,
           {"actual": _actual, "gold": _entry["zones_m2"]})
    _dock_perimeter = sum(
        next(z["perimeter_lm"] for z in _castle_reads[f"External Markup Unit-{n}.pdf"]["zones"]
             if z["category"] == "dock") for n in range(1, 5)
    )
    ck("Castle Dock polygon perimeter reproduces client BOQ 967 Lm",
       abs(_dock_perimeter - 967) / 967 * 100 <= 1, _dock_perimeter)
    _unit3_path = _castle_dir / "External Markup Unit-3.pdf"
    ck("Castle Unit-3 channel-vs-BOQ mismatch emits assessor flag",
       any("channel measured 545.36 Lm" in flag
           for flag in _zone_reference_flags_test(str(_unit3_path),
                                                  _castle_reads[_unit3_path.name]["zones"])))
except _FixtureNotPresent as _e:
    print(f"  [SKIP] {_e} — fixture not present")

print("raw external zone measurement — Yard/Dock split + low-confidence state cap")
import numpy as _np_multi_yard
from takeoff_unmarked import segment_hatch as _segment_hatch_multi_yard
_multi_yard_image = _np_multi_yard.full((100, 100, 3), 255, dtype=_np_multi_yard.uint8)
_multi_yard_image[10:25, 10:30] = (180, 180, 180)   # 300 m2 primary
_multi_yard_image[45:55, 10:22] = (180, 180, 180)   # 120 m2 second unit (<200)
_multi_yard_image[70:72, 80:84] = (180, 180, 180)   # 8 m2 legend chip
_multi_yard_diag = {}
_multi_yard_mask = _segment_hatch_multi_yard(
    _multi_yard_image, (180, 180, 180), tol=0, close=1, k=1.0, S=1.0,
    exclude_border=False, legend_exclusion_bbox=[79, 69, 85, 73],
    full_rgb=True, _diag=_multi_yard_diag)
ck("same-tint segmentation retains a real second unit below the single-yard 200 m2 floor",
   int(_multi_yard_mask.sum()) == 420 and
   len(_multi_yard_diag.get("_retained_component_masks", [])) == 2,
   {"pixels": int(_multi_yard_mask.sum()),
    "components": _multi_yard_diag.get("component_candidates")})
ck("matched legend chip is geometrically excluded, never promoted as another Yard",
   not _multi_yard_mask[70:72, 80:84].any(),
    {"excluded_legend_m2": _multi_yard_diag.get("excluded_legend_m2"),
     "components": _multi_yard_diag.get("component_candidates")})

# Raster boundaries can be serrated by internal linework.  Perimeter is permitted only when
# an independently encoded closed CAD path strongly overlaps the segmented tint.  This fixture
# has two disjoint regions and deliberately nicks raster pixels along one edge; neither expected
# perimeter is inferred from an area or a client value.
import fitz as _fitz_native_boundary
from takeoff_unmarked import _native_boundary_for_mask as _native_boundary_for_mask_test

def _closed_line_drawing(_points):
    _closed = _points + [_points[0]]
    return {
        "rect": _fitz_native_boundary.Rect(
            min(p.x for p in _points), min(p.y for p in _points),
            max(p.x for p in _points), max(p.y for p in _points)),
        "items": [("l", start, end) for start, end in zip(_closed, _closed[1:])],
        "closePath": True,
    }

class _BoundaryPage:
    rotation_matrix = _fitz_native_boundary.Matrix(1, 1)
    def __init__(self, drawings):
        self._drawings = drawings
    def get_drawings(self):
        return self._drawings

_boundary_a = [_fitz_native_boundary.Point(10, 10), _fitz_native_boundary.Point(60, 10),
               _fitz_native_boundary.Point(60, 40), _fitz_native_boundary.Point(10, 40)]
_boundary_b = [_fitz_native_boundary.Point(90, 55), _fitz_native_boundary.Point(140, 55),
               _fitz_native_boundary.Point(140, 90), _fitz_native_boundary.Point(120, 90),
               _fitz_native_boundary.Point(120, 75), _fitz_native_boundary.Point(90, 75)]
_boundary_page = _BoundaryPage([
    _closed_line_drawing(_boundary_a), _closed_line_drawing(_boundary_b)])
_boundary_mask_a = _np_multi_yard.zeros((120, 170), dtype=bool)
_boundary_mask_a[10:41, 10:61] = True
_boundary_mask_a[10:12, 20:25] = False  # raster-only nick: must not become extra perimeter
_boundary_mask_b = _np_multi_yard.zeros((120, 170), dtype=_np_multi_yard.uint8)
import cv2 as _cv2_native_boundary
_cv2_native_boundary.fillPoly(
    _boundary_mask_b, [_np_multi_yard.array(_boundary_b, dtype=_np_multi_yard.int32)], 1)
_native_a, _native_a_reason = _native_boundary_for_mask_test(
    _boundary_page, _boundary_mask_a, S=1.0, k=0.1)
_native_b, _native_b_reason = _native_boundary_for_mask_test(
    _boundary_page, _boundary_mask_b.astype(bool), S=1.0, k=0.1)
ck("native CAD perimeter ignores raster/internal-edge serration on each disjoint region",
   _native_a is not None and _native_b is not None and
   abs(_native_a["perimeter_lm"] - 16.0) < 0.01 and
   abs(_native_b["perimeter_lm"] - 17.0) < 0.01,
   {"region_a": _native_a or _native_a_reason,
    "region_b": _native_b or _native_b_reason})
_unresolved_boundary, _unresolved_reason = _native_boundary_for_mask_test(
    _BoundaryPage([]), _boundary_mask_a, S=1.0, k=0.1)
ck("perimeter refuses when no corroborating native boundary exists",
   _unresolved_boundary is None and "no explicit closed" in _unresolved_reason,
   _unresolved_reason)

from takeoff_unmarked import (
    _transition_candidates_from_surface_mask as _transition_candidates_from_mask_test,
    _native_boundary_stream_budget as _native_boundary_stream_budget_test,
    MAX_NATIVE_BOUNDARY_STREAM_BYTES as _native_boundary_stream_limit_test,
)
ck("dense native-vector pages refuse optional perimeter parsing before the robustness timeout",
   _native_boundary_stream_budget_test(_native_boundary_stream_limit_test) and
   not _native_boundary_stream_budget_test(_native_boundary_stream_limit_test + 1))
_transition_surface = _np_multi_yard.zeros((100, 160), dtype=bool)
_transition_surface[41:46, 10:61] = True
_transition_yard = [{
    "region_id": "yard-region-1",
    "polygon_pts": [[10, 10], [60, 10], [60, 40], [10, 40]],
    "perimeter_confidence": "high",
}]
_transition_prefills, _transition_prefill_reasons = \
    _transition_candidates_from_mask_test(
        _transition_yard, _transition_surface, k=0.1, S=1.0)
ck("legend-surface adjacency creates one assisted Transition prefill outside measured zones",
   len(_transition_prefills) == 1 and
   _transition_prefills[0]["proposed_length_lm"] == 5.0 and
   _transition_prefills[0]["assumed"] is True and
   "length_lm" not in _transition_prefills[0] and
   _transition_prefills[0]["category"] == "transition",
   _transition_prefills or _transition_prefill_reasons)
_transition_ambiguous_surface = _transition_surface.copy()
_transition_ambiguous_surface[4:10, 10:61] = True
_transition_ambiguous, _transition_ambiguous_reasons = \
    _transition_candidates_from_mask_test(
        _transition_yard, _transition_ambiguous_surface, k=0.1, S=1.0)
ck("two disjoint adjacent surface runs refuse instead of guessing a Yard entrance",
   not _transition_ambiguous and
   any("disjoint" in reason for reason in _transition_ambiguous_reasons),
   _transition_ambiguous_reasons)

# Source: Fortel's Yard Markup.pdf supplied for 2165 Tanro Voltage Business Park:
# Unit-1&2 = 235.37 m² and Unit-3 = 133.79 m².  Production sees only a temporary
# annotation-stripped copy; these client answers stay here in the validation assertion.
try:
    import glob as _glob_tanro_regions
    import tempfile as _tempfile_tanro_regions
    from accuracy_report import strip_annotations as _strip_tanro_regions
    from pathlib import Path as _Path_tanro_regions
    import takeoff_unmarked as _takeoff_unmarked_tanro_regions
    _tanro_matches = _glob_tanro_regions.glob(
        "drawings/aryan_drive/**/2165 Tanro- Voltage Business Park/Markup/Yard Markup.pdf",
        recursive=True,
    )
    if not _tanro_matches:
        raise _FixtureNotPresent("Tanro Yard Markup.pdf")
    with _tempfile_tanro_regions.TemporaryDirectory() as _tanro_tmp:
        _tanro_raw = _Path_tanro_regions(_tanro_tmp) / "Yard raw.pdf"
        _strip_tanro_regions(_Path_tanro_regions(_tanro_matches[0]), _tanro_raw)
        _tanro_result = _takeoff_unmarked_tanro_regions.takeoff(str(_tanro_raw))
        _tanro_regions = _tanro_result.get("yard_regions") or []
        _tanro_areas = sorted(region.get("area_m2") for region in _tanro_regions)
        _tanro_truth = sorted([235.37, 133.79])
        ck("Tanro raw Yard separates both client unit regions within 5%",
           len(_tanro_areas) == 2 and all(
               abs(actual - expected) / expected * 100 <= 5
               for actual, expected in zip(_tanro_areas, _tanro_truth)
           ), {"actual": _tanro_areas, "truth": _tanro_truth})
        ck("every retained raw Yard region surfaces its own area, bbox and perimeter",
           all(region.get("area_m2") and region.get("bbox_pdf_pts")
               and region.get("perimeter_lm") for region in _tanro_regions),
           _tanro_regions)
        _tanro_actual = sorted(
            (region["area_m2"], region["perimeter_lm"]) for region in _tanro_regions)
        _tanro_perimeter_truth = sorted([(235.37, 61.54), (133.79, 49.80)])
        ck("Tanro raw Yard native per-region perimeters are within 5%",
           len(_tanro_actual) == 2 and all(
               abs(actual_perimeter - truth_perimeter) / truth_perimeter * 100 <= 5
               for (_, actual_perimeter), (_, truth_perimeter)
               in zip(_tanro_actual, _tanro_perimeter_truth)
           ), {"actual": _tanro_actual, "truth": _tanro_perimeter_truth})
        _tanro_transition_truth = sorted([14.25, 12.30])
        _tanro_transition_candidates = _tanro_result.get("transition_candidates") or []
        _tanro_transition_actual = sorted(
            candidate.get("proposed_length_lm")
            for candidate in _tanro_transition_candidates)
        ck("Tanro raw macadam/Yard adjacency prefills both client Transition runs within 5%",
           len(_tanro_transition_actual) == 2 and all(
               abs(actual - expected) / expected * 100 <= 5
               for actual, expected in zip(
                   _tanro_transition_actual, _tanro_transition_truth)
           ), {"actual": _tanro_transition_actual,
               "truth": _tanro_transition_truth})
        ck("raw Transition prefills never leak into measured zones or the measured total",
           all(candidate.get("assumed") is True and
               "length_lm" not in candidate
               for candidate in _tanro_transition_candidates) and
           not any(zone.get("category") == "transition"
                   for zone in (_tanro_result.get("zones") or [])),
           {"candidates": _tanro_transition_candidates,
            "zones": _tanro_result.get("zones")})
except _FixtureNotPresent as _e:
    print(f"  [SKIP] {_e} — fixture not present")
try:
    _external_marked_paths = [
        _castle_dir / f"External Markup Unit-{number}.pdf" for number in range(1, 5)
    ]
    for _path in _external_marked_paths:
        _require_fixture(_path, "Castle Donington raw Yard/Dock validation")
    from accuracy_report import strip_annotations as _strip_annotations_external
    from takeoff_pipeline import takeoff as _pipeline_takeoff_external
    import takeoff_unmarked as _takeoff_unmarked_external
    import os as _os_external

    _raw_external_results = {}
    try:
        for _unit_number, _marked_path in enumerate(_external_marked_paths, start=1):
            _raw_path = Path(f"/tmp/ci_external_unit_{_unit_number}_stripped.pdf")
            _strip_annotations_external(_marked_path, _raw_path)
            _truth = _read_marked_zones(str(_marked_path))
            _truth_by_zone = {
                zone["category"]: zone
                for zone in _truth["zones"]
            }
            _measured = _pipeline_takeoff_external(
                str(_raw_path), send_approval=False, auto_extract_spec=False)
            _raw_external_results[_unit_number] = _measured
            _measured_by_zone = {
                zone["category"]: zone
                for zone in _measured.get("zones", [])
            }
            _yard_ok = (
                "external_yard" in _measured_by_zone
                and abs(
                    _measured_by_zone["external_yard"]["area_m2"]
                    - _truth_by_zone["external_yard"]["area_m2"]
                ) / _truth_by_zone["external_yard"]["area_m2"] * 100 <= 5
            )
            _dock_ok = (
                "dock" in _measured_by_zone
                and abs(
                    _measured_by_zone["dock"]["area_m2"]
                    - _truth_by_zone["dock"]["area_m2"]
                ) / _truth_by_zone["dock"]["area_m2"] * 100 <= 5
            )
            _zone_sum = sum(
                zone.get("area_m2") or 0
                for zone in _measured.get("zones", [])
                if zone.get("category") in ("external_yard", "dock")
            )
            ck(f"raw External Unit-{_unit_number}: one Yard + one Dock, each within 5%",
               _yard_ok and _dock_ok and
               [zone["category"] for zone in _measured.get("zones", [])].count(
                   "external_yard") == 1 and
               [zone["category"] for zone in _measured.get("zones", [])].count("dock") == 1,
               {"truth": {key: value.get("area_m2")
                          for key, value in _truth_by_zone.items()},
                "measured": {key: value.get("area_m2")
                             for key, value in _measured_by_zone.items()}})
            ck(f"raw External Unit-{_unit_number}: zone total reconciles exactly, no double count",
               abs(_zone_sum - _measured["zones_total_area_m2"]) < 0.01
               and _measured["area_m2"] == _measured_by_zone["external_yard"]["area_m2"],
               {"zones": _zone_sum,
                "zone_total": _measured["zones_total_area_m2"],
                "legacy_yard_area": _measured["area_m2"]})
    finally:
        for _unit_number in range(1, 5):
            try:
                _os_external.remove(f"/tmp/ci_external_unit_{_unit_number}_stripped.pdf")
            except FileNotFoundError:
                pass

    ck("raw channel assumptions stay outside measured zones and accuracy totals",
       all(
           not any(zone.get("category") in ("channel", "transition")
                   for zone in result.get("zones", []))
           and len(result.get("channel_proposals", [])) == 2
           and {proposal.get("component") for proposal in result["channel_proposals"]} ==
               {"dock_retaining_wall", "yard_longest_contained_run"}
           and all(proposal.get("assumed") is True
                   and proposal.get("requires_assessor_confirmation") is True
                   and len(proposal.get("polyline_pts", [])) == 2
                   and (abs(proposal["polyline_pts"][0][0] -
                            proposal["polyline_pts"][1][0]) <= 0.01
                        or abs(proposal["polyline_pts"][0][1] -
                               proposal["polyline_pts"][1][1]) <= 0.01)
                   for proposal in result["channel_proposals"])
           and any("Channel MEASUREMENT not attempted" in flag
                   for flag in result.get("flags", []))
           for result in _raw_external_results.values()
       ))

    _marked_with_real_channel = _pipeline_takeoff_external(
        str(_external_marked_paths[0]), send_approval=False, auto_extract_spec=False)
    ck("real marked Channel linework wins; no assumed proposal is created",
       any(zone.get("category") == "channel" and zone.get("length_lm")
           for zone in _marked_with_real_channel.get("zones", []))
       and not _marked_with_real_channel.get("channel_proposals"),
       {"zones": _marked_with_real_channel.get("zones"),
        "proposals": _marked_with_real_channel.get("channel_proposals")})

    _dock_only_proposals, _dock_only_flags = _takeoff_unmarked_external.propose_channels(
        [], {
            "loading_face_lm": 42.0,
            "loading_face_pts": [[10.0, 20.0], [10.0, 62.0]],
        }, 1.0, scale_verified=True)
    ck("unconfident Yard geometry refuses its run instead of guessing",
       len(_dock_only_proposals) == 1
       and _dock_only_proposals[0]["component"] == "dock_retaining_wall"
       and any("Yard run refused" in flag for flag in _dock_only_flags),
       {"proposals": _dock_only_proposals, "flags": _dock_only_flags})

    _channel_enabled_before = _takeoff_unmarked_external.CHANNEL_PROPOSALS_ENABLED
    try:
        _takeoff_unmarked_external.CHANNEL_PROPOSALS_ENABLED = False
        _disabled_proposals, _disabled_flags = _takeoff_unmarked_external.propose_channels(
            [[0,0],[100,0],[100,100],[0,100]], {
                "loading_face_lm": 42.0,
                "loading_face_pts": [[10.0,20.0],[10.0,62.0]],
            }, 1.0, scale_verified=True)
    finally:
        _takeoff_unmarked_external.CHANNEL_PROPOSALS_ENABLED = _channel_enabled_before
    ck("one gate disables every channel proposal without touching measurement",
       not _disabled_proposals
       and _disabled_flags == [
           "CHANNEL PROPOSALS disabled by CHANNEL_PROPOSALS_ENABLED"],
       {"proposals": _disabled_proposals, "flags": _disabled_flags})

    # Direct takeoff proves the state cap is caused by the non-grey swatch fallback,
    # independently of Unit 3's downstream portal/job handling.
    _unit3_raw = Path("/tmp/ci_external_unit_3_state_cap.pdf")
    try:
        _strip_annotations_external(_external_marked_paths[2], _unit3_raw)
        _unit3_direct = _takeoff_unmarked_external.takeoff(str(_unit3_raw))
        ck("non-grey/white swatch fallback can never reach MEASURED_VERIFIED",
           _unit3_direct.get("region_confidence") == "low" and
           _unit3_direct.get("measurement_state") == "MEASURED_UNVERIFIED" and
           _unit3_direct.get("needs_assessor") is True,
           {"confidence": _unit3_direct.get("region_confidence"),
            "state": _unit3_direct.get("measurement_state")})
    finally:
        try:
            _os_external.remove(_unit3_raw)
        except FileNotFoundError:
            pass

    # An ambiguous native loading-face signal must remain visible in the zone
    # contract and must cap approval; it cannot be silently folded into Yard.
    _ambiguous_raw = Path("/tmp/ci_external_ambiguous_dock.pdf")
    _real_dock_detector = _takeoff_unmarked_external.detect_raw_dock_zone
    try:
        _strip_annotations_external(_external_marked_paths[0], _ambiguous_raw)
        _takeoff_unmarked_external.detect_raw_dock_zone = lambda *args, **kwargs: {
            "zone": None,
            "reason": "multiple plausible loading faces",
            "evidence_seen": True,
        }
        _ambiguous_result = _takeoff_unmarked_external.takeoff(str(_ambiguous_raw))
        ck("ambiguous raw Dock stays unclassified + flagged and cannot be VERIFIED",
           any(zone.get("category") == "unclassified"
               and zone.get("needs_assessor") is True
               for zone in _ambiguous_result.get("zones", []))
           and _ambiguous_result.get("measurement_state") == "MEASURED_UNVERIFIED"
           and any("classify/trace Dock zone" in flag
                   for flag in _ambiguous_result.get("flags", [])),
           {"zones": _ambiguous_result.get("zones"),
            "state": _ambiguous_result.get("measurement_state")})
    finally:
        _takeoff_unmarked_external.detect_raw_dock_zone = _real_dock_detector
        try:
            _os_external.remove(_ambiguous_raw)
        except FileNotFoundError:
            pass
except _FixtureNotPresent as _e:
    print(f"  [SKIP] {_e} — fixture not present")

print("raw multi-region Yard fixture + assessor keep/exclude loop")
try:
    _tanro_marked = Path(
        "drawings/inderjit_markups_31jul/2165 Tanro- Voltage Business Park/Markup/Yard Markup.pdf")
    _require_fixture(_tanro_marked, "Tanro multi-region Yard validation")
    from accuracy_report import strip_annotations as _strip_tanro_annotations
    from takeoff_pipeline import takeoff as _takeoff_tanro_multi
    _tanro_raw = Path("/tmp/ci_tanro_multi_region_stripped.pdf")
    try:
        _strip_tanro_annotations(_tanro_marked, _tanro_raw)
        _tanro_result = _takeoff_tanro_multi(
            str(_tanro_raw), send_approval=False, auto_extract_spec=False)
        _tanro_regions = _tanro_result.get("yard_regions", [])
        ck("Tanro raw drawing retains both distant unit Yards as separate visible regions",
           len(_tanro_regions) == 2 and all(region.get("bbox_pdf_pts")
                                            for region in _tanro_regions),
           [{"area_m2":region.get("area_m2"), "bbox":region.get("bbox_pdf_pts")}
            for region in _tanro_regions])
        ck("Tanro candidate total sums both retained regions within 5% of client truth",
           abs(_tanro_result.get("area_m2", 0) - 369.2) / 369.2 * 100 <= 5 and
           _tanro_result.get("yard_region_review_required") is True,
           {"measured": _tanro_result.get("area_m2"), "truth": 369.2})
    finally:
        try:
            _tanro_raw.unlink()
        except FileNotFoundError:
            pass
except _FixtureNotPresent as _e:
    print(f"  [SKIP] {_e} — fixture not present")

print("assessor Yard-region exclusion endpoint")
import tempfile as _tempfile_yard_review
import approval_server as _AS_yard_review
_yard_review_tmp = Path(_tempfile_yard_review.mkdtemp(prefix="ci_yard_regions_"))
_yard_review_jobs_before = _AS_yard_review.JOBS_FILE
_yard_review_rates_before = _AS_yard_review.CLIENT_RATES_FILE
try:
    _AS_yard_review.JOBS_FILE = _yard_review_tmp / "jobs.json"
    _AS_yard_review.CLIENT_RATES_FILE = _yard_review_tmp / "client_rates.json"
    _yard_review_id = "bbbbbbbb-bbbb-4bbb-8bbb-bbbbbbbbbbbb"
    _yard_review_regions = [
        {"region_id":"yard-region-1", "area_m2":234.2, "included":True,
         "bbox_pdf_pts":[10,10,30,25], "polygon_pts":[[10,10],[30,10],[30,25],[10,25]]},
        {"region_id":"yard-region-2", "area_m2":132.0, "included":True,
         "bbox_pdf_pts":[60,45,72,55], "polygon_pts":[[60,45],[72,45],[72,55],[60,55]]},
    ]
    _yard_review_zones = [{
        "zone_key":"external_yard", "category":"external_yard", "area_m2":366.0,
        "measurement_kind":"area", "needs_assessor":True,
    }]
    _AS_yard_review.save_jobs({_yard_review_id: {
        "id":_yard_review_id, "status":"pending", "decision":None,
        "area_m2":366.0, "measurement_state":"MEASURED_UNVERIFIED",
        "scale_k":0.1, "yard_regions":_yard_review_regions,
        "yard_region_review_required":True, "zones":_yard_review_zones,
        "flags":["YARD REGION REVIEW REQUIRED: synthetic"],
        "result":{"file":"Multi Yard.pdf", "area_m2":366.0, "scale_k":0.1,
                  "measurement_state":"MEASURED_UNVERIFIED",
                  "yard_regions":_yard_review_regions,
                  "yard_region_review_required":True, "zones":_yard_review_zones,
                  "flags":["YARD REGION REVIEW REQUIRED: synthetic"]},
    }})
    _yard_review_client = _AS_yard_review.app.test_client()
    _partial_yard_review = _yard_review_client.post(
        f"/yard-regions/{_yard_review_id}",
        json={"decisions":[{"region_id":"yard-region-1", "action":"keep"}]})
    ck("Yard-region review must cover every retained component, never silently omit one",
       _partial_yard_review.status_code == 409, _partial_yard_review.get_json())
    _complete_yard_review = _yard_review_client.post(
        f"/yard-regions/{_yard_review_id}", json={"decisions":[
            {"region_id":"yard-region-1", "action":"keep"},
            {"region_id":"yard-region-2", "action":"exclude"},
        ]})
    _reviewed_yard_job = _AS_yard_review.load_jobs()[_yard_review_id]
    ck("assessor exclusion removes exactly that component from Yard and its BOQ zone",
       _complete_yard_review.status_code == 200 and
       _complete_yard_review.get_json().get("area_m2") == 234.2 and
       _reviewed_yard_job["area_m2"] == 234.2 and
       _reviewed_yard_job["zones"][0]["area_m2"] == 234.2 and
       _reviewed_yard_job["yard_regions"][1]["included"] is False,
       {"response":_complete_yard_review.get_json(),
        "regions":_reviewed_yard_job.get("yard_regions")})
    ck("completed Yard-region review clears only its gate and preserves four-state review",
       not _reviewed_yard_job.get("yard_region_review_required") and
       "same-tint Yard regions" not in
       (_AS_yard_review._approve_block_reason(_reviewed_yard_job) or "") and
       _reviewed_yard_job.get("measurement_state") == "MEASURED_UNVERIFIED",
       _AS_yard_review._approve_block_reason(_reviewed_yard_job))
finally:
    _AS_yard_review.JOBS_FILE = _yard_review_jobs_before
    _AS_yard_review.CLIENT_RATES_FILE = _yard_review_rates_before
    shutil.rmtree(_yard_review_tmp, ignore_errors=True)

print("channel proposal geometry — retaining-wall adjacent and never diagonal")
from takeoff_unmarked import (
    _longest_contained_yard_run as _channel_yard_run,
    propose_channels as _propose_channels_axis,
)
_channel_rect = [[0,0],[1000,0],[1000,600],[0,600]]
_horizontal_wall = [
    {"polyline_pts":[[0,1],[1000,1]], "grey_wall_evidence":False},
    {"polyline_pts":[[0,1],[1000,1]], "grey_wall_evidence":False},
]
_horizontal_run, _horizontal_refusal = _channel_yard_run(
    _channel_rect, 0.1, wall_segments=_horizontal_wall)
ck("corroborated horizontal retaining-wall run is emitted exactly axis-aligned",
   _horizontal_run is not None and _horizontal_refusal is None and
   _horizontal_run["length_lm"] == 100.0 and
   _horizontal_run["polyline_pts"][0][1] == _horizontal_run["polyline_pts"][1][1],
   (_horizontal_run, _horizontal_refusal))
_vertical_wall = [
    {"polyline_pts":[[1,0],[1,600]], "grey_wall_evidence":True},
]
_vertical_run, _vertical_refusal = _channel_yard_run(
    _channel_rect, 0.1, wall_segments=_vertical_wall)
ck("corroborated vertical retaining-wall run is emitted exactly axis-aligned",
   _vertical_run is not None and _vertical_refusal is None and
   _vertical_run["length_lm"] == 60.0 and
   _vertical_run["polyline_pts"][0][0] == _vertical_run["polyline_pts"][1][0],
   (_vertical_run, _vertical_refusal))
_diagonal_run, _diagonal_refusal = _channel_yard_run(
    _channel_rect, 0.1, wall_segments=[
        {"polyline_pts":[[0,0],[1000,600]], "grey_wall_evidence":True},
        {"polyline_pts":[[0,0],[1000,600]], "grey_wall_evidence":True},
    ])
ck("diagonal-only evidence refuses the Yard component instead of falling back",
   _diagonal_run is None and "non-diagonal" in _diagonal_refusal,
   _diagonal_refusal)
_diagonal_dock_proposals, _diagonal_dock_flags = _propose_channels_axis(
    _channel_rect,
    {"loading_face_lm":100.0, "loading_face_pts":[[0,0],[1000,600]]},
    0.1, wall_segments=_horizontal_wall)
ck("diagonal Dock loading-face evidence refuses the whole assumption, not just the Yard run",
   not _diagonal_dock_proposals and
   any("not a straight non-diagonal" in flag for flag in _diagonal_dock_flags),
   _diagonal_dock_flags)
_ambiguous_wall_run, _ambiguous_wall_refusal = _channel_yard_run(
    _channel_rect, 0.1, wall_segments=[
        {"polyline_pts":[[0,1],[1000,1]], "grey_wall_evidence":False},
        {"polyline_pts":[[0,1],[1000,1]], "grey_wall_evidence":False},
        {"polyline_pts":[[0,599],[970,599]], "grey_wall_evidence":False},
        {"polyline_pts":[[0,599],[970,599]], "grey_wall_evidence":False},
    ])
ck("near-equal competing retaining-wall runs refuse rather than guessing an edge",
   _ambiguous_wall_run is None and "within 5%" in _ambiguous_wall_refusal,
   _ambiguous_wall_refusal)
_no_dock_proposals, _no_dock_flags = _propose_channels_axis(
    _channel_rect, None, 0.1, scale_verified=True,
    dock_presence="absent", wall_segments=_horizontal_wall)
ck("unit with no Dock level gets exactly one full-Yard-width assumed channel",
   len(_no_dock_proposals) == 1 and
   _no_dock_proposals[0]["component"] == "yard_longest_contained_run" and
   _no_dock_proposals[0]["channel_case"] == "no_dock_level_one_run" and
   any("one full-Yard-width" in flag for flag in _no_dock_flags),
   (_no_dock_proposals, _no_dock_flags))
_with_dock_proposals, _with_dock_flags = _propose_channels_axis(
    _channel_rect,
    {"loading_face_lm":60.0, "loading_face_pts":[[1,0],[1,600]]},
    0.1, scale_verified=True, dock_presence="present",
    wall_segments=_horizontal_wall, access_road_interruption=True)
ck("unit with a Dock level gets dock-level + full-width runs and access-road assumption",
   len(_with_dock_proposals) == 2 and
   all(proposal["channel_case"] == "with_dock_level_two_runs"
       for proposal in _with_dock_proposals) and
   _with_dock_proposals[1]["access_road_interruption"] is True and
   any("split/edit" in reason
       for reason in _with_dock_proposals[1]["confidence_reasons"]),
   (_with_dock_proposals, _with_dock_flags))
_ambiguous_dock_proposals, _ambiguous_dock_flags = _propose_channels_axis(
    _channel_rect, None, 0.1, dock_presence="ambiguous",
    wall_segments=_horizontal_wall)
ck("ambiguous Dock presence refuses instead of silently choosing one- or two-run rule",
   not _ambiguous_dock_proposals and
   any("one-channel versus two-channel" in flag for flag in _ambiguous_dock_flags),
   _ambiguous_dock_flags)

try:
    _office_marked_paths = [
        _castle_dir / f"Office Floors Unit-{number}.pdf" for number in range(1, 5)
    ]
    _office_stripped_paths = [
        _castle_dir / "_stripped" / path.name for path in _office_marked_paths
    ]
    for _path in _office_marked_paths + _office_stripped_paths:
        _require_fixture(_path, "Castle Donington stripped-office validation")
    import json as _json_office_gold
    from takeoff_pipeline import takeoff as _pipeline_takeoff_office
    _office_gold = _json_office_gold.loads(Path("gold.json").read_text())
    _outside_gate = []
    for _marked_path, _stripped_path in zip(_office_marked_paths, _office_stripped_paths):
        with _fitz_zones.open(_marked_path) as _marked_doc, _fitz_zones.open(_stripped_path) as _stripped_doc:
            _marked_annots = sum(1 for page in _marked_doc for _ in (page.annots() or []))
            _stripped_annots = sum(1 for page in _stripped_doc for _ in (page.annots() or []))
        ck(f"stripped Office fixture removes every annotation: {_marked_path.name}",
           _marked_annots > 0 and _stripped_annots == 0,
           {"marked": _marked_annots, "stripped": _stripped_annots})

        _assisted = _pipeline_takeoff_office(
            str(_stripped_path), send_approval=False, auto_extract_spec=False)
        _candidate_levels = [
            candidate["level"] for candidate in _assisted.get("candidate_polygons", [])
        ]
        with _fitz_zones.open(_stripped_path) as _office_title_doc:
            from office_candidates import _level_titles as _office_level_titles
            _expected_levels = sorted(
                title["level"] for title in _office_level_titles(_office_title_doc[0]))
        ck(f"every Office level appears exactly once (no duplicate rows): {_marked_path.name}",
           sorted(_candidate_levels) == _expected_levels and
           len(_candidate_levels) == len(set(_candidate_levels)),
           {"expected": _expected_levels, "actual": _candidate_levels})

        _level_areas = {}
        for _candidate in _assisted.get("candidate_polygons", []):
            _candidate_area = 0.0
            _regions = _candidate.get("regions") or [_candidate.get("polygon_pts", [])]
            _region_holes = _candidate.get("region_holes") or [[] for _ in _regions]
            for _region_index, _region in enumerate(_regions):
                if len(_region) < 3:
                    continue
                _region_area, _ = measure_regions(
                    [_region], _assisted["scale_k"],
                    holes={0: _region_holes[_region_index]})
                _candidate_area += _region_area
            _level = _candidate["level"]
            _level_areas[_level] = _candidate_area
        _detected_total = round(sum(_level_areas.values()), 2)
        _markup_total = _read_marked_zones(str(_marked_path))["area_m2"]
        _boq_total = _office_gold[str(_marked_path)]["net_m2"]
        _delta_pct = (_detected_total - _markup_total) / _markup_total * 100
        _outside_gate.append(abs(_delta_pct) > 5)
        ck(f"Office candidates remain assisted geometry only: {_marked_path.name}",
           _assisted.get("area_m2") is None and
           _assisted.get("measurement_state") == "UNMEASURED" and
           _assisted.get("needs_assessor") is True and
           _assisted.get("costing") is None and _assisted.get("polygon_pts") is None and
           bool(_assisted.get("candidate_polygons")) and
           all("area_m2" not in candidate
               for candidate in _assisted.get("candidate_polygons", [])),
           {"diagnostic_candidate_total": _detected_total, "markup": _markup_total,
            "boq": _boq_total, "delta_pct": round(_delta_pct, 2)})
    ck("Office auto-measure bar remains closed unless every unit is within 5%",
       len(_outside_gate) == 4 and any(_outside_gate), _outside_gate)
except _FixtureNotPresent as _e:
    print(f"  [SKIP] {_e} — fixture not present")

